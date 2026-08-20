"""可编辑规则和测试案例 Excel 模板。

本空白模板保留给未来人工复核或自定义新版本使用。项目当前可直接使用代码内的
资料综合版 V1.0，因此运行现有版本不再依赖老师填写此表。
"""

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .daily_color_rule_config import (
    PUBLIC_FACTOR_CODE_TO_FIELD,
    RELATION_CODE_TO_FIELD,
    SOLAR_TERMS,
    ElementRelationScores,
    PublicFactorWeights,
    PublicRuleConfiguration,
    TendencyThresholds,
)


FACTOR_LABELS = {
    "PUBLIC_YEAR_STEM": "年干",
    "PUBLIC_YEAR_BRANCH": "年支主五行",
    "PUBLIC_MONTH_STEM": "月干",
    "PUBLIC_MONTH_BRANCH": "月支主五行",
    "PUBLIC_DAY_STEM": "日干",
    "PUBLIC_DAY_BRANCH": "日支主五行",
    "PUBLIC_SOLAR_TERM": "当前节气",
}

RELATION_LABELS = {
    "SAME_ELEMENT": "候选五行与参考五行相同",
    "CANDIDATE_GENERATES_REFERENCE": "候选五行生参考五行",
    "REFERENCE_GENERATES_CANDIDATE": "参考五行生候选五行",
    "CANDIDATE_CONTROLS_REFERENCE": "候选五行克参考五行",
    "REFERENCE_CONTROLS_CANDIDATE": "参考五行克候选五行",
}


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254A3C")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_rule_workbook() -> Workbook:
    """构造一份空白规则工作簿，所有专业值必须由老师填写。"""
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    guide_rows = [
        ("用途", "本表用于把人工复核或自定义的公共五色规则转成程序可执行配置。"),
        ("重要", "示例文档中的+15、+20、-15不是正式权重，本模板不预填任何分数。"),
        ("时间口径", "全部按北京时间；技术标识Asia/Shanghai；午夜换日；不做真太阳时修正。"),
        ("审核要求", "至少填写并确认10个测试案例，记录审核人、审核时间和资料依据后，才可进入expert_reviewed。"),
        ("个人规则", "个人规则问题与测试案例单独填写；公共规则通过不代表个人规则自动通过。"),
        ("启用方式", "填写完成后交给后台校验。未通过校验或仍为test_only的配置不能用于正式发布。"),
    ]
    _header(guide, ["项目", "说明"])
    for row in guide_rows:
        guide.append(row)
    _fit(guide, [20, 100])

    version = workbook.create_sheet("版本信息")
    _header(version, ["字段", "填写值", "说明"])
    for row in [
        ("版本号", "", "例如 wuse-public-v1.0；内容修改必须升级版本"),
        ("状态", "test_only", "test_only / source_reviewed / expert_reviewed / active"),
        ("生效日期", "", "active 时必填，格式 YYYY-MM-DD"),
        ("算法说明", "", "说明采用的统一流派、因素和总体计算原则"),
        ("专业依据", "", "每行一条，写明老师口径或资料依据"),
        ("审核人", "", "expert_reviewed 或 active 时必填"),
        ("审核时间", "", "格式 YYYY-MM-DD HH:MM:SS"),
    ]:
        version.append(row)
    _fit(version, [22, 44, 76])

    factors = workbook.create_sheet("公共因子权重")
    _header(factors, ["因子编号", "因子名称", "权重", "专业说明"])
    for code in PUBLIC_FACTOR_CODE_TO_FIELD:
        factors.append((code, FACTOR_LABELS[code], "", "由规则制定者填写；七项权重合计必须为100"))
    _fit(factors, [32, 22, 14, 70])

    relations = workbook.create_sheet("五行关系分值")
    _header(relations, ["关系编号", "关系含义", "基础分", "专业说明"])
    for code in RELATION_CODE_TO_FIELD:
        relations.append((code, RELATION_LABELS[code], "", "由规则制定者统一定义，不按具体日期临时改变"))
    _fit(relations, [40, 38, 14, 65])

    terms = workbook.create_sheet("节气五行映射")
    _header(terms, ["节气", "五行", "专业说明"])
    for term in SOLAR_TERMS:
        terms.append((term, "", "填写金、木、水、火、土之一"))
    _fit(terms, [18, 14, 75])

    ties = workbook.create_sheet("同分排序")
    _header(ties, ["优先顺序", "颜色", "说明"])
    for rank in range(1, 6):
        ties.append((rank, "", "五色得分相同时按本顺序处理；五种颜色必须各填一次"))
    _fit(ties, [16, 18, 78])

    thresholds = workbook.create_sheet("趋势阈值")
    _header(thresholds, ["趋势编号", "用户层含义", "分数下限", "说明"])
    for row in [
        ("strong_support_min", "今天很顺", "", "最高档下限"),
        ("support_min", "比较合适", "", "第二档下限"),
        ("balanced_min", "平稳一般", "", "第三档下限"),
        ("caution_min", "会比较累", "", "第四档下限；低于此值为成效偏弱"),
    ]:
        thresholds.append(row)
    _fit(thresholds, [28, 24, 16, 68])

    cases = workbook.create_sheet("公共测试案例")
    _header(cases, ["序号", "目标日期", "预期第1名", "预期第2名", "预期第3名", "预期第4名", "预期第5名", "老师确认", "确认人", "说明"])
    for index in range(1, 21):
        cases.append((index, "", "", "", "", "", "", "否", "", "至少确认10条；颜色必须完整且不重复"))
    _fit(cases, [10, 18, 16, 16, 16, 16, 16, 16, 18, 48])

    personal = workbook.create_sheet("个人规则待确认")
    _header(personal, ["问题编号", "必须确认的问题", "老师填写", "技术说明"])
    questions = [
        ("P01", "个人算法采用扶抑、调候还是其他统一口径？", "禁止把多套体系临时混用"),
        ("P02", "日主强弱如何计算，月令、藏干和各柱权重是多少？", "需要明确公式和边界"),
        ("P03", "个人结构与当日五行的生、克、泄、耗分别如何计分？", "需要完整关系分值表"),
        ("P04", "未知出生时辰时删去哪部分规则，结果措辞降低到什么粒度？", "程序固定不补造时柱"),
        ("P05", "个人五色同分时如何排序？", "必须给出确定性顺序"),
        ("P06", "穿搭、办事、沟通、香品是否共用一套排名？", "若不共用需要拆分场景规则"),
    ]
    for code, question, note in questions:
        personal.append((code, question, "", note))
    _fit(personal, [16, 65, 55, 55])

    personal_cases = workbook.create_sheet("个人测试案例")
    _header(personal_cases, ["序号", "测试档案编号", "出生日期", "出生时间", "是否已知时辰", "目标日期", "预期五色顺序", "老师确认", "确认人", "说明"])
    for index in range(1, 21):
        personal_cases.append((index, f"CASE-{index:02d}", "", "", "", "", "", "否", "", "只用虚构测试档案，不填写真实用户姓名或手机号"))
    _fit(personal_cases, [10, 20, 18, 18, 18, 18, 35, 16, 18, 55])
    return workbook


def build_rule_template() -> bytes:
    """生成可下载或写入 docs 的 Excel 二进制内容。"""
    stream = io.BytesIO()
    build_rule_workbook().save(stream)
    return stream.getvalue()


def _required_int(value, label: str) -> int:
    """读取老师填写的整数，并把空值/小数转换为明确错误。"""
    if value is None or str(value).strip() == "":
        raise ValueError(f"{label}不能为空")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}必须是整数") from exc
    if isinstance(value, float) and value != number:
        raise ValueError(f"{label}必须是整数")
    return number


def _optional_date(value, label: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"{label}格式必须是YYYY-MM-DD") from exc


def _optional_datetime(value, label: str) -> datetime | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"{label}格式必须是YYYY-MM-DD HH:MM:SS") from exc


def parse_rule_template(content: bytes) -> PublicRuleConfiguration:
    """解析填写后的 Excel，并交给强类型配置模型做完整专业流程校验。"""
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError("规则文件不是可读取的.xlsx工作簿") from exc

    required_sheets = {
        "版本信息", "公共因子权重", "五行关系分值", "节气五行映射",
        "同分排序", "趋势阈值", "公共测试案例",
    }
    missing = required_sheets - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"规则文件缺少工作表：{', '.join(sorted(missing))}")

    version_values = {
        str(row[0]).strip(): row[1]
        for row in workbook["版本信息"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    version = str(version_values.get("版本号") or "").strip()
    if not version:
        raise ValueError("版本号不能为空")
    status = str(version_values.get("状态") or "test_only").strip()
    algorithm_summary = str(version_values.get("算法说明") or "").strip()
    if not algorithm_summary:
        raise ValueError("算法说明不能为空")
    references = [line.strip() for line in str(version_values.get("专业依据") or "").splitlines() if line.strip()]

    factor_rows = {
        str(row[0]).strip(): row[2]
        for row in workbook["公共因子权重"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    if set(factor_rows) != set(PUBLIC_FACTOR_CODE_TO_FIELD):
        raise ValueError("公共因子权重表的因子编号不完整或包含未知编号")
    factor_values = {
        field: _required_int(factor_rows[code], f"{code}权重")
        for code, field in PUBLIC_FACTOR_CODE_TO_FIELD.items()
    }

    relation_rows = {
        str(row[0]).strip(): row[2]
        for row in workbook["五行关系分值"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    if set(relation_rows) != set(RELATION_CODE_TO_FIELD):
        raise ValueError("五行关系分值表的关系编号不完整或包含未知编号")
    relation_values = {
        field: _required_int(relation_rows[code], f"{code}基础分")
        for code, field in RELATION_CODE_TO_FIELD.items()
    }

    term_values = {
        str(row[0]).strip(): str(row[1] or "").strip()
        for row in workbook["节气五行映射"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    if set(term_values) != set(SOLAR_TERMS):
        raise ValueError("节气五行映射必须完整包含二十四节气")
    for term, element in term_values.items():
        if element not in {"金", "木", "水", "火", "土"}:
            raise ValueError(f"{term}的五行必须填写金、木、水、火、土之一")

    tie_rows = sorted(
        (row for row in workbook["同分排序"].iter_rows(min_row=2, values_only=True) if row[0] is not None),
        key=lambda row: _required_int(row[0], "同分优先顺序"),
    )
    tie_order = [str(row[1] or "").strip() for row in tie_rows]

    threshold_rows = {
        str(row[0]).strip(): row[2]
        for row in workbook["趋势阈值"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    threshold_names = {"strong_support_min", "support_min", "balanced_min", "caution_min"}
    if set(threshold_rows) != threshold_names:
        raise ValueError("趋势阈值编号不完整或包含未知编号")
    thresholds = {name: _required_int(threshold_rows[name], name) for name in threshold_names}

    approved_cases = 0
    for row_number, row in enumerate(workbook["公共测试案例"].iter_rows(min_row=2, values_only=True), start=2):
        if str(row[7] or "").strip() != "是":
            continue
        colors = [str(value or "").strip() for value in row[2:7]]
        if set(colors) != {"白金", "绿金", "黑金", "红金", "黄金"}:
            raise ValueError(f"公共测试案例第{row_number}行的预期五色必须完整且不重复")
        if not row[1]:
            raise ValueError(f"公共测试案例第{row_number}行目标日期不能为空")
        if not str(row[8] or "").strip():
            raise ValueError(f"公共测试案例第{row_number}行确认人不能为空")
        approved_cases += 1

    return PublicRuleConfiguration(
        version=version,
        status=status,
        effective_from=_optional_date(version_values.get("生效日期"), "生效日期"),
        algorithm_summary=algorithm_summary,
        factor_weights=PublicFactorWeights(**factor_values),
        relation_scores=ElementRelationScores(**relation_values),
        solar_term_elements=term_values,
        tie_break_color_order=tie_order,
        tendency_thresholds=TendencyThresholds(**thresholds),
        professional_references=references,
        approved_test_case_count=approved_cases,
        reviewed_by=str(version_values.get("审核人") or "").strip() or None,
        reviewed_at=_optional_datetime(version_values.get("审核时间"), "审核时间"),
    )
