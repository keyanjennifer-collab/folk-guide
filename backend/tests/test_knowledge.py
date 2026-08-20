import io
from pathlib import Path

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.database import Base, SessionLocal, engine
from app.knowledge_routes import STORAGE_ROOT
from app.main import app
from app.models import KnowledgeChunk, KnowledgeDocument


ADMIN_HEADERS = {"X-Admin-Key": "dev-admin-key", "X-Admin-Name": "knowledge-reviewer"}


def cleanup_hash(file_hash: str) -> None:
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        document = db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.file_hash == file_hash))
        if document:
            path = STORAGE_ROOT.parent / document.storage_key
            db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id).delete()
            db.delete(document)
            db.commit()
            if path.exists():
                path.unlink()


def upload(client: TestClient, filename: str, content: bytes, mime: str = "text/plain"):
    return client.post(
        "/api/admin/knowledge/documents",
        headers=ADMIN_HEADERS,
        data={
            "title": "二十四节气测试资料",
            "source_name": "项目原创测试",
            "author": "测试作者",
            "copyright_status": "original",
        },
        files={"file": (filename, content, mime)},
    )


def text_pdf_bytes(text: str) -> bytes:
    """构造一个只含基础拉丁文字层的最小PDF，避免测试依赖外部生成工具。"""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content_stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content_stream)).encode("ascii") + b" >>\nstream\n" + content_stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(payload)


def test_markdown_upload_parse_review_search_disable():
    content = ("# 二十四节气\n\n节气是观察太阳周年运动形成的时间知识。\n\n"
               "## 立春\n\n立春表示春季相关时间节律的开始。" * 20).encode("utf-8")
    import hashlib
    digest = hashlib.sha256(content).hexdigest()
    cleanup_hash(digest)
    with TestClient(app) as client:
        uploaded = upload(client, "solar-terms.md", content, "text/markdown")
        assert uploaded.status_code == 200
        document_id = uploaded.json()["id"]
        assert uploaded.json()["status"] == "uploaded"

        duplicate = upload(client, "again.md", content, "text/markdown")
        assert duplicate.status_code == 409

        parsed = client.post(f"/api/admin/knowledge/documents/{document_id}/parse", headers=ADMIN_HEADERS)
        assert parsed.status_code == 200
        assert parsed.json()["status"] == "reviewing"
        assert parsed.json()["chunk_count"] >= 2

        chunks = client.get(f"/api/admin/knowledge/documents/{document_id}/chunks", headers=ADMIN_HEADERS)
        assert chunks.status_code == 200
        first_chunk = chunks.json()[0]
        assert first_chunk["heading"] in {"二十四节气", "立春"}

        updated = client.put(
            f"/api/admin/knowledge/chunks/{first_chunk['id']}",
            headers=ADMIN_HEADERS,
            json={"heading": "节气概述", "content": first_chunk["content"] + "\n人工审核补充。"},
        )
        assert updated.status_code == 200
        assert updated.json()["review_status"] == "pending"

        approved = client.post(
            f"/api/admin/knowledge/documents/{document_id}/approve",
            headers=ADMIN_HEADERS,
            json={"approved_chunk_ids": None},
        )
        assert approved.status_code == 200
        assert approved.json()["approved_chunk_count"] == approved.json()["chunk_count"]

        results = client.get("/api/admin/knowledge/search", headers=ADMIN_HEADERS, params={"q": "立春 节气"})
        assert results.status_code == 200
        assert results.json()
        assert results.json()[0]["source_name"] == "项目原创测试"

        blocked_edit = client.put(
            f"/api/admin/knowledge/chunks/{first_chunk['id']}",
            headers=ADMIN_HEADERS,
            json={"content": "不允许直接修改"},
        )
        assert blocked_edit.status_code == 409

        disabled = client.post(f"/api/admin/knowledge/documents/{document_id}/disable", headers=ADMIN_HEADERS)
        assert disabled.status_code == 200
        assert client.get("/api/admin/knowledge/search", headers=ADMIN_HEADERS, params={"q": "立春"}).json() == []
    cleanup_hash(digest)


def test_docx_and_txt_parsers():
    docx = DocxDocument()
    docx.add_heading("道德经资料", level=1)
    docx.add_paragraph("上善若水。水善利万物而不争。")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "原文"
    table.cell(0, 1).text = "来源"
    table.cell(1, 0).text = "上善若水"
    table.cell(1, 1).text = "第八章"
    stream = io.BytesIO()
    docx.save(stream)

    contents = [
        ("dao.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("notes.txt", "五行色彩资料。\n\n传统文化内容仅供了解。".encode("utf-8"), "text/plain"),
    ]
    import hashlib
    with TestClient(app) as client:
        for filename, content, mime in contents:
            digest = hashlib.sha256(content).hexdigest()
            cleanup_hash(digest)
            response = upload(client, filename, content, mime)
            assert response.status_code == 200
            document_id = response.json()["id"]
            parsed = client.post(f"/api/admin/knowledge/documents/{document_id}/parse", headers=ADMIN_HEADERS)
            assert parsed.status_code == 200
            assert parsed.json()["chunk_count"] >= 1
            cleanup_hash(digest)


def test_text_pdf_parse_preserves_page_metadata():
    content = text_pdf_bytes("Spring timing and five element color knowledge for verified PDF extraction.")
    import hashlib
    digest = hashlib.sha256(content).hexdigest()
    cleanup_hash(digest)
    with TestClient(app) as client:
        uploaded = upload(client, "classic-text.pdf", content, "application/pdf")
        assert uploaded.status_code == 200
        document_id = uploaded.json()["id"]

        parsed = client.post(
            f"/api/admin/knowledge/documents/{document_id}/parse",
            headers=ADMIN_HEADERS,
            params={"parser_mode": "native"},
        )
        assert parsed.status_code == 200
        assert parsed.json()["page_count"] == 1
        assert parsed.json()["extraction_method"] == "pdf_native"
        assert parsed.json()["needs_ocr"] is False

        chunks = client.get(
            f"/api/admin/knowledge/documents/{document_id}/chunks", headers=ADMIN_HEADERS
        ).json()
        assert chunks
        assert chunks[0]["page_start"] == 1
        assert chunks[0]["page_end"] == 1
        assert chunks[0]["extraction_method"] == "pdf_native"
        assert "five element color" in chunks[0]["content"]
    cleanup_hash(digest)


def test_pdf_without_text_layer_is_marked_ocr_required():
    import hashlib
    import pypdf

    stream = io.BytesIO()
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(stream)
    content = stream.getvalue()
    digest = hashlib.sha256(content).hexdigest()
    cleanup_hash(digest)
    with TestClient(app) as client:
        uploaded = upload(client, "scanned-classic.pdf", content, "application/pdf")
        document_id = uploaded.json()["id"]
        parsed = client.post(
            f"/api/admin/knowledge/documents/{document_id}/parse", headers=ADMIN_HEADERS
        )
        assert parsed.status_code == 422
        assert "需要OCR" in parsed.json()["detail"]

        document = client.get(
            f"/api/admin/knowledge/documents/{document_id}", headers=ADMIN_HEADERS
        ).json()
        assert document["status"] == "ocr_required"
        assert document["page_count"] == 1
        assert document["needs_ocr"] is True
        assert document["approved_chunk_count"] == 0
    cleanup_hash(digest)


def test_low_confidence_ocr_chunk_requires_manual_correction():
    import hashlib

    content = ("五行古籍OCR测试正文。" * 50).encode("utf-8")
    digest = hashlib.sha256(content).hexdigest()
    cleanup_hash(digest)
    with TestClient(app) as client:
        uploaded = upload(client, "ocr-review.txt", content)
        document_id = uploaded.json()["id"]
        assert client.post(
            f"/api/admin/knowledge/documents/{document_id}/parse", headers=ADMIN_HEADERS
        ).status_code == 200

        with SessionLocal() as db:
            chunk = db.scalar(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id))
            chunk.extraction_method = "pdf_mineru"
            chunk.ocr_confidence = 0.42
            chunk_id = chunk.id
            db.commit()

        blocked = client.post(
            f"/api/admin/knowledge/documents/{document_id}/approve",
            headers=ADMIN_HEADERS,
            json={"approved_chunk_ids": None},
        )
        assert blocked.status_code == 422
        assert "OCR置信度过低" in blocked.json()["detail"]

        corrected = client.put(
            f"/api/admin/knowledge/chunks/{chunk_id}",
            headers=ADMIN_HEADERS,
            json={"content": "人工逐字核对后的五行古籍正文。"},
        )
        assert corrected.status_code == 200
        assert corrected.json()["extraction_method"] == "manual_review"
        assert corrected.json()["ocr_confidence"] is None
        approved = client.post(
            f"/api/admin/knowledge/documents/{document_id}/approve",
            headers=ADMIN_HEADERS,
            json={"approved_chunk_ids": None},
        )
        assert approved.status_code == 200
    cleanup_hash(digest)


def test_upload_rejects_unsupported_file_and_bad_copyright():
    with TestClient(app) as client:
        unsupported = upload(client, "unsafe.pdf", b"not a pdf", "application/pdf")
        assert unsupported.status_code == 422
        bad_copyright = client.post(
            "/api/admin/knowledge/documents",
            headers=ADMIN_HEADERS,
            data={"title": "测试", "source_name": "测试", "copyright_status": "unknown"},
            files={"file": ("safe.txt", b"content", "text/plain")},
        )
        assert bad_copyright.status_code == 422


def test_knowledge_admin_page_is_available():
    with TestClient(app) as client:
        response = client.get("/admin/knowledge")
        assert response.status_code == 200
        assert "AI国学知识库后台" in response.text
