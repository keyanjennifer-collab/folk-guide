"""把最新项目交接 Markdown 生成便于阅读和归档的 Word 版本。

Markdown 是新聊天最稳定的机器可读交接源，Word 供项目负责人日常查看。生成器只读取
项目内的脱敏报告，不读取 .env，也不会把任何密钥写入文档。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "五色知时小程序_新聊天项目交接与继续开发报告_V1.0.md"
# 当前Windows执行环境对Python直接创建中文docx文件名存在编码兼容问题。
# Word正文仍全部使用中文，文件名改用稳定ASCII，避免生成过程因路径编码失败。
OUTPUT = ROOT / "docs" / "WuseZhishi_NewChat_Project_Handoff_V1.0.docx"


def set_cell_shading(cell, fill: str) -> None:
    """设置表格单元格底色。"""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_repeat_table_header(row) -> None:
    """让长表格跨页时重复显示表头。"""
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_run_font(run, name: str, size: float | None = None, color: str | None = None) -> None:
    """同时设置西文字体和中文字体，避免Word回退到不一致的字体。"""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str) -> None:
    """处理报告中常用的粗体和行内代码；其余Markdown保持为普通文字。"""
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    for part in token_pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_run_font(run, "微软雅黑")
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.5, "244A3F")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "微软雅黑")


def add_page_number(paragraph) -> None:
    """插入Word自动页码域。"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run_font(run, "微软雅黑", 9, "7A817C")


def configure_document(document: Document) -> None:
    """统一纸张、正文、标题、页眉和页脚样式。"""
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.35

    heading_colors = {1: "183F35", 2: "385F4D", 3: "8B7444"}
    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string(heading_colors[level])
        style.font.size = Pt({1: 17, 2: 14, 3: 12}[level])
        style.font.bold = True
        style.paragraph_format.space_before = Pt({1: 15, 2: 11, 3: 8}[level])
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("五色知时小程序｜新聊天项目交接与继续开发报告 V1.0")
    set_run_font(header_run, "微软雅黑", 8.5, "75817B")
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, title: str) -> None:
    """创建简洁的东方绿金封面。"""
    document.add_paragraph()
    document.add_paragraph()
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker_run = kicker.add_run("WUSE ZHISHI · PROJECT HANDOFF")
    set_run_font(kicker_run, "Georgia", 10, "9A8354")

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title.replace("｜", "\n"))
    set_run_font(title_run, "宋体", 27, "183F35")
    title_run.bold = True
    title_paragraph.paragraph_format.space_before = Pt(30)
    title_paragraph.paragraph_format.space_after = Pt(24)

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule_run = rule.add_run("五色应时，知时而行")
    set_run_font(rule_run, "宋体", 14, "8B7444")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("当前代码状态 · 继续开发路线 · 上线边界\n2026年8月20日（北京时间）")
    set_run_font(meta_run, "微软雅黑", 10, "66716B")
    meta.paragraph_format.space_before = Pt(90)

    document.add_page_break()


def parse_table(lines: list[str], start: int, document: Document) -> int:
    """读取一段Markdown表格并插入Word表格，返回下一行位置。"""
    table_lines: list[str] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1

    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
    if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        for raw in table_lines:
            paragraph = document.add_paragraph()
            add_inline(paragraph, raw)
        return index

    headers = rows[0]
    body = rows[2:]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for position, value in enumerate(headers):
        cell = table.rows[0].cells[position]
        set_cell_shading(cell, "DCE7DF")
        paragraph = cell.paragraphs[0]
        add_inline(paragraph, value)
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("183F35")

    for row_values in body:
        cells = table.add_row().cells
        for position in range(len(headers)):
            value = row_values[position] if position < len(row_values) else ""
            add_inline(cells[position].paragraphs[0], value)
    document.add_paragraph()
    return index


def markdown_to_docx(markdown: str) -> Document:
    """把本报告使用的Markdown子集转换为Word。"""
    lines = markdown.splitlines()
    title = lines[0].removeprefix("# ").strip()
    document = Document()
    configure_document(document)
    add_cover(document, title)
    index = 1
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.45)
                paragraph.paragraph_format.right_indent = Cm(0.25)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(7)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F0F3EF")
                paragraph._p.get_or_add_pPr().append(shading)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, "Consolas", 9, "244A3F")
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            index = parse_table(lines, index, document)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_heading(level=level)
            add_inline(paragraph, heading.group(2))
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(7)
            set_cell = OxmlElement("w:shd")
            set_cell.set(qn("w:fill"), "EFE9DC")
            paragraph._p.get_or_add_pPr().append(set_cell)
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped.rstrip("  "))
        index += 1

    return document


def main() -> None:
    """读取交接源、生成Word，并输出文件路径。"""
    markdown = SOURCE.read_text(encoding="utf-8")
    document = markdown_to_docx(markdown)
    document.core_properties.title = "五色知时小程序｜新聊天项目交接与继续开发报告 V1.0"
    document.core_properties.subject = "项目状态、代码入口、待办路线、上线与安全边界"
    document.core_properties.author = "Codex"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
