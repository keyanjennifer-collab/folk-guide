r"""生成资料综合版每日五色 Excel、Word 与固定测试样本。

运行方式（项目根目录）：

    backend\.venv\Scripts\python.exe tools\generate_daily_color_research_v1.py

脚本只使用虚构档案，不读取数据库、.env、API Key、AppSecret 或真实用户信息。
"""

from __future__ import annotations

import json
import sys
from dataclasses import asdict
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
DOCS = ROOT / "docs"
sys.path.insert(0, str(BACKEND))

from app.calendar_service import apply_calendar_calculation  # noqa: E402
from app.daily_color_context import build_personal_rule_input, build_public_rule_input  # noqa: E402
from app.daily_color_personal_engine import calculate_personal_rule  # noqa: E402
from app.daily_color_research_v1 import (  # noqa: E402
    PERSONAL_RESEARCH_CONFIG,
    PUBLIC_RESEARCH_CONFIG,
)
from app.daily_color_rule_engine import calculate_public_rule  # noqa: E402
from app.models import BirthProfile  # noqa: E402
from app.schemas import BirthProfileInput  # noqa: E402


SAMPLE_START_DATE = date(2026, 8, 19)
XLSX_PATH = DOCS / "今日五色资料综合规则与样本_V1.0.xlsx"
DOCX_PATH = DOCS / "今日五色资料综合算法与参数说明_V1.0.docx"
JSON_PATH = DOCS / "今日五色资料综合固定测试样本_V1.0.json"

HEADER_FILL = PatternFill("solid", fgColor="254A3C")
SECTION_FILL = PatternFill("solid", fgColor="DDE8DF")


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fit(sheet, widths: list[int]) -> None:
    """设置列宽、自动换行和首行冻结。"""
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _make_profile(*, birth_date: date, time_known: bool, birth_time: str | None) -> BirthProfile:
    """创建不落库的虚构档案，用于可复现样本。"""
    payload = BirthProfileInput(
        birth_date=birth_date,
        time_known=time_known,
        birth_time=birth_time,
        birth_city="杭州",
        gender="unspecified",
    )
    profile = BirthProfile(user_id=999999, profile_version=1, **payload.model_dump())
    apply_calendar_calculation(profile, payload)
    return profile


def build_samples() -> dict:
    """生成未来7天公共样本和两组未来3天个人样本。"""
    public_samples = []
    for offset in range(7):
        target = SAMPLE_START_DATE + timedelta(days=offset)
        rule_input = build_public_rule_input(target, PUBLIC_RESEARCH_CONFIG.version)
        calculation = calculate_public_rule(rule_input, PUBLIC_RESEARCH_CONFIG)
        public_samples.append({
            "target_date": target.isoformat(),
            "pillars": {
                "year": rule_input.calendar.pillars.year.text,
                "month": rule_input.calendar.pillars.month.text,
                "day": rule_input.calendar.pillars.day.text,
            },
            "solar_term": rule_input.calendar.solar_term_on_day or rule_input.calendar.previous_solar_term.name,
            "ranking": [item.model_dump(mode="json") for item in calculation.result.ranking.items],
        })

    profiles = [
        ("CASE-FOUR-PILLARS", date(1995, 6, 18), True, "14:30"),
        ("CASE-THREE-PILLARS", date(1982, 11, 3), False, None),
    ]
    personal_samples = []
    for case_code, birthday, time_known, birth_time in profiles:
        profile = _make_profile(birth_date=birthday, time_known=time_known, birth_time=birth_time)
        for offset in range(3):
            target = SAMPLE_START_DATE + timedelta(days=offset)
            rule_input = build_personal_rule_input(profile, target, PERSONAL_RESEARCH_CONFIG.version)
            calculation = calculate_personal_rule(
                rule_input,
                PERSONAL_RESEARCH_CONFIG,
                PUBLIC_RESEARCH_CONFIG,
            )
            personal_samples.append({
                "case_code": case_code,
                "birth_date": birthday.isoformat(),
                "birth_time": birth_time,
                "precision_mode": rule_input.precision_mode,
                "target_date": target.isoformat(),
                "birth_pillars": rule_input.birth_context.pillars.model_dump(mode="json"),
                "element_distribution": calculation.element_distribution,
                "strength": asdict(calculation.strength),
                "ranking": [item.model_dump(mode="json") for item in calculation.result.ranking.items],
            })

    return {
        "sample_start_date": SAMPLE_START_DATE.isoformat(),
        "timezone": "Asia/Shanghai",
        "time_standard": "北京时间",
        "public_rule_version": PUBLIC_RESEARCH_CONFIG.version,
        "personal_rule_version": PERSONAL_RESEARCH_CONFIG.version,
        "public_config_fingerprint": PUBLIC_RESEARCH_CONFIG.fingerprint(),
        "personal_config_fingerprint": PERSONAL_RESEARCH_CONFIG.fingerprint(),
        "public_7_days": public_samples,
        "personal_3_days": personal_samples,
    }


def build_excel(samples: dict) -> Workbook:
    """生成已填写、可直接检查的研究版工作簿。"""
    workbook = Workbook()
    version = workbook.active
    version.title = "版本说明"
    _header(version, ["字段", "内容"])
    rows = [
        ("公共规则版本", PUBLIC_RESEARCH_CONFIG.version),
        ("个人规则版本", PERSONAL_RESEARCH_CONFIG.version),
        ("状态", "source_reviewed（资料已核对；不是专家背书，也不是active正式版本）"),
        ("时间口径", "北京时间 / Asia/Shanghai / 午夜换日"),
        ("公共算法", PUBLIC_RESEARCH_CONFIG.algorithm_summary),
        ("个人算法", PERSONAL_RESEARCH_CONFIG.algorithm_summary),
        ("公共配置指纹", PUBLIC_RESEARCH_CONFIG.fingerprint()),
        ("个人配置指纹", PERSONAL_RESEARCH_CONFIG.fingerprint()),
        ("重要边界", "古籍没有每日穿衣颜色百分制。权重、分值、倍率、阈值和融合比例是本项目工程参数。"),
    ]
    for row in rows:
        version.append(row)
    _fit(version, [26, 115])

    sources = workbook.create_sheet("文献与边界")
    _header(sources, ["资料", "本版本采用的原则", "是否直接提供数值", "处理方式"])
    source_rows = [
        ("《黄帝内经·素问·阴阳应象大论》", "木—青/苍、火—赤、土—黄、金—白、水—黑的五色对应", "否", "转为绿金、红金、黄金、白金、黑金的产品命名"),
        ("《五行大义》卷二", "相生相克、干支五行、四时旺相休囚死", "否", "关系可考；分值与倍率由项目版本化"),
        ("《子平真诠评注》", "专求月令，同时结合年日时根气判断强弱", "否", "月支设为个人柱权重最高单项"),
        ("《滴天髓》通行本", "扶抑得宜、损益取中", "否", "只约束策略方向，不冒充原典数值"),
        ("《协纪辨方书》", "历法、月建、择日义例", "否", "本版不把书中宜忌直接换算成颜色分"),
    ]
    for row in source_rows:
        sources.append(row)
    _fit(sources, [32, 60, 22, 62])

    factors = workbook.create_sheet("公共因子权重")
    _header(factors, ["因子", "权重", "层级", "说明"])
    factor_labels = {
        "PUBLIC_YEAR_STEM": ("年干", "年层"),
        "PUBLIC_YEAR_BRANCH": ("年支主五行", "年层"),
        "PUBLIC_MONTH_STEM": ("月干", "月令季节层"),
        "PUBLIC_MONTH_BRANCH": ("月支主五行", "月令季节层"),
        "PUBLIC_DAY_STEM": ("日干", "日层"),
        "PUBLIC_DAY_BRANCH": ("日支主五行", "日层"),
        "PUBLIC_SOLAR_TERM": ("当前节气", "月令季节层"),
    }
    for code, weight in PUBLIC_RESEARCH_CONFIG.factor_weights.by_code().items():
        label, layer = factor_labels[code]
        factors.append((code, weight, layer, f"{label}；权重是工程参数"))
    _fit(factors, [34, 14, 24, 60])

    relations = workbook.create_sheet("公共关系分值")
    _header(relations, ["关系", "分值", "产品解释", "性质"])
    relation_explanations = {
        "SAME_ELEMENT": "候选色与环境同类",
        "CANDIDATE_GENERATES_REFERENCE": "候选色生环境，个人输出较多",
        "REFERENCE_GENERATES_CANDIDATE": "环境生候选色，支持度最高",
        "CANDIDATE_CONTROLS_REFERENCE": "候选色克环境，可用但不列最高",
        "REFERENCE_CONTROLS_CANDIDATE": "环境克候选色，列为谨慎",
    }
    for code, score in PUBLIC_RESEARCH_CONFIG.relation_scores.by_code().items():
        relations.append((code, score, relation_explanations[code], "项目工程分值，非古籍原数值"))
    _fit(relations, [42, 14, 52, 46])

    terms = workbook.create_sheet("节气映射")
    _header(terms, ["节气", "五行", "月建说明"])
    for term, element in PUBLIC_RESEARCH_CONFIG.solar_term_elements.items():
        terms.append((term, element, "按寅卯木、辰土、巳午火、未土、申酉金、戌土、亥子水、丑土归类"))
    _fit(terms, [16, 14, 82])

    pillar = workbook.create_sheet("个人柱权重")
    _header(pillar, ["位置编号", "权重", "说明"])
    for code, weight in PERSONAL_RESEARCH_CONFIG.pillar_weights.by_code().items():
        pillar.append((code, weight, "未知时辰时删除PERSON_TIME_STEM/BRANCH，再将其余结构归一化"))
    _fit(pillar, [34, 14, 80])

    hidden = workbook.create_sheet("藏干比例")
    _header(hidden, ["地支", "藏干", "五行", "比例", "性质"])
    from app.daily_color_schemas import HEAVENLY_STEM_ELEMENT
    for branch, shares in PERSONAL_RESEARCH_CONFIG.hidden_stems.items():
        for item in shares:
            hidden.append((branch, item.stem, HEAVENLY_STEM_ELEMENT[item.stem], item.share, "常见启发式工程比例；古籍无统一百分数"))
    _fit(hidden, [14, 14, 14, 14, 62])

    seasons = workbook.create_sheet("旺衰倍率")
    _header(seasons, ["状态", "倍率", "关系定义", "性质"])
    descriptions = {
        "旺": "与当月主气相同",
        "相": "当月主气所生",
        "休": "生当月主气",
        "囚": "克当月主气",
        "死": "被当月主气所克",
    }
    for state, multiplier in PERSONAL_RESEARCH_CONFIG.seasonal_multipliers.by_state().items():
        seasons.append((state, multiplier, descriptions[state], "状态名称有传统依据；倍率为工程参数"))
    _fit(seasons, [14, 14, 42, 54])

    strategy = workbook.create_sheet("强弱与策略")
    _header(strategy, ["状态", "判断区间", "生我", "同我", "我生", "我克", "克我", "说明"])
    configs = PERSONAL_RESEARCH_CONFIG.strategy_scores
    for regime, interval, score in [
        ("偏弱", "支持比 < 42%", configs.weak),
        ("相对平衡", "42% ≤ 支持比 ≤ 58%", configs.balanced),
        ("偏强", "支持比 > 58%", configs.strong),
    ]:
        strategy.append((regime, interval, score.resource, score.peer, score.output, score.wealth, score.officer, "角色分均为工程参数"))
    _fit(strategy, [18, 30, 14, 14, 14, 14, 14, 42])

    fusion = workbook.create_sheet("融合与阈值")
    _header(fusion, ["参数", "数值", "说明"])
    fusion_rows = [
        ("出生结构权重", f"{PERSONAL_RESEARCH_CONFIG.birth_structure_weight}%", "个人五行需要分"),
        ("公共环境权重", f"{PERSONAL_RESEARCH_CONFIG.public_environment_weight}%", "当天公共五色分"),
        ("补缺目标", PERSONAL_RESEARCH_CONFIG.balance_target_percent, "每种五行的均衡参考占比"),
        ("补缺系数", PERSONAL_RESEARCH_CONFIG.deficiency_weight, "(20%-实际占比)×系数"),
        ("补缺修正限幅", PERSONAL_RESEARCH_CONFIG.deficiency_adjustment_limit, "防止极端结构把角色策略完全覆盖"),
        ("strong_support下限", PERSONAL_RESEARCH_CONFIG.tendency_thresholds.strong_support_min, "内部趋势，不直接承诺吉凶"),
        ("support下限", PERSONAL_RESEARCH_CONFIG.tendency_thresholds.support_min, "内部趋势"),
        ("balanced下限", PERSONAL_RESEARCH_CONFIG.tendency_thresholds.balanced_min, "内部趋势"),
        ("caution下限", PERSONAL_RESEARCH_CONFIG.tendency_thresholds.caution_min, "低于此值为restrained"),
    ]
    for row in fusion_rows:
        fusion.append(row)
    _fit(fusion, [32, 20, 80])

    public_cases = workbook.create_sheet("公共7日样本")
    _header(public_cases, ["日期", "年柱", "月柱", "日柱", "节气", "第1名", "分", "第2名", "分", "完整顺序"])
    for item in samples["public_7_days"]:
        ranking = item["ranking"]
        public_cases.append((
            item["target_date"], item["pillars"]["year"], item["pillars"]["month"], item["pillars"]["day"], item["solar_term"],
            ranking[0]["color"], ranking[0]["rule_score"], ranking[1]["color"], ranking[1]["rule_score"],
            " > ".join(f'{row["color"]}({row["rule_score"]})' for row in ranking),
        ))
    _fit(public_cases, [16, 12, 12, 12, 14, 14, 10, 14, 10, 65])

    personal_cases = workbook.create_sheet("个人3日样本")
    _header(personal_cases, ["案例", "出生日期", "出生时间", "精度", "目标日期", "日主状态", "支持比", "五行分布", "完整顺序"])
    for item in samples["personal_3_days"]:
        personal_cases.append((
            item["case_code"], item["birth_date"], item["birth_time"] or "未知", item["precision_mode"], item["target_date"],
            item["strength"]["regime"], item["strength"]["support_ratio"],
            "、".join(f"{key}{value:.2f}%" for key, value in item["element_distribution"].items()),
            " > ".join(f'{row["color"]}({row["rule_score"]})' for row in item["ranking"]),
        ))
    _fit(personal_cases, [24, 16, 16, 18, 16, 16, 16, 60, 70])
    return workbook


def _set_doc_font(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_table(document: Document, headers: list[str], rows: list[tuple]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


def build_word(samples: dict) -> Document:
    """生成面向产品、开发和后续审核的中文说明书。"""
    document = Document()
    _set_doc_font(document)
    title = document.add_heading("今日五色资料综合算法与参数说明 V1.0", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("五色应时，知时而行｜北京时间口径｜资料综合试运行版")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("一、这次已经给出的版本", level=1)
    document.add_paragraph(
        "本版本不再等待老师填写空表，已经提供可运行的公共五色与个人五色配置、"
        "评分引擎、逐项审计明细、未来7天公共样本和两组未来3天个人样本。"
        "规则状态是 source_reviewed：表示公开资料已经核对、工程参数已经固定；"
        "它不是对外可宣传的专家背书，也暂未标记为 active。"
    )
    document.add_paragraph(
        "最重要的边界：古籍没有“每日穿衣颜色百分制公式”。五行五色、生克、月令、"
        "旺衰、扶抑属于资料依据；权重、分值、倍率、阈值、藏干比例和70/30融合属于"
        "本项目的现代可解释模型。"
    )

    document.add_heading("二、资料依据与采用边界", level=1)
    _add_table(document, ["资料", "采用原则", "本版处理"], [
        ("《黄帝内经·素问·阴阳应象大论》", "木对应青/苍、火赤、土黄、金白、水黑", "转为绿金、红金、黄金、白金、黑金；现代产品命名不冒充原词"),
        ("《五行大义》卷二", "木火土金水相生；干支五行；四时旺相休囚死", "用来确定关系和季节状态；不从原文虚构百分数"),
        ("《子平真诠评注》", "专求月令；得时为旺、失时为衰，同时结合年日时根气", "个人月支为最高单项权重，仍计算其他柱"),
        ("《滴天髓》通行本", "扶抑得宜、损益取中", "只用于偏弱生扶、偏强泄耗克、平衡补缺的方向约束"),
        ("《协纪辨方书》", "历法、月建和择日义例", "本版不把书中宜忌直接换算成个人颜色分，避免混用体系"),
    ])
    document.add_paragraph("公开核对路径：")
    for ref in PUBLIC_RESEARCH_CONFIG.professional_references:
        document.add_paragraph(ref, style="List Bullet")
    document.add_paragraph(
        "《滴天髓》《协纪辨方书》版本较多、公开站点访问也可能受限；本版没有让二者"
        "承担任何独有数值，只采用通行原则或明确排除不相干的择日宜忌，因此不影响"
        "算法的可复现性。正式出版或宣传引用前仍应按确定底本复核页码。"
    )

    document.add_heading("三、公共五色算法", level=1)
    document.add_paragraph("公共总分 = Σ（候选颜色与参考五行的关系分 × 因子权重 ÷ 100）。")
    _add_table(document, ["层级", "因子", "权重"], [
        ("年层", "年干 / 年支", "5% / 5%"),
        ("月令季节层", "月干 / 月支 / 节气", "15% / 25% / 5%"),
        ("日层", "日干 / 日支", "25% / 20%"),
    ])
    document.add_paragraph("设计意图：日层45%保证每天确实变化；月令季节层45%落实重月令；年层10%只作长期背景。")
    _add_table(document, ["候选色相对环境", "分值", "内部解释"], [
        ("环境生候选", "+30", "支持最高"),
        ("同五行", "+20", "同类协调"),
        ("候选克环境", "+5", "可用但不列最高"),
        ("候选生环境", "-10", "输出较多"),
        ("环境克候选", "-25", "谨慎使用"),
    ])
    document.add_paragraph(
        "节气按月建归类：立春至春分以木为主，清明谷雨为辰土；立夏至夏至为火，"
        "小暑大暑为未土；立秋至秋分为金，寒露霜降为戌土；立冬至冬至为水，"
        "小寒大寒为丑土。这样保留四个季末土月，不机械四等分。"
    )

    document.add_heading("四、个人五色算法", level=1)
    document.add_paragraph("个人计算共五步：")
    for text in [
        "以出生日干为日主；未知时辰只用年、月、日三柱。",
        "天干按固定五行计，地支按版本化藏干比例拆分；月支权重28%，为最高单项。",
        "按出生月支确定主气，再用旺1.40、相1.20、休1.00、囚0.80、死0.60修正。",
        "生我＋同我的结构占比小于42%为偏弱，大于58%为偏强，中间为相对平衡。",
        "出生结构需求分占70%，当天公共环境分占30%，得到最终个人五色顺序。",
    ]:
        document.add_paragraph(text, style="List Number")
    _add_table(document, ["柱位置", "年干", "年支", "月干", "月支", "日干", "日支", "时干", "时支"], [
        ("权重", 8, 12, 12, 28, 15, 15, 5, 5),
    ])
    _add_table(document, ["状态", "生我", "同我", "我生", "我克", "克我"], [
        ("偏弱", 35, 25, -8, -15, -25),
        ("相对平衡", 2, 0, 4, 2, -2),
        ("偏强", -25, -20, 30, 20, 12),
    ])
    document.add_paragraph(
        "补缺修正 = (20% - 该五行实际占比) × 0.8，并限制在±15内。"
        "个人最终分 = (角色基础分 + 补缺修正) × 70% + 公共环境分 × 30%。"
    )
    document.add_paragraph(
        "未知时辰处理：完全删除时干5%和时支5%的贡献，再把其余六项经过季节修正后的"
        "总量重新归一化；输出precision_mode=three_pillars，并加入PERSON_MISSING_TIME依据。"
    )

    document.add_heading("五、固定样本摘要", level=1)
    public_rows = []
    for item in samples["public_7_days"]:
        public_rows.append((
            item["target_date"], item["pillars"]["day"], item["solar_term"],
            " > ".join(f'{row["color"]}({row["rule_score"]})' for row in item["ranking"]),
        ))
    _add_table(document, ["日期", "日柱", "节气", "公共五色顺序"], public_rows)
    document.add_paragraph(
        "个人样本详见同目录 Excel 的“个人3日样本”和 JSON 文件。样本固定从2026-08-19"
        "开始，不会因运行脚本当天日期变化而变化，可直接作为回归测试基线。"
    )

    document.add_heading("六、产品使用和合规措辞", level=1)
    for text in [
        "面向用户表述为传统文化与生活方式建议，不使用保证招财、改命、消灾等承诺。",
        "排名描述用更适合、可搭配、谨慎选用，不输出疾病、灾祸、收益等确定性判断。",
        "未知时辰必须告知精度较低，不能补造时柱或输出精确时段断语。",
        "模型只负责把结构化结果写成自然语言，不能修改日期、干支、分数和排名。",
        "任何参数修改都升级规则版本并重新生成7日/3日样本，旧缓存按版本失效。",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("七、下一阶段接口方向", level=1)
    document.add_paragraph(
        "下一阶段应把本版规则接入每日内容生成任务：公共内容每天提前缓存未来7天；"
        "用户保存或修改档案后，按profile_version提前缓存未来3天。缓存键至少包含日期、"
        "规则版本、配置指纹、用户档案版本和精度模式。只有规则结果成功后，才调用大模型"
        "生成穿搭、事务和香品文案；模型失败时保留结构化排名，不写入半成品缓存。"
    )
    return document


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    samples = build_samples()
    JSON_PATH.write_text(json.dumps(samples, ensure_ascii=False, indent=2), encoding="utf-8")
    build_excel(samples).save(XLSX_PATH)
    build_word(samples).save(DOCX_PATH)
    print(f"generated: {XLSX_PATH}")
    print(f"generated: {DOCX_PATH}")
    print(f"generated: {JSON_PATH}")


if __name__ == "__main__":
    main()
