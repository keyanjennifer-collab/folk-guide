"""生成正式 AI 国学问答框架的实施与学习报告。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from generate_code_guide_doc import ROOT, set_styles


OUTPUT = ROOT / "docs" / "AI国学正式问答框架_实施与代码学习报告_V1.0.docx"


def bullets(doc, values):
    """向文档添加统一样式的项目符号。"""
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


doc = Document()
set_styles(doc)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("AI国学正式问答框架\n实施与代码学习报告")
subtitle = doc.add_paragraph("五行五色财库香小程序｜V1.0｜2026年8月12日")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("一、本阶段交付", level=1)
bullets(doc, [
    "正式 POST /api/ai/chat 接口。",
    "新客账号创建后赠送3天：普通问答每日20次，七日比较每日2次。",
    "预留正式30天权益：普通问答每日50次，七日比较每日5次。",
    "安全分类在检索和模型调用之前执行，高风险问题不消耗次数。",
    "混合检索审核知识片段，并返回文档、章节和片段引用。",
    "本人历史记录、点赞/点踩反馈和跨用户数据隔离。",
    "全项目自动测试：21 passed。",
])

doc.add_heading("二、关键代码", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "文件"
table.rows[0].cells[1].text = "职责"
for path, duty in [
    ("backend/app/ai_routes.py", "问答、次数、历史和反馈HTTP接口"),
    ("backend/app/ai_service.py", "安全分类、权益计算、检索及消息保存"),
    ("backend/app/ai_schemas.py", "输入输出校验模型"),
    ("backend/app/llm_provider.py", "可替换大模型接口与本地开发实现"),
    ("backend/app/models.py", "AI权益和正式问答记录数据表"),
    ("backend/tests/test_ai_chat.py", "权益、安全和权限隔离测试"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = path, duty

doc.add_heading("三、接口说明", level=1)
bullets(doc, [
    "GET /api/ai/quota：查询权益、到期时间及今日剩余次数。",
    "POST /api/ai/chat：正式AI国学问答。",
    "GET /api/ai/history：读取本人问答历史。",
    "PUT /api/ai/messages/{id}/feedback：对本人回答点赞或点踩。",
])

doc.add_heading("四、问答数据流", level=1)
doc.add_paragraph("JWT登录 → 问题安全分类 → 权益和次数校验 → 混合检索已审核资料 → 模型适配器生成回答 → 保存答案与引用 → 返回剩余次数。")
doc.add_paragraph("用户聊天不会进入公共知识库；反馈接口必须同时匹配 message_id 和当前 user_id，防止跨用户访问。")

doc.add_heading("五、当前限制和下一步", level=1)
doc.add_paragraph("尚未选择正式大模型、Embedding和外部向量库。因此当前 local-context-preview-v1 只整理检索片段并明确标记为开发预览，不冒充真正AI回答。下一步需要选择服务商、配置密钥、实现供应商适配器，并把小程序AI国学页面连接到这些正式接口。")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
