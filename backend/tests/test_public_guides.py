import io
from datetime import date, datetime, timedelta, timezone

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.database import SessionLocal
from app.models import PublicGuide, PublicGuideAudit
from sqlalchemy import select


ADMIN_HEADERS = {"X-Admin-Key": "dev-admin-key", "X-Admin-Name": "pytest-admin"}


def cleanup_guide(day: date) -> None:
    with SessionLocal() as db:
        guide = db.scalar(select(PublicGuide).where(PublicGuide.guide_date == day))
        if guide:
            db.query(PublicGuideAudit).filter(PublicGuideAudit.guide_id == guide.id).delete()
            db.delete(guide)
            db.commit()


def guide_payload(day: date) -> dict:
    fixed = [
        (1, "绿金", "木", "今天很顺", "GREEN"),
        (2, "黑金", "水", "比较合适", "BLACK"),
        (3, "黄金", "土", "平稳一般", "GOLD"),
        (4, "白金", "金", "会比较累", "WHITE"),
        (5, "红金", "火", "成效偏弱", "RED"),
    ]
    return {
        "guide_date": day.isoformat(),
        "weekday": "星期一",
        "lunar_date": "农历测试日期",
        "solar_term": "节气测试",
        "day_ganzhi": "甲子",
        "items": [
            {
                "rank": rank,
                "color": color,
                "element": element,
                "smoothness": smoothness,
                "suitable": ["合作", "沟通"],
                "resistance": "可能需要更多耐心",
                "advice": "先确认重点再行动",
                "product_code": product,
                "incense_name": f"{color}财库香",
                "scent": "香气描述",
            }
            for rank, color, element, smoothness, product in fixed
        ],
        "share_title": "今日五色排名",
        "share_summary": "今日完整建议已更新",
        "push_summary": "今日五色已更新",
        "rule_version": "manual-test-v1",
    }


def test_public_guide_workflow_and_audit():
    day = date.today() + timedelta(days=20)
    cleanup_guide(day)
    with TestClient(app) as client:
        unauthorized = client.post("/api/admin/public-guides", json=guide_payload(day))
        assert unauthorized.status_code == 403

        created = client.post("/api/admin/public-guides", headers=ADMIN_HEADERS, json=guide_payload(day))
        assert created.status_code == 200
        assert created.json()["status"] == "draft"

        reviewed = client.post(f"/api/admin/public-guides/{day}/review", headers=ADMIN_HEADERS)
        assert reviewed.status_code == 200
        assert reviewed.json()["status"] == "reviewing"

        published = client.post(f"/api/admin/public-guides/{day}/publish", headers=ADMIN_HEADERS)
        assert published.status_code == 200
        assert published.json()["status"] == "published"

        public = client.get(f"/api/public-guides/{day}")
        assert public.status_code == 200
        assert [item["rank"] for item in public.json()["items"]] == [1, 2, 3, 4, 5]

        audits = client.get(f"/api/admin/public-guides/{day}/audits", headers=ADMIN_HEADERS)
        assert audits.status_code == 200
        assert [item["action"] for item in audits.json()] == ["create", "submit_review", "publish"]

        withdrawn = client.post(f"/api/admin/public-guides/{day}/withdraw", headers=ADMIN_HEADERS)
        assert withdrawn.status_code == 200
        assert client.get(f"/api/public-guides/{day}").status_code == 404
    cleanup_guide(day)


def test_excel_template_and_import():
    day = date.today() + timedelta(days=21)
    cleanup_guide(day)
    with TestClient(app) as client:
        template = client.get("/api/admin/public-guides/template.xlsx", headers=ADMIN_HEADERS)
        assert template.status_code == 200
        workbook = load_workbook(io.BytesIO(template.content))
        sheet = workbook.active
        for row in range(2, 7):
            sheet.cell(row=row, column=1, value=day.isoformat())
        output = io.BytesIO()
        workbook.save(output)

        imported = client.post(
            "/api/admin/public-guides/import",
            headers=ADMIN_HEADERS,
            files={"file": ("guides.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 200
        assert imported.json()["created"] == 1
        detail = client.get(f"/api/admin/public-guides/{day}", headers=ADMIN_HEADERS)
        assert detail.status_code == 200
        assert detail.json()["status"] == "draft"
    cleanup_guide(day)


def test_scheduled_content_auto_publishes():
    day = date.today() + timedelta(days=22)
    cleanup_guide(day)
    with TestClient(app) as client:
        client.post("/api/admin/public-guides", headers=ADMIN_HEADERS, json=guide_payload(day))
        client.post(f"/api/admin/public-guides/{day}/review", headers=ADMIN_HEADERS)
        schedule = client.post(
            f"/api/admin/public-guides/{day}/schedule",
            headers=ADMIN_HEADERS,
            json={"scheduled_at": (datetime.now(timezone.utc) - timedelta(seconds=1)).isoformat()},
        )
        assert schedule.status_code == 200
        listed = client.get("/api/admin/public-guides", headers=ADMIN_HEADERS)
        target = next(item for item in listed.json() if item["guide_date"] == day.isoformat())
        assert target["status"] == "published"
    cleanup_guide(day)


def test_rejects_incomplete_or_wrong_color_mapping():
    day = date.today() + timedelta(days=23)
    cleanup_guide(day)
    payload = guide_payload(day)
    payload["items"][0]["element"] = "火"
    with TestClient(app) as client:
        response = client.post("/api/admin/public-guides", headers=ADMIN_HEADERS, json=payload)
        assert response.status_code == 422
    cleanup_guide(day)


def test_admin_page_is_available():
    with TestClient(app) as client:
        response = client.get("/admin/public-guides")
        assert response.status_code == 200
        assert "今日五色运营后台" in response.text


def test_today_returns_rule_cache_when_no_manual_content(monkeypatch):
    """没有人工发布内容时，首页应返回资料综合规则缓存而不是空白或演示数据。"""
    fixed_day = date(2036, 8, 19)
    cleanup_guide(fixed_day)
    monkeypatch.setattr("app.public_guide_routes.beijing_today", lambda: fixed_day)

    with TestClient(app) as client:
        response = client.get("/api/public-guides/today")

    assert response.status_code == 200
    payload = response.json()
    assert payload["guide_date"] == fixed_day.isoformat()
    assert payload["rule_version"] == "wuse-public-research-v1.0"
    assert len(payload["items"]) == 5
    assert {item["color"] for item in payload["items"]} == {"白金", "绿金", "黑金", "红金", "黄金"}


def test_today_returns_published_content_for_beijing_date(monkeypatch):
    """今日接口应按统一的北京时间读取当天已经发布的记录。"""
    fixed_day = date(2036, 8, 20)
    cleanup_guide(fixed_day)
    monkeypatch.setattr("app.public_guide_routes.beijing_today", lambda: fixed_day)

    with TestClient(app) as client:
        client.post("/api/admin/public-guides", headers=ADMIN_HEADERS, json=guide_payload(fixed_day))
        client.post(f"/api/admin/public-guides/{fixed_day}/review", headers=ADMIN_HEADERS)
        client.post(f"/api/admin/public-guides/{fixed_day}/publish", headers=ADMIN_HEADERS)
        response = client.get("/api/public-guides/today")

    assert response.status_code == 200
    assert response.json()["guide_date"] == fixed_day.isoformat()
    cleanup_guide(fixed_day)
