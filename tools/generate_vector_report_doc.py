from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from generate_code_guide_doc import ROOT, set_styles


OUTPUT = ROOT / "docs" / "AI国学向量检索基础层_实施与代码学习报告_V1.0.docx"


def add_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


doc = Document()
set_styles(doc)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("AI国学向量检索基础层\n实施与代码学习报告")
p = doc.add_paragraph("五行五色财库香小程序｜V1.0｜2026年8月12日")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(11)

doc.add_heading("一、本阶段完成内容", level=1)
add_list(doc, [
    "为知识库解析、切片、审核、文件安全和数据库模型补充了中文注释。",
    "建立 VectorStore 统一协议：upsert、query、delete、health_check。",
    "加入 MemoryVectorStore 本地实现，不产生模型费用，便于先开发和测试。",
    "新增审核片段向量同步，以及关键词与向量结果的混合排序。",
    "文档停用时删除向量映射；未审核文档禁止同步。",
    "全项目自动测试结果：19 passed。",
])

doc.add_heading("二、主要代码文件", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "文件"
table.rows[0].cells[1].text = "作用"
for path, purpose in [
    ("backend/app/vector_store.py", "向量库统一接口和本地开发实现"),
    ("backend/app/retrieval_service.py", "同步、删除和混合检索业务逻辑"),
    ("backend/app/knowledge_routes.py", "管理端同步与检索 HTTP 接口"),
    ("backend/app/knowledge_service.py", "文档解析、切片和关键词检索"),
    ("backend/tests/test_vector_retrieval.py", "审核边界、同步和停用测试"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = path, purpose

doc.add_heading("三、数据流程", level=1)
doc.add_paragraph("上传文档 → 解析切片 → 人工审核 → 同步 approved 片段 → 混合检索 → 返回正文和来源。")
doc.add_paragraph("重要边界：上传或待审核内容不能进入向量库；文档停用后不能继续被问答召回。")

doc.add_heading("四、新增接口", level=1)
add_list(doc, [
    "POST /api/admin/knowledge/documents/{id}/sync-vectors：同步审核通过的片段。",
    "GET /api/admin/knowledge/hybrid-search?q=问题：预览混合检索结果。",
])

doc.add_heading("五、为什么暂时使用本地实现", level=1)
doc.add_paragraph("外部向量库和 Embedding 服务商尚未确定。业务层先依赖统一协议，可以完整验证审核和检索流程；以后接入云服务只需增加适配器，不必重写文档处理与问答业务。local-bigram-v1 仅用于开发测试，不属于真正的语义 Embedding，也不会作为正式生产效果。")

doc.add_heading("六、下一阶段", level=1)
add_list(doc, [
    "确定向量库与 Embedding 服务商，并实现持久化适配器。",
    "新增正式 /api/ai/chat：问题分类、检索、提示词、模型回答和引用。",
    "加入新客3天与正式服务的问答次数限制。",
    "把小程序 AI 国学页面的模拟答案替换为真实接口。",
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
