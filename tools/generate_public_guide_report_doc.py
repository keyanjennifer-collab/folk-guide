from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from generate_code_guide_doc import ROOT, bullets, code, page_break, set_styles, shade, steps, table


OUTPUT = ROOT / "docs" / "今日五色数据库后台与Excel导入_实施报告_V1.0.docx"


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    set_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("今日五色数据库、运营后台\n与Excel批量导入实施报告")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("五行五色财库香项目｜开发交付 V1.0")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(113, 128, 120)
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1)
    shade(box.cell(0, 0), "E9F0EB")
    box.cell(0, 0).text = "本阶段已经完成“人工录入每日五色→保存草稿→提交审核→定时或立即发布→小程序公开读取→撤回与审计”的完整后端闭环，并提供Excel模板和批量导入能力。AI知识库Word解析与向量化属于下一阶段，尚未在本次混入。"
    doc.add_paragraph()

    doc.add_heading("01｜本次完成内容", level=1)
    bullets(doc, [
        "新增公共每日五色数据库表。",
        "新增后台操作审计表。",
        "新增每天正好五种颜色、排名1至5的结构化校验。",
        "固定白金=金、绿金=木、黑金=水、红金=火、黄金=土。",
        "限制顺畅度为：今天很顺、比较合适、平稳一般、会比较累、成效偏弱。",
        "实现草稿、待审核、待发布、已发布和已撤回状态。",
        "实现立即发布与到期自动发布。",
        "实现公开今日接口和按日期接口，不要求用户登录。",
        "实现管理员Key保护、操作人记录和状态审计。",
        "实现Excel模板下载、文件大小/格式校验和批量导入草稿。",
        "实现浏览器运营管理页面。",
        "小程序首页已连接公开今日接口，未发布时显示标记为演示的兜底内容。",
    ])

    doc.add_heading("02｜新增文件", level=1)
    table(doc, ["文件", "作用"], [
        ["backend/app/public_guide_schemas.py", "每日五色数据结构和业务校验"],
        ["backend/app/public_guide_service.py", "保存、审计、发布、Excel生成与解析"],
        ["backend/app/public_guide_routes.py", "公开接口、管理员接口和运营页面路由"],
        ["backend/app/admin_auth.py", "X-Admin-Key管理员凭证检查"],
        ["backend/app/static/admin_public_guides.html", "本地运营录入与发布页面"],
        ["backend/tests/test_public_guides.py", "工作流、Excel、排期、校验和页面测试"],
        ["docs/今日五色批量导入模板.xlsx", "运营人员批量填写模板"],
    ], [7.0, 8.7])
    doc.add_paragraph("同时修改models.py、main.py、config.py、requirements.txt、.env.example以及小程序home/index.ts与index.wxml。")

    doc.add_heading("03｜数据库", level=1)
    table(doc, ["表", "关键字段", "用途"], [
        ["public_guides", "guide_date、status、version、payload_json、scheduled_at、published_at", "每天一条公共内容及版本状态"],
        ["public_guide_audits", "guide_id、action、operator、before_status、after_status、detail_json", "记录创建、修改、审核、发布和撤回"],
    ], [3.7, 7.0, 5.0])
    doc.add_paragraph("当前继续使用SQLite原型数据库；新表由SQLAlchemy启动时创建。进入多人协作和生产阶段前，应切换PostgreSQL并使用Alembic迁移。")

    doc.add_heading("04｜内容字段", level=1)
    table(doc, ["层级", "字段"], [
        ["每日公共信息", "日期、星期、农历、节气、当日干支、规则版本"],
        ["每色排名", "排名、颜色、五行、顺畅度"],
        ["行动建议", "适合事项、可能阻力、行动建议"],
        ["商品连接", "商品编码、香品名称、香气描述"],
        ["分发", "分享标题、分享摘要、推送摘要"],
        ["系统", "状态、版本、排期时间、发布时间、创建和更新时间"],
    ], [4.0, 11.7])

    doc.add_heading("05｜状态流程", level=1)
    code(doc, "draft 草稿\n  ↓ 提交审核\nreviewing 待审核\n  ↓ 设置排期\nscheduled 待发布\n  ↓ 到期自动发布或人工立即发布\npublished 已发布\n  ↓ 撤回\nwithdrawn 已撤回\n  ↓ 修改后重新成为草稿")
    bullets(doc, [
        "已发布内容不能直接编辑，必须先撤回。",
        "发布后的内容才允许公共接口读取。",
        "每次创建、修改、审核、发布和撤回都写入审计。",
        "定时任务当前由访问接口时检查到期内容；生产阶段应增加独立调度器。",
    ])

    doc.add_heading("06｜公开接口", level=1)
    table(doc, ["接口", "认证", "作用"], [
        ["GET /api/public-guides/today", "不需要", "按北京时间读取今日已发布内容"],
        ["GET /api/public-guides/{date}", "不需要", "读取指定日期已发布内容"],
    ], [7.0, 2.5, 6.2])
    doc.add_paragraph("今日没有已发布内容时返回404及“今日详细指南正在更新”提示。小程序当前捕获错误并显示演示内容，同时标记“演示内容”。")

    doc.add_heading("07｜管理员接口", level=1)
    table(doc, ["方法与路径", "作用"], [
        ["GET /api/admin/public-guides", "列表，可按状态筛选"],
        ["POST /api/admin/public-guides", "创建或更新某日草稿"],
        ["GET/PUT /api/admin/public-guides/{date}", "读取或编辑某日内容"],
        ["POST /{date}/review", "提交审核"],
        ["POST /{date}/schedule", "设置UTC排期时间"],
        ["POST /{date}/publish", "立即发布"],
        ["POST /{date}/withdraw", "撤回已发布内容"],
        ["GET /{date}/audits", "查看操作审计"],
        ["GET /template.xlsx", "下载Excel模板"],
        ["POST /import", "上传.xlsx并导入草稿"],
    ], [8.2, 7.5])
    doc.add_paragraph("管理员请求必须携带X-Admin-Key和X-Admin-Name。开发默认Key为dev-admin-key，正式环境必须通过ADMIN_API_KEY替换成随机强密钥。")

    doc.add_heading("08｜运营页面使用方法", level=1)
    code(doc, "启动后端后访问：\nhttp://127.0.0.1:8000/admin/public-guides")
    steps(doc, [
        "填写管理员Key和操作人。",
        "选择目标日期，填写星期、农历、节气和干支。",
        "确认五种颜色的排名、顺畅度、事项、阻力、建议和香品。",
        "填写分享标题、摘要、推送摘要和规则版本。",
        "保存草稿。",
        "提交审核。",
        "选择立即发布，或填写时间后设置定时发布。",
        "需要修改已发布内容时，先撤回再保存草稿。",
    ])

    doc.add_heading("09｜Excel导入", level=1)
    doc.add_paragraph("模板文件：docs/今日五色批量导入模板.xlsx。一天必须填写五行，每行对应一个颜色。适合事项可用顿号、逗号或分号分隔。")
    table(doc, ["校验", "结果"], [
        ["文件不是.xlsx", "拒绝"],
        ["文件超过5MB", "拒绝"],
        ["缺少模板列", "提示具体缺少列"],
        ["日期格式错误", "提示错误行号"],
        ["排名缺失或重复", "拒绝整日内容"],
        ["缺少任一固定颜色", "拒绝整日内容"],
        ["颜色与五行错误", "例如绿金填成火时拒绝"],
        ["顺畅度不在允许值中", "拒绝"],
        ["已发布日期", "要求先撤回，避免覆盖线上内容"],
        ["全部通过", "保存为草稿，不直接发布"],
    ], [6.0, 9.7])

    doc.add_heading("10｜小程序首页接入", level=1)
    code(doc, "页面onShow\n→ GET /api/public-guides/today\n→ 成功：转换成前端五色卡片并显示“今日已发布”\n→ 失败：保留演示数据并显示“演示内容”")
    doc.add_paragraph("分享标题使用后台返回的share_title；后续还需要接入真实分享图片和来源统计。")

    doc.add_heading("11｜测试结果", level=1)
    doc.add_paragraph("本阶段完成后，全项目测试结果为：14 passed。")
    table(doc, ["测试", "覆盖内容"], [
        ["管理员权限", "缺少或错误Key返回403"],
        ["工作流", "草稿、审核、发布、公开读取、撤回"],
        ["审计", "创建、审核和发布动作顺序"],
        ["Excel", "模板下载、修改日期、上传和保存草稿"],
        ["定时发布", "到期内容自动变为published"],
        ["结构校验", "错误颜色五行关系返回422"],
        ["运营页面", "HTML页面可访问"],
        ["回归", "档案、历法、登录占位和问答接口保持通过"],
    ], [4.5, 11.2])

    doc.add_heading("12｜仍需后续完成", level=1)
    bullets(doc, [
        "把开发管理员Key升级成真正的后台账号、角色和登录。",
        "将SQLite切换PostgreSQL并加入Alembic。",
        "生产环境使用独立调度器执行定时发布，不依赖接口访问触发。",
        "增加后台预览、内容复制、分页、筛选和批量审核。",
        "增加访问、分享、登录和商品点击统计。",
        "由国学老师确认真正的每日排名和规则版本。",
        "下一阶段实现AI知识文档上传、解析、切片、审核和外部向量库。",
    ])

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "今日五色｜数据库后台与Excel导入实施报告"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.text = "五行五色财库香项目 V1.0"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build()
    document.save(OUTPUT)
    print(OUTPUT)
