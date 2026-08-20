from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from generate_code_guide_doc import ROOT, bullets, code, page_break, set_styles, steps, table


OUTPUT = ROOT / "docs" / "五行五色小程序前端代码学习手册_V1.0.docx"


def source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def heading_note(doc, title, text):
    doc.add_heading(title, level=2)
    doc.add_paragraph(text)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "五行五色小程序｜前端代码学习手册"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(145, 155, 149)
        footer = section.footer.paragraphs[0]
        footer.text = "初学者学习版 V1.0｜配合 D:\\soft\\folk-guide 使用"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(145, 155, 149)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    set_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("五行五色小程序\n前端代码学习手册")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("从零理解 TypeScript、WXML、WXSS、JSON 和页面交互")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(113, 128, 120)
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1)
    cell = box.cell(0, 0)
    cell.text = "这不是后端API文档。本手册只讲你目前不熟悉的前端部分，并直接使用当前项目代码作为例子。建议一边打开 Word，一边在微信开发者工具或 VS Code 中查看 D:\\soft\\folk-guide\\miniprogram。"
    from generate_code_guide_doc import shade
    shade(cell, "E9F0EB")
    doc.add_paragraph()
    table(doc, ["适合读者", "学习目标"], [
        ["懂一点Python或后端，但看不懂前端", "能判断一个小程序页面由哪些文件组成"],
        ["没学过TypeScript", "能读懂变量、数组、对象、函数和事件"],
        ["没写过WXML/WXSS", "能修改文字、数据、按钮、布局和颜色"],
        ["准备继续参与本项目", "能区分演示数据、真实接口和业务状态"],
    ], [5.3, 10.4])

    doc.add_heading("建议学习顺序", level=1)
    steps(doc, [
        "先读第1至第5章，理解一个页面由四种文件组成。",
        "打开今日五色页面，对照第6章修改一处文字并刷新预览。",
        "学习第7章的数据绑定和事件，再看AI国学的固定问题。",
        "学习第8章的列表、条件和购物车，再看财库香。",
        "学习第9章的页面状态，再看“我的”。",
        "最后做第15章练习，不要一开始背全部语法。",
    ])

    doc.add_heading("目录", level=1)
    for item in [
        "01 前端到底是什么", "02 微信小程序四种文件", "03 TypeScript最小知识",
        "04 WXML页面结构", "05 WXSS页面样式", "06 JSON与全局导航",
        "07 数据绑定与用户点击", "08 循环、条件和弹层", "09 页面跳转与生命周期",
        "10 今日五色逐步阅读", "11 AI国学逐步阅读", "12 财库香逐步阅读",
        "13 我的逐步阅读", "14 浏览器预览与小程序的区别", "15 动手练习",
        "16 常见错误和调试", "17 当前项目前端代码地图", "18 下一阶段学习路线",
        "附录A 常用语法速查", "附录B 当前核心前端源码",
    ]:
        doc.add_paragraph(item)

    page_break(doc)
    doc.add_heading("01｜前端到底是什么", level=1)
    doc.add_paragraph("后端API负责数据、规则和安全；前端负责用户在手机上看到什么、点击后发生什么，以及如何把后端数据展示出来。")
    code(doc, "用户看到页面\n   ↓ 点击按钮、填写内容\n微信小程序前端\n   ↓ 发送 HTTP 请求\nPython 后端 API\n   ↓ 查询数据库、计算规则\n返回 JSON 数据\n   ↓\n前端更新页面")
    table(doc, ["事情", "前端负责", "后端负责"], [
        ["显示今日五色", "颜色、卡片、展开动画、文字排版", "返回当天真实排名和建议"],
        ["AI问答", "输入框、消息气泡、加载状态、引用展示", "检索向量库、调用模型、次数扣减"],
        ["加入购物车", "按钮、商品数量、临时显示", "校验价格、库存、优惠和登录用户"],
        ["支付", "发起支付并显示结果", "创建订单、签名、验支付回调、发货"],
        ["服务卡", "展示剩余天数和操作按钮", "发卡、开启、到期、赠礼和冻结"],
    ], [3.0, 6.1, 6.6])
    bullets(doc, [
        "前端传来的价格不能作为最终价格，因为用户可能修改前端数据。",
        "前端隐藏按钮不等于真正限制权限，后端仍必须检查。",
        "当前项目很多前端数据是模拟数据，刷新页面会恢复。",
    ])

    doc.add_heading("02｜微信小程序四种文件", level=1)
    doc.add_paragraph("一个页面通常由四个同名文件组成。以今日五色为例：")
    code(doc, "pages/home/\n├─ index.ts      页面逻辑和数据\n├─ index.wxml    页面结构\n├─ index.wxss    页面样式\n└─ index.json    页面配置")
    table(doc, ["文件", "可以把它理解成", "主要写什么"], [
        ["index.ts", "大脑", "数据、函数、点击事件、调用API"],
        ["index.wxml", "骨架", "标题、卡片、按钮、列表、输入框"],
        ["index.wxss", "外观", "颜色、大小、间距、圆角、布局"],
        ["index.json", "页面设置", "导航栏标题、分享、下拉刷新"],
    ], [3.0, 4.2, 8.5])
    doc.add_paragraph("学习时不要一次看四个文件。推荐先看WXML知道页面有哪些区域，再看TS知道数据从哪里来，最后看WXSS理解外观。")

    doc.add_heading("03｜TypeScript最小知识", level=1)
    doc.add_paragraph("TypeScript是带类型提示的JavaScript。你不需要先学完整语言，只要掌握本项目用到的几个结构。")
    heading_note(doc, "3.1 变量", "const表示这个变量名称不会重新指向其他值；let表示它可以改变。")
    code(doc, 'const name = "绿金";\nlet remainingQuestions = 20;\nremainingQuestions = 19;')
    heading_note(doc, "3.2 基本类型", "string是文字，number是数字，boolean是真/假，null表示当前没有值。")
    code(doc, 'const title: string = "今日五色";\nconst price: number = 199;\nconst loggedIn: boolean = true;\nconst selectedProduct = null;')
    heading_note(doc, "3.3 数组", "数组是按顺序排列的一组值，使用方括号。")
    code(doc, 'const suitable: string[] = ["合作沟通", "项目启动", "计划推进"];\n\n// 第一个值，下标从0开始\nconsole.log(suitable[0]);')
    heading_note(doc, "3.4 对象", "对象把一件事的多个字段放在一起。")
    code(doc, 'const product = {\n  id: "green",\n  name: "绿金",\n  price: 59,\n  spec: "30支／管"\n};\n\nconsole.log(product.name); // 绿金')
    heading_note(doc, "3.5 type类型定义", "type告诉编辑器一个对象应该有哪些字段。如果漏字段或类型错，TypeScript会提前报错。")
    code(doc, 'type Product = {\n  id: string;\n  name: string;\n  price: number;\n};')
    heading_note(doc, "3.6 函数", "函数是一段可以反复执行的操作。括号中是输入，花括号中是具体动作。")
    code(doc, 'function add(a: number, b: number): number {\n  return a + b;\n}\n\nconst total = add(199, 59);')
    heading_note(doc, "3.7 箭头函数", "find、map、filter、reduce中经常出现箭头函数。它是函数的简写。")
    code(doc, 'const green = products.find((item) => item.id === "green");\n\nconst active = items.filter((item) => item.quantity > 0);\n\nconst total = items.reduce(\n  (sum, item) => sum + item.price * item.quantity,\n  0\n);')
    table(doc, ["方法", "作用", "项目例子"], [
        ["find", "找到第一个符合条件的值", "根据商品id找商品"],
        ["map", "把每个值转换后形成新数组", "修改购物车某个商品数量"],
        ["filter", "只保留符合条件的值", "删除数量为0的商品"],
        ["reduce", "把数组汇总成一个值", "计算购物车总价和总件数"],
    ], [3.0, 5.0, 7.7])

    doc.add_heading("04｜WXML页面结构", level=1)
    doc.add_paragraph("WXML类似HTML，但使用微信组件。最常用的是view、text、button、input、textarea、scroll-view和picker。")
    code(doc, '<view class="card">\n  <view class="title">今日五色</view>\n  <text class="muted">每日更新</text>\n  <button bindtap="showDetail">查看详情</button>\n</view>')
    table(doc, ["组件", "用途", "近似HTML"], [
        ["view", "容器、卡片、行、区域", "div"],
        ["text", "短文字，可放在一行中", "span"],
        ["button", "按钮", "button"],
        ["input", "单行输入", "input"],
        ["textarea", "多行输入", "textarea"],
        ["scroll-view", "可滚动区域", "带overflow的div"],
        ["picker", "日期、时间或选择器", "select/date input"],
        ["switch", "开关", "checkbox"],
    ], [3.2, 7.3, 5.2])
    doc.add_paragraph("class只是给样式起名字。WXML中的class=\"card\"会找到WXSS中的.card规则。")

    doc.add_heading("05｜WXSS页面样式", level=1)
    doc.add_paragraph("WXSS类似CSS。它决定页面怎么排列、颜色是什么、间距多大。微信小程序常用rpx，它会根据手机屏幕宽度自动缩放。")
    code(doc, '.card {\n  margin-bottom: 24rpx;\n  padding: 32rpx;\n  border-radius: 24rpx;\n  background: #ffffff;\n}\n\n.title {\n  color: #27322d;\n  font-size: 38rpx;\n  font-weight: 700;\n}')
    table(doc, ["属性", "含义", "例子"], [
        ["padding", "内容与边框之间的内边距", "padding: 24rpx"],
        ["margin", "当前元素与外部元素的距离", "margin-bottom: 20rpx"],
        ["background", "背景色或渐变", "background: #fff"],
        ["color", "文字颜色", "color: #385f4d"],
        ["font-size", "文字大小", "font-size: 28rpx"],
        ["border-radius", "圆角", "border-radius: 20rpx"],
        ["display:flex", "横向或纵向弹性布局", "配合justify-content"],
        ["display:grid", "网格布局", "商品两列排列"],
        ["position:fixed", "固定在屏幕某处", "购物车浮动按钮"],
    ], [3.5, 7.0, 5.2])
    heading_note(doc, "5.1 Flex布局", "本项目顶部左右排列、价格与按钮并排，大多使用Flex。")
    code(doc, '.price-line {\n  display: flex;\n  justify-content: space-between;\n  align-items: center;\n}')
    heading_note(doc, "5.2 Grid布局", "单色香两列商品卡和用户中心四个订单入口使用Grid。")
    code(doc, '.product-grid {\n  display: grid;\n  grid-template-columns: 1fr 1fr;\n  gap: 18rpx;\n}')

    doc.add_heading("06｜JSON与全局导航", level=1)
    doc.add_paragraph("app.json是整个小程序的总目录。pages决定有哪些页面，tabBar决定底部四个入口。")
    code(doc, source("miniprogram/app.json"), "当前 miniprogram/app.json")
    bullets(doc, [
        "pages数组第一项是启动时首先打开的页面。",
        "tabBar中的pagePath必须已经出现在pages中。",
        "底部Tab页面之间使用wx.switchTab。",
        "普通子页面使用wx.navigateTo。",
        "单页index.json只影响当前页面，例如enableShareAppMessage。",
    ])

    page_break(doc)
    doc.add_heading("07｜数据绑定与用户点击", level=1)
    heading_note(doc, "7.1 Page与data", "每个页面通过Page创建。data中的值可以显示在WXML里。")
    code(doc, 'Page({\n  data: {\n    title: "今日五色",\n    expandedRank: 1\n  }\n});')
    code(doc, '<view>{{title}}</view>\n<view>当前展开：{{expandedRank}}</view>')
    doc.add_paragraph("双花括号{{ }}表示从页面data中读取值，这叫数据绑定。")
    heading_note(doc, "7.2 setData", "不能只修改普通变量来更新页面。需要调用setData，微信才会重新渲染WXML。")
    code(doc, 'this.setData({ expandedRank: 2 });')
    heading_note(doc, "7.3 bindtap点击事件", "WXML用bindtap指定用户点击后调用哪个TS函数。")
    code(doc, '<button bindtap="showRules">服务说明</button>\n\n// index.ts\nshowRules() {\n  wx.showModal({ title: "服务说明", content: "……" });\n}')
    heading_note(doc, "7.4 dataset传参数", "同一个函数处理多个商品或颜色时，可以把id放到data-id。")
    code(doc, '<button data-id="{{item.id}}" bindtap="addProduct">加入</button>\n\naddProduct(event) {\n  const id = event.currentTarget.dataset.id;\n  // 根据id找到商品\n}')
    doc.add_paragraph("data-id在TS中对应dataset.id；data-question对应dataset.question。")
    heading_note(doc, "7.5 输入框", "输入时，event.detail.value就是用户输入。")
    code(doc, '<textarea value="{{question}}" bindinput="setQuestion" />\n\nsetQuestion(event) {\n  this.setData({ question: event.detail.value });\n}')

    doc.add_heading("08｜循环、条件和弹层", level=1)
    heading_note(doc, "8.1 wx:for循环", "有五种颜色或多条消息时，不需要复制五套WXML。使用wx:for遍历数组。")
    code(doc, '<view wx:for="{{guides}}" wx:key="rank">\n  <text>第{{item.rank}}名</text>\n  <text>{{item.name}}</text>\n</view>')
    bullets(doc, [
        "item代表当前这一个数组元素。",
        "index代表当前位置，从0开始。",
        "wx:key帮助微信判断哪一项发生了变化。",
        "嵌套循环可以用wx:for-item改名，避免两个item冲突。",
    ])
    heading_note(doc, "8.2 wx:if条件显示", "只有条件为真时才生成这个区域。")
    code(doc, '<view wx:if="{{cartCount}}">购物车有{{cartCount}}件</view>\n<view wx:else>购物车为空</view>')
    heading_note(doc, "8.3 模态弹层", "财库香详情和购物车使用固定遮罩＋底部面板。")
    code(doc, '<view wx:if="{{selectedProduct}}" class="overlay" catchtap="closeProduct">\n  <view class="product-sheet" catchtap="noop">\n    商品详情\n  </view>\n</view>')
    doc.add_paragraph("外层点击关闭，内层使用catchtap阻止点击继续传到外层，否则点击商品详情本身也会关闭。")
    heading_note(doc, "8.4 wx.showModal", "简单提示不需要自己画弹层，可以调用微信API。")
    code(doc, 'wx.showModal({\n  title: "登录后结算",\n  content: "提交订单时需要微信登录",\n  confirmText: "微信登录",\n  cancelText: "继续逛逛"\n});')

    doc.add_heading("09｜页面跳转与生命周期", level=1)
    table(doc, ["方法", "使用场景", "项目例子"], [
        ["wx.switchTab", "跳到四个底部Tab之一", "今日五色→财库香"],
        ["wx.navigateTo", "打开普通子页面并保留上一页", "我的→生辰档案"],
        ["wx.navigateBack", "返回上一页", "档案保存后返回"],
        ["wx.reLaunch", "关闭当前页面栈并重新打开", "注销后回首页"],
    ], [3.3, 6.2, 6.2])
    heading_note(doc, "9.1 onLoad", "页面第一次创建时执行，适合读取路径参数或初始数据。")
    heading_note(doc, "9.2 onShow", "页面每次显示都会执行，从其他页返回时也会执行，适合刷新数据。")
    heading_note(doc, "9.3 onShareAppMessage", "用户点击微信分享时返回标题和目标路径。")
    code(doc, 'onShareAppMessage() {\n  return {\n    title: "今日五色排名与生活建议",\n    path: "/pages/home/index?from=share"\n  };\n}')

    page_break(doc)
    doc.add_heading("10｜今日五色逐步阅读", level=1)
    doc.add_paragraph("阅读顺序：先打开index.wxml看页面区域，再打开index.ts找guides数组和事件，最后看index.wxss。")
    table(doc, ["页面区域", "WXML关键词", "TS数据/方法"], [
        ["日期和分享", "hero、open-type=share", "onShareAppMessage"],
        ["五色排名", "wx:for=guides", "guides"],
        ["展开详情", "wx:if=expandedRank", "toggleGuide"],
        ["对应香品", "了解香品按钮", "toIncense"],
        ["个人结果", "personal-card", "showPersonalLogin"],
        ["订阅提醒", "action-card", "showSubscribe"],
    ], [3.5, 5.8, 6.4])
    code(doc, 'toggleGuide(event) {\n  const rank = Number(event.currentTarget.dataset.rank);\n  this.setData({\n    expandedRank: this.data.expandedRank === rank ? 0 : rank\n  });\n}')
    doc.add_paragraph("三元表达式“条件 ? 值A : 值B”表示：如果当前已经展开这个排名，就改成0收起；否则展开用户点击的排名。")
    doc.add_paragraph("当前日期和guides写死在前端，仅用于视觉演示。接入后台后，应在onShow中请求当天接口，再this.setData({ guides: result.ranking })。")

    doc.add_heading("11｜AI国学逐步阅读", level=1)
    table(doc, ["区域", "数据", "作用"], [
        ["体验头部", "trialDaysLeft、dailyLimit", "显示3天体验和次数"],
        ["个人结果", "showFullResult", "控制完整结果展开"],
        ["固定问题", "questions", "四组问题横向滚动"],
        ["消息列表", "messages", "用户和AI消息气泡"],
        ["输入框", "question", "保存正在输入的问题"],
        ["反馈", "feedback", "有帮助或需要纠正"],
    ], [3.2, 5.0, 7.5])
    code(doc, 'type ChatItem = {\n  role: "user" | "assistant";\n  text: string;\n  references?: string[];\n  feedback?: "helpful" | "incorrect";\n};')
    doc.add_paragraph("竖线表示只能从几个值中选择；问号表示这个字段可以没有。例如用户消息可以没有references。")
    code(doc, 'this.setData({\n  messages: [\n    ...this.data.messages,\n    { role: "user", text: question }\n  ]\n});')
    doc.add_paragraph("...this.data.messages叫展开语法：先复制原有消息，再在末尾加入新消息。直接修改数组可能导致页面更新不可靠。")
    doc.add_paragraph("当前ANSWERS是演示答案表，正式接入向量库后应删除，sendQuestion改为调用自己的Python API。")

    doc.add_heading("12｜财库香逐步阅读", level=1)
    table(doc, ["功能", "主要方法", "关键理解"], [
        ["打开商品详情", "showProduct", "根据data-id从PRODUCTS查找商品"],
        ["加入套装", "addFlagship", "创建套装对象后复用addItem"],
        ["加入单品", "addProduct", "避免每种香写不同函数"],
        ["合并购物车", "addItem", "相同id增加quantity"],
        ["数量加减", "changeQuantity", "map修改、filter删除0件"],
        ["计算金额", "reduce", "价格×数量后汇总"],
        ["模拟结算", "checkout", "只显示登录说明，没有真实支付"],
    ], [3.3, 4.0, 8.4])
    code(doc, 'const items = this.data.cartItems.map((item) =>\n  item.id === id\n    ? { ...item, quantity: item.quantity + delta }\n    : item\n+).filter((item) => item.quantity > 0);')
    doc.add_paragraph("这段代码先逐个检查商品：如果id相同，就复制商品并修改数量；其他商品保持原样。最后只保留数量大于0的项目。")
    code(doc, 'const cartTotal = items.reduce(\n  (total, item) => total + item.price * item.quantity,\n  0\n);')
    doc.add_paragraph("0是初始总额。reduce遍历每个商品，把价格乘数量累加。真实结算必须由后端重新计算，不能相信这个前端总额。")

    doc.add_heading("13｜我的逐步阅读", level=1)
    doc.add_paragraph("“我的”页面的难点不是语法，而是状态多。当前data集中模拟了登录、档案、体验、问答次数、服务卡和订单数量。")
    code(doc, 'data: {\n  isLoggedIn: true,\n  profileCompleteness: 85,\n  trialDaysLeft: 3,\n  questionRemaining: 20,\n  pendingCards: 1,\n  orderCounts: {\n    unpaid: 0,\n    unshipped: 1,\n    shipped: 0,\n    afterSale: 0\n  }\n}')
    doc.add_paragraph("WXML使用wx:if=\"{{!isLoggedIn}}\"显示游客页面，wx:else显示登录后的用户中心。感叹号!表示取反。")
    table(doc, ["交互", "方法", "当前行为"], [
        ["档案管理", "toProfile", "打开生辰档案页"],
        ["订单入口", "showOrders", "弹出模拟说明"],
        ["查看卡包", "showCards", "显示模拟卡片数量"],
        ["开启服务卡", "activateCard", "显示规则说明"],
        ["赠礼", "giftCard", "显示模拟赠礼说明"],
        ["提醒开关", "togglePush", "只更新页面内存"],
        ["历史记录", "showRecords", "显示待接数据库说明"],
    ], [3.5, 4.3, 7.9])

    doc.add_heading("14｜浏览器预览与小程序的区别", level=1)
    table(doc, ["项目", "preview/index.html", "真实微信小程序"], [
        ["运行环境", "普通浏览器", "微信小程序运行环境"],
        ["页面结构", "HTML", "WXML"],
        ["页面样式", "CSS", "WXSS"],
        ["登录", "只能模拟", "wx.login＋Python后端"],
        ["分享", "alert模拟", "open-type=share"],
        ["支付", "不能验证", "wx.requestPayment"],
        ["用途", "快速讨论视觉", "真正开发、审核和上线"],
    ], [3.1, 6.1, 6.5])
    doc.add_paragraph("现在浏览器预览和微信小程序代码是两套页面，需要人工同步。取得正确的小程序AppID后，应以微信开发者工具为主要预览方式。")

    page_break(doc)
    doc.add_heading("15｜动手练习", level=1)
    doc.add_paragraph("每次只改一处，保存后先看TypeScript是否报错，再刷新预览。建议修改前复制原值，遇到问题容易恢复。")
    steps(doc, [
        "在home/index.ts中把绿金的status改成“比较合适”，观察页面文字变化。",
        "在home/index.wxml中把“今日公共五色排名”改为“今日五色总览”。",
        "在home/index.wxss中把.hero-title的font-size从62rpx改为54rpx。",
        "在chat/index.ts的questions中新增一个固定问题。",
        "在caikuxiang/index.ts中把单色价格临时改为69，观察前端合计变化，然后改回59。",
        "在settings/index.ts中把profileCompleteness改成100，观察进度条。",
        "给“我的”页面新增一个“联系客服”菜单行和点击提示。",
        "运行tsc --noEmit检查类型。",
    ])
    doc.add_heading("练习：把演示数据抽离", level=2)
    doc.add_paragraph("当你能读懂页面后，可以把PRODUCTS、guides、questions分别移到miniprogram/mock目录。页面通过import导入，避免文件过长。这是接真实API前很好的过渡。")
    code(doc, '// miniprogram/mock/products.ts\nexport const PRODUCTS = [\n  { id: "green", name: "绿金", price: 59 }\n];\n\n// 页面中\nimport { PRODUCTS } from "../../mock/products";')

    doc.add_heading("16｜常见错误和调试", level=1)
    table(doc, ["现象", "常见原因", "检查方法"], [
        ["页面空白", "WXML语法错误或页面未注册", "看开发者工具Console和app.json"],
        ["点击没反应", "bindtap名称与TS方法不一致", "搜索方法名并检查拼写"],
        ["文字不更新", "修改普通变量但没setData", "确认调用this.setData"],
        ["列表重复或错位", "wx:key不合适", "使用稳定唯一id或rank"],
        ["Tab跳转失败", "对Tab页面用了navigateTo", "改为wx.switchTab"],
        ["样式没生效", "class拼错或WXSS选择器错误", "对照WXML class与WXSS"],
        ["接口访问失败", "127.0.0.1、域名校验、后端未启动", "先访问/health"],
        ["TypeScript红线", "字段类型不一致或方法不存在", "运行tsc --noEmit"],
        ["浏览器是新页面，小程序没变", "两套预览未同步", "确认改的是preview还是miniprogram"],
    ], [3.5, 6.2, 6.0])
    code(doc, 'cd D:\\soft\\folk-guide\n.\\node_modules\\.bin\\tsc.cmd --noEmit', "前端类型检查")
    bullets(doc, [
        "编译错误先看第一条，后面的错误可能由第一条连锁造成。",
        "不要一次改十个文件；小步修改后立即检查。",
        "Console不要输出AppSecret、Token、手机号或完整生辰信息。",
    ])

    doc.add_heading("17｜当前项目前端代码地图", level=1)
    table(doc, ["想修改什么", "主要文件"], [
        ["底部四个入口和页面顺序", "miniprogram/app.json"],
        ["全局背景、卡片、按钮基础样式", "miniprogram/app.wxss"],
        ["今日五色排名和详情", "pages/home/index.ts、index.wxml、index.wxss"],
        ["AI国学问题、次数和消息", "pages/chat/index.ts、index.wxml、index.wxss"],
        ["商品、购物车和详情", "pages/caikuxiang/index.ts、index.wxml、index.wxss"],
        ["用户、服务卡、订单和设置", "pages/settings/index.ts、index.wxml、index.wxss"],
        ["生辰日期、时间和城市", "pages/profile/index.ts、index.wxml"],
        ["所有后端请求", "miniprogram/services/api.ts"],
        ["浏览器演示", "preview/index.html"],
        ["微信开发者工具AppID", "project.config.json"],
    ], [6.0, 9.7])

    doc.add_heading("18｜下一阶段学习路线", level=1)
    steps(doc, [
        "掌握对象、数组、函数、map/filter/reduce。",
        "掌握WXML的{{ }}、wx:for、wx:if、bindtap和data-*。",
        "掌握Flex、Grid、间距、字体和圆角。",
        "掌握Page的data、setData、onLoad和onShow。",
        "掌握Promise、async/await和wx.request。",
        "学习把页面演示数据替换成Python API返回值。",
        "学习加载、错误、空数据、未登录和无权限状态。",
        "最后再学习组件拆分、状态管理和性能优化。",
    ])
    doc.add_paragraph("对这个项目来说，最有价值的下一次实践是：先把“今日五色”的静态guides替换成一个Mock API，再替换成真正的Python接口。这样可以贯通前端学习和你已经理解的后端API。")

    page_break(doc)
    doc.add_heading("附录A｜常用语法速查", level=1)
    table(doc, ["写法", "含义"], [
        ["{{name}}", "在WXML显示data中的name"],
        ["wx:if=\"{{condition}}\"", "条件为真时显示"],
        ["wx:for=\"{{items}}\"", "循环显示数组"],
        ["bindtap=\"method\"", "点击时调用method"],
        ["data-id=\"{{item.id}}\"", "把id传给点击事件"],
        ["this.setData({ name: value })", "修改页面数据并刷新显示"],
        ["array.find(...) ", "找到一个项目"],
        ["array.map(...) ", "转换每个项目"],
        ["array.filter(...) ", "筛选项目"],
        ["array.reduce(...) ", "汇总金额或数量"],
        ["...array", "展开并复制数组内容"],
        ["condition ? a : b", "条件成立取a，否则取b"],
        ["async / await", "等待异步请求完成"],
        ["wx.navigateTo", "打开普通子页面"],
        ["wx.switchTab", "切换底部Tab"],
        ["wx.showModal", "显示微信弹窗"],
        ["wx.showToast", "显示短提示"],
    ], [6.4, 9.3])

    doc.add_heading("附录B｜当前核心前端源码", level=1)
    doc.add_paragraph("以下代码直接读取当前项目，方便对照学习。完整WXML和WXSS仍建议在VS Code中打开，因为编辑器能提供语法高亮和文件跳转。")
    files = [
        "miniprogram/app.json",
        "miniprogram/services/api.ts",
        "miniprogram/pages/home/index.ts",
        "miniprogram/pages/chat/index.ts",
        "miniprogram/pages/caikuxiang/index.ts",
        "miniprogram/pages/settings/index.ts",
        "miniprogram/pages/profile/index.ts",
    ]
    for path in files:
        doc.add_heading(path, level=2)
        code(doc, source(path))

    add_header_footer(doc)
    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build()
    document.save(OUTPUT)
    print(OUTPUT)
