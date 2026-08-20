from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "五行五色财库香小程序_代码与详细讲解_V1.0.docx"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def cell_text(cell, text: str, bold=False, color=None, size=9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = False
    for i, header in enumerate(headers):
        cell_text(result.rows[0].cells[i], header, True, "FFFFFF", 9)
        shade(result.rows[0].cells[i], "385F4D")
        if widths:
            result.rows[0].cells[i].width = Cm(widths[i])
    for row_no, values in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(values):
            cell_text(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])
            if row_no % 2:
                shade(cells[i], "F3F6F4")
    doc.add_paragraph()
    return result


def bullets(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def steps(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Number")


def code(doc, value: str, caption: str | None = None):
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        run.bold = True
        run.font.color.rgb = RGBColor(77, 102, 89)
    box = doc.add_table(rows=1, cols=1)
    box.autofit = False
    cell = box.cell(0, 0)
    shade(cell, "F1F4F2")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    lines = value.rstrip().splitlines() or [""]
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line.replace("\t", "    "))
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(47, 65, 56)
    doc.add_paragraph()


def page_break(doc):
    doc.add_page_break()


def set_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for name, size, color in [
        ("Title", 24, "27322D"),
        ("Heading 1", 16, "385F4D"),
        ("Heading 2", 12.5, "4D6659"),
        ("Heading 3", 10.5, "6D7E74"),
    ]:
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "五行五色财库香小程序｜代码说明书"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(145, 155, 149)
        footer = section.footer.paragraphs[0]
        footer.text = "内部开发说明 V1.0｜2026年8月"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(145, 155, 149)


def source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    set_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("五行五色财库香小程序\n代码与详细讲解")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("微信小程序前端＋Python FastAPI 后端｜内部开发说明 V1.0")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(113, 128, 120)
    doc.add_paragraph()
    intro = doc.add_table(rows=1, cols=1)
    shade(intro.cell(0, 0), "E9F0EB")
    p = intro.cell(0, 0).paragraphs[0]
    p.add_run("文档说明：").bold = True
    p.add_run("本文解释当前项目源码、页面交互、后端接口和后续接入点。当前四大板块以前端演示为主，Python后端只完成基础登录、档案、每日示例和问答占位，尚未连接真实规则、支付、订单、权益、向量库和大模型。")
    doc.add_paragraph()
    table(doc, ["项目", "内容"], [
        ["项目目录", str(ROOT)],
        ["前端技术", "微信原生小程序：TypeScript、WXML、WXSS、JSON"],
        ["后端技术", "Python 3.12、FastAPI、SQLAlchemy、SQLite、JWT"],
        ["浏览器预览", "preview/index.html，本机 http://127.0.0.1:8088"],
        ["当前测试", "后端2项测试通过；前端TypeScript类型检查通过"],
        ["当前阶段", "四大板块可交互前端原型＋基础后端框架"],
    ], [4.2, 11.5])

    doc.add_heading("目录", level=1)
    for item in [
        "01 项目目标与当前完成度", "02 总体架构", "03 目录结构", "04 本地运行方式",
        "05 小程序全局配置", "06 前端请求与微信登录", "07 今日五色页面", "08 AI国学页面",
        "09 财库香页面", "10 我的页面", "11 生辰档案页面", "12 Python后端入口与接口",
        "13 数据库模型", "14 登录与安全", "15 每日内容与问答占位", "16 测试",
        "17 浏览器预览", "18 当前限制与风险", "19 后续开发顺序", "附录A API速查",
        "附录B 数据字段建议", "附录C 当前核心源码",
    ]:
        doc.add_paragraph(item)

    page_break(doc)
    doc.add_heading("01｜项目目标与当前完成度", level=1)
    doc.add_paragraph("项目以“今日五色”为免费日活和分享入口，以“AI国学”为个人数字服务，以“财库香”为商品成交和复购入口，以“我的”为档案、订单和权益管理中心。")
    table(doc, ["模块", "业务作用", "当前代码状态", "尚未接入"], [
        ["今日五色", "免费日活、分享、香品导流", "五色排名、详情、分享、提醒、登录提示前端完成", "每日后台、规则引擎、真实分享图、订阅消息"],
        ["AI国学", "个人化、知识问答、30天服务", "3天体验、次数、个人结果、固定问题、引用和反馈前端完成", "向量库、模型、档案计算、权益与限流"],
        ["财库香", "199元套装、59元补充和复购", "商品、详情、购物车和模拟结算前端完成", "库存、支付、订单、物流、退款与发卡"],
        ["我的", "用户、档案、卡包、订单和记录", "用户中心、卡包、订单、提醒和记录入口前端完成", "真实用户、状态机、赠礼、历史与注销"],
    ], [2.5, 3.5, 5.1, 4.8])
    doc.add_paragraph("重要理解：当前页面中大部分业务数据仍写在TypeScript演示数据中。浏览器能看到完整体验，不代表真实业务闭环已经完成。")

    doc.add_heading("02｜总体架构", level=1)
    code(doc, "微信小程序前端\n├─ 今日五色\n├─ AI国学\n├─ 财库香\n└─ 我的\n       │ HTTPS / JSON\n       ▼\nPython FastAPI 后端\n├─ 微信登录与JWT\n├─ 生辰档案\n├─ 每日内容占位\n├─ AI问答占位\n└─ 账号删除\n       │\n       ├─ SQLite（开发）/ PostgreSQL（生产建议）\n       ├─ 外部向量库（后续）\n       ├─ 大模型API（后续）\n       ├─ 微信支付与物流（后续）\n       └─ 运营后台与规则引擎（后续）")
    bullets(doc, [
        "前端只负责展示、输入、页面状态和调用后端，不能保存AppSecret或模型密钥。",
        "确定性排盘和五色排名应由Python规则引擎计算，不让AI临时决定。",
        "AI国学应先检索审核知识库，再结合结构化个人结果回答。",
        "商品订单、退款和服务卡必须使用后端状态机，不能由前端自行发放。",
    ])

    doc.add_heading("03｜目录结构", level=1)
    code(doc, "folk-guide/\n├─ backend/\n│  ├─ app/                 FastAPI应用\n│  ├─ tests/               后端测试\n│  ├─ requirements.txt     Python依赖\n│  └─ Dockerfile           容器部署\n├─ miniprogram/\n│  ├─ pages/home/          今日五色\n│  ├─ pages/chat/          AI国学\n│  ├─ pages/caikuxiang/    财库香\n│  ├─ pages/settings/      我的\n│  ├─ pages/profile/       生辰档案\n│  └─ services/api.ts      后端请求封装\n├─ preview/index.html      浏览器交互预览\n├─ docs/                   项目Word文档\n├─ tools/                  文档生成脚本\n├─ project.config.json     微信开发者工具配置\n├─ package.json            前端开发依赖\n└─ README.md               启动说明")
    table(doc, ["文件类型", "作用"], [
        [".ts", "页面逻辑、模拟数据、事件处理、后端请求"],
        [".wxml", "微信小程序页面结构，类似HTML"],
        [".wxss", "微信小程序样式，rpx为响应式尺寸单位"],
        [".json", "页面标题、分享能力、底部导航等配置"],
        [".py", "Python接口、业务服务、数据模型和测试"],
    ], [3.5, 12.2])

    doc.add_heading("04｜本地运行方式", level=1)
    doc.add_heading("4.1 启动Python后端", level=2)
    code(doc, "cd D:\\soft\\folk-guide\\backend\n.\\.venv\\Scripts\\Activate.ps1\nuvicorn app.main:app --reload")
    doc.add_paragraph("启动后访问 http://127.0.0.1:8000/health 检查状态，访问 http://127.0.0.1:8000/docs 查看Swagger接口文档。")
    doc.add_heading("4.2 启动浏览器预览", level=2)
    code(doc, "cd D:\\soft\\folk-guide\n.\\backend\\.venv\\Scripts\\python.exe -m http.server 8088 --directory preview")
    doc.add_paragraph("然后访问 http://127.0.0.1:8088。浏览器预览是独立HTML演示，不等于微信小程序真实运行环境。")
    doc.add_heading("4.3 前端类型检查", level=2)
    code(doc, "cd D:\\soft\\folk-guide\n.\\node_modules\\.bin\\tsc.cmd --noEmit")
    doc.add_heading("4.4 后端测试", level=2)
    code(doc, "cd D:\\soft\\folk-guide\\backend\n.\\.venv\\Scripts\\python.exe -m pytest -q")

    page_break(doc)
    doc.add_heading("05｜小程序全局配置", level=1)
    doc.add_paragraph("miniprogram/app.json声明页面加载顺序、导航栏和底部四个Tab。第一个页面是启动页，因此今日五色必须排在首位。")
    code(doc, source("miniprogram/app.json"), "文件：miniprogram/app.json")
    bullets(doc, [
        "pages数组中的每个页面都必须真实存在。",
        "tabBar页面之间使用wx.switchTab跳转，不能使用wx.navigateTo。",
        "profile不是Tab页面，因此从“我的”进入档案时使用wx.navigateTo。",
        "project.config.json中的AppID当前为空，导入微信工具前必须填写真正的小程序AppID。",
    ])

    doc.add_heading("06｜前端请求与微信登录", level=1)
    doc.add_paragraph("miniprogram/services/api.ts统一封装wx.request和登录。页面不应重复实现Token拼接和错误处理。")
    code(doc, source("miniprogram/services/api.ts"), "文件：miniprogram/services/api.ts")
    steps(doc, [
        "页面调用ensureLogin。",
        "前端调用wx.login取得一次性code。",
        "code发送到POST /api/auth/wechat。",
        "Python后端向微信换取openid，创建本地用户。",
        "后端签发JWT，前端存入微信本地存储。",
        "后续请求携带Authorization: Bearer <token>。",
    ])
    doc.add_paragraph("生产注意：API_BASE不能继续使用127.0.0.1，必须替换为已备案HTTPS域名，并加入微信小程序request合法域名。")

    doc.add_heading("07｜今日五色页面", level=1)
    doc.add_heading("7.1 页面职责", level=2)
    bullets(doc, [
        "不登录也能完整查看公共五色。",
        "展示日期、农历、节气、当日干支、第一至第五排名。",
        "点击颜色展开适合事项、阻力、行动建议和对应香品。",
        "个人结果、收藏、记录和提醒等行为再触发登录。",
        "分享进入当天公共页面，不先显示登录墙。",
    ])
    doc.add_heading("7.2 数据定义", level=2)
    code(doc, "type ColorGuide = {\n  rank: number;\n  name: string;\n  element: string;\n  tone: string;\n  status: string;\n  suitable: string[];\n  resistance: string;\n  advice: string;\n  incense: string;\n  scent: string;\n};")
    doc.add_paragraph("当前guides数组是演示数据。未来应由GET /api/public-guides/today返回，前端不能继续写死日期和排名。")
    doc.add_heading("7.3 主要交互", level=2)
    table(doc, ["方法", "作用", "后续变化"], [
        ["toggleGuide", "展开或收起某个颜色详情", "保持前端处理"],
        ["showPersonalLogin", "个人结果入口显示登录提示", "连接真实登录和权益检查"],
        ["showSubscribe", "显示提醒说明", "连接wx.requestSubscribeMessage"],
        ["toIncense", "跳转财库香Tab", "传入商品ID可进一步定位商品"],
        ["onShareAppMessage", "设置微信分享标题和路径", "增加当日内容版本与分享图"],
    ], [3.5, 5.0, 7.2])
    code(doc, source("miniprogram/pages/home/index.ts"), "当前逻辑：miniprogram/pages/home/index.ts")

    doc.add_heading("08｜AI国学页面", level=1)
    doc.add_heading("8.1 当前前端规则", level=2)
    table(doc, ["用户类型", "服务期", "问答次数", "七日比较", "个人结果"], [
        ["新用户体验", "3天", "每日20次", "每日2次", "每日生成1次，可无限查看"],
        ["正式服务用户", "30天", "每日50次", "每日5次", "每日生成1次，可无限查看"],
    ], [3.3, 2.4, 3.2, 3.1, 4.0])
    doc.add_paragraph("当前页面默认显示新客体验第1天。次数只存在页面内存里，刷新后会恢复；后续必须由后端按用户、日期和功能类型统一计数。")
    doc.add_heading("8.2 页面组成", level=2)
    bullets(doc, [
        "3天体验与剩余次数。",
        "档案完整度与管理入口。",
        "个人主色、辅助色、建议配色和完整个人结果。",
        "今日行动、事业合作、财务事务、文化与档案四组固定问题。",
        "自由输入、模拟回答、参考资料和反馈纠错。",
    ])
    doc.add_heading("8.3 问答流程", level=2)
    code(doc, "用户问题\n  ↓\n身份与权益检查\n  ↓\n问题分类与安全检测\n  ↓\n读取今日个人结构化结果（按需）\n  ↓\n检索外部向量知识库\n  ↓\n组合系统提示词与引用片段\n  ↓\n调用模型\n  ↓\n输出安全检测与引用校验\n  ↓\n扣减次数、保存记录、返回页面")
    doc.add_paragraph("当前ANSWERS常量只是前端演示，正式接入后必须删除前端答案表，改为POST /api/ai/chat。")
    code(doc, source("miniprogram/pages/chat/index.ts"), "当前逻辑：miniprogram/pages/chat/index.ts")

    doc.add_heading("09｜财库香页面", level=1)
    doc.add_heading("9.1 商品结构", level=2)
    table(doc, ["商品", "规格", "价格", "权益"], [
        ["五色旗舰套装", "五色各30支，共150支", "199元包邮", "确认收货后发30天服务卡"],
        ["黑金／黄金／绿金／红金／白金", "任一颜色30支", "59元包邮", "单色累计确认收货实付满199元发卡"],
    ], [4.4, 4.5, 3.1, 4.6])
    doc.add_heading("9.2 前端购物车", level=2)
    doc.add_paragraph("addItem在页面内存中合并相同商品、计算数量和金额；changeQuantity负责加减数量并删除数量为0的项目。")
    code(doc, "const existing = items.find((item) => item.id === product.id);\nif (existing) existing.quantity += 1;\nelse items.push({ ...product, quantity: 1 });\n\nconst cartTotal = items.reduce(\n  (total, item) => total + item.price * item.quantity,\n  0\n);")
    doc.add_paragraph("生产环境不能相信前端价格、库存和权益。结算时前端只提交SKU和数量，Python后端必须重新读取当前价格、优惠、库存、运费和用户资格。")
    doc.add_heading("9.3 订单后续流程", level=2)
    code(doc, "购物车\n→ 后端创建待支付订单\n→ 微信支付\n→ 支付回调验签\n→ 待发货\n→ 填写物流\n→ 用户/系统确认收货\n→ 判断退款和异常\n→ 发放30天服务卡")
    code(doc, source("miniprogram/pages/caikuxiang/index.ts"), "当前逻辑：miniprogram/pages/caikuxiang/index.ts")

    doc.add_heading("10｜我的页面", level=1)
    doc.add_paragraph("“我的”集中管理身份、档案、订单、服务卡、赠礼、提醒、结果记录、问答记录和隐私。当前所有数字均为演示状态。")
    table(doc, ["区域", "当前表现", "未来接口"], [
        ["用户头部", "昵称、体验天数、问答次数、待开启卡", "GET /api/me/summary"],
        ["本人档案", "完整度85%", "GET /api/profiles/current"],
        ["服务卡", "体验卡使用中＋30天卡待开启", "GET /api/entitlements"],
        ["订单", "待付款、待发货、待收货、售后", "GET /api/orders"],
        ["每日提醒", "开关和偏好时间", "PUT /api/preferences/notification"],
        ["历史记录", "结果与问答入口", "GET /api/results/history、GET /api/chat/history"],
    ], [3.3, 6.0, 6.3])
    code(doc, source("miniprogram/pages/settings/index.ts"), "当前逻辑：miniprogram/pages/settings/index.ts")

    doc.add_heading("11｜生辰档案页面", level=1)
    doc.add_paragraph("档案页目前支持公历出生日期、可选出生时间和可选出生城市。它通过Python接口保存到SQLite。")
    bullets(doc, [
        "出生日期必填且不能晚于今天。",
        "出生时间不知道时允许留空，后续生成简化结果。",
        "出生地第一版只收城市，不收医院或详细地址。",
        "真实姓名和手机号不应作为排盘必填项。",
        "后续需要增加公历/农历、性别可选、时区和隐私说明。",
    ])
    code(doc, source("miniprogram/pages/profile/index.ts"), "文件：miniprogram/pages/profile/index.ts")

    page_break(doc)
    doc.add_heading("12｜Python后端入口与接口", level=1)
    doc.add_paragraph("backend/app/main.py创建FastAPI应用、配置CORS、建表并声明接口。开发阶段使用SQLite，启动时自动create_all。生产环境应改用Alembic迁移。")
    table(doc, ["方法", "路径", "作用", "认证"], [
        ["GET", "/health", "健康检查", "否"],
        ["POST", "/api/auth/wechat", "微信code换登录Token", "否"],
        ["GET", "/api/profiles/current", "读取本人档案", "是"],
        ["PUT", "/api/profiles/current", "创建或修改本人档案", "是"],
        ["DELETE", "/api/profiles/current", "删除本人档案", "是"],
        ["GET", "/api/daily", "读取并缓存每日演示内容", "是"],
        ["POST", "/api/chat", "基础问答与风险词占位", "是"],
        ["DELETE", "/api/account", "删除账号及关联演示数据", "是"],
    ], [1.6, 4.3, 6.8, 2.0])
    code(doc, source("backend/app/main.py"), "文件：backend/app/main.py")

    doc.add_heading("13｜数据库模型", level=1)
    table(doc, ["数据表", "作用", "主要字段"], [
        ["users", "微信用户", "id、openid、created_at"],
        ["birth_profiles", "本人出生档案", "user_id、日期、时间、城市、性别"],
        ["daily_guidance", "按用户和日期缓存每日内容", "user_id、guidance_date、payload_json"],
        ["chat_messages", "问答记录", "user_id、question、answer、created_at"],
    ], [3.2, 5.0, 7.5])
    doc.add_paragraph("这些表只覆盖最早的基础骨架，与当前完整产品相比还缺少公共每日内容、规则版本、商品、SKU、购物车、订单、支付、物流、退款、权益卡、赠礼、订阅偏好、反馈和知识库引用等表。")
    code(doc, source("backend/app/models.py"), "文件：backend/app/models.py")

    doc.add_heading("14｜登录与安全", level=1)
    doc.add_heading("14.1 JWT签发", level=2)
    code(doc, source("backend/app/auth.py"), "文件：backend/app/auth.py")
    bullets(doc, [
        "JWT_SECRET生产环境必须使用高强度随机值并放在环境变量。",
        "微信AppSecret和模型API Key只能存在服务器，不能写入小程序。",
        "当前Token有效期为7天，正式上线应考虑刷新机制、封禁和注销失效。",
        "生产环境CORS不能使用通配符。",
        "手机号不是基础功能必需信息时不应强制收集。",
    ])
    doc.add_heading("14.2 开发模拟登录", level=2)
    doc.add_paragraph("未配置微信AppID且ENVIRONMENT=development时，后端把传入code哈希成模拟openid，便于本地联调。生产环境必须关闭这一分支。")

    doc.add_heading("15｜每日内容与问答占位", level=1)
    doc.add_paragraph("backend/app/services.py包含微信code交换、每日示例生成、安全词占位和JSON序列化。每日示例只根据生日和日期生成稳定配色，不是正式五色算法。")
    code(doc, source("backend/app/services.py"), "文件：backend/app/services.py")
    doc.add_paragraph("正式开发时建议将services.py拆分为wechat_service.py、calendar_service.py、public_guide_service.py、bazi_service.py、knowledge_service.py、llm_service.py和safety_service.py。")

    doc.add_heading("16｜测试", level=1)
    doc.add_paragraph("test_api_flow.py验证登录、缺档案、保存档案、每日内容、普通问答、高风险问答和注销的完整流程。")
    code(doc, source("backend/tests/test_api_flow.py"), "文件：backend/tests/test_api_flow.py")
    bullets(doc, [
        "当前测试使用项目SQLite文件，不是隔离测试数据库；后续应改成临时数据库fixture。",
        "需要补充未授权、Token过期、非法日期、重复请求和并发写入测试。",
        "商城必须覆盖重复支付回调、退款、部分退款、确认收货、发卡和撤卡测试。",
        "AI必须覆盖次数限制、敏感问题、引用缺失、模型超时和向量库不可用测试。",
    ])

    doc.add_heading("17｜浏览器预览", level=1)
    doc.add_paragraph("preview/index.html把四个页面做成一个独立HTML，用于无需AppID快速讨论视觉和交互。它没有复用WXML/WXSS，也不会自动和小程序同步，因此每次正式修改小程序页面后，需要同步更新预览或逐步停止维护该预览。")
    bullets(doc, [
        "优点：打开快、方便演示、无需微信账号。",
        "缺点：不是微信运行环境，不能验证wx.login、分享、订阅消息、支付和Tab行为。",
        "建议：取得真实小程序AppID后，以微信开发者工具模拟器为主，浏览器预览只用于早期视觉讨论。",
    ])

    doc.add_heading("18｜当前限制与风险", level=1)
    table(doc, ["问题", "当前情况", "处理建议"], [
        ["小程序AppID", "project.config.json为空", "注册全新小程序并填写正确AppID"],
        ["日期和五色", "前端写死演示日期与排名", "连接公共每日内容接口"],
        ["AI国学", "前端模拟答案与次数", "后端权益、分类、RAG、模型和限流"],
        ["商城", "购物车只在页面内存", "后端价格校验、订单、支付、物流和退款"],
        ["权益卡", "页面模拟1张待开启卡", "建立权益状态机与订单事件发卡"],
        ["推送", "仅显示说明和偏好", "验证微信订阅消息模板与次数规则"],
        ["数据库", "SQLite＋自动建表", "PostgreSQL＋Alembic＋备份"],
        ["个人信息", "基础档案已存在", "隐私政策、单独同意、删除和审计"],
        ["合规表达", "已有免责声明和温和文案", "后台审核、禁用词、举报和人工纠错"],
    ], [3.0, 5.1, 7.8])

    doc.add_heading("19｜后续开发顺序", level=1)
    steps(doc, [
        "取得真正的小程序AppID，在微信开发者工具中完成真机前端验证。",
        "将四个页面中的演示数据抽成统一Mock服务，避免散落在页面文件。",
        "实现公共每日五色后台：人工录入、审核、定时发布和首页接口。",
        "完善用户、档案和新客3天体验的真实状态。",
        "设计AI分类、安全、次数、向量库和模型接口，再接AI国学页面。",
        "建立商品、SKU、库存、购物车、订单和微信支付。",
        "实现确认收货发卡、单色累计发卡、卡包、激活、冻结和赠礼。",
        "实现订阅提醒、分享卡、统计、历史记录和用户数据删除。",
        "切换PostgreSQL、加入Alembic、日志、监控、备份和正式部署。",
        "进行种子用户、支付退款、权益和内容合规全流程测试。",
    ])

    page_break(doc)
    doc.add_heading("附录A｜API速查与建议新增接口", level=1)
    table(doc, ["模块", "建议接口", "说明"], [
        ["公共五色", "GET /api/public-guides/today", "返回今天已发布的公共五色"],
        ["公共五色", "GET /api/public-guides/{date}", "历史内容，按权限开放"],
        ["个人结果", "GET /api/personal-results/today", "检查档案与权益后返回缓存结果"],
        ["AI问答", "POST /api/ai/chat", "分类、检索、生成、扣次数"],
        ["AI次数", "GET /api/ai/quota", "返回问答和七日比较剩余次数"],
        ["商品", "GET /api/products", "商品、SKU、价格、库存和上下架"],
        ["购物车", "GET/PUT /api/cart", "登录用户购物车"],
        ["订单", "POST /api/orders", "后端重新计算价格并创建订单"],
        ["支付", "POST /api/payments/wechat", "生成微信支付参数"],
        ["服务卡", "GET /api/entitlements", "卡包和使用中服务"],
        ["服务卡", "POST /api/entitlements/{id}/activate", "开启服务卡"],
        ["赠礼", "POST /api/entitlements/{id}/gift", "生成赠礼Token"],
        ["我的", "GET /api/me/summary", "聚合用户中心数据"],
    ], [3.0, 6.5, 6.4])

    doc.add_heading("附录B｜建议新增的数据表", level=1)
    table(doc, ["表", "关键字段", "用途"], [
        ["public_guides", "date、ranking_json、status、rule_version", "公共每日五色"],
        ["products / skus", "name、price、stock、status", "商品与库存"],
        ["orders / order_items", "status、amount、payment_no", "订单与明细"],
        ["shipments / refunds", "tracking_no、refund_status", "物流与售后"],
        ["entitlements", "owner、status、issued_at、service_end_at", "30天服务卡"],
        ["entitlement_gifts", "gift_token、claimed_by、claimed_at", "赠礼领取"],
        ["personal_results", "user、date、result_json、rule_version", "个人每日结果缓存"],
        ["ai_usage", "user、date、type、count、cost", "次数和成本"],
        ["knowledge_references", "source、chunk_id、copyright", "回答引用追溯"],
        ["user_feedback", "message_id、type、comment", "回答反馈与纠错"],
        ["notification_preferences", "enabled、preferred_time", "提醒偏好"],
    ], [3.4, 6.2, 6.3])

    page_break(doc)
    doc.add_heading("附录C｜当前核心源码", level=1)
    doc.add_paragraph("以下代码从当前项目文件自动读取，便于离线查看。WXML和WXSS篇幅较长，本附录收录核心TypeScript、Python及全局配置；完整页面结构和样式以项目源码为准。")
    appendix_files = [
        "miniprogram/app.json",
        "miniprogram/services/api.ts",
        "miniprogram/pages/home/index.ts",
        "miniprogram/pages/chat/index.ts",
        "miniprogram/pages/caikuxiang/index.ts",
        "miniprogram/pages/settings/index.ts",
        "miniprogram/pages/profile/index.ts",
        "backend/app/config.py",
        "backend/app/database.py",
        "backend/app/schemas.py",
        "backend/app/auth.py",
        "backend/app/models.py",
        "backend/app/services.py",
        "backend/app/main.py",
        "backend/tests/test_api_flow.py",
    ]
    for path in appendix_files:
        doc.add_heading(path, level=2)
        code(doc, source(path))

    add_header_footer(doc)
    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build()
    document.save(OUTPUT)
    print(OUTPUT)
