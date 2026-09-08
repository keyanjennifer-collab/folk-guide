"""专业配置校验和公共五色规则执行骨架测试。"""

from datetime import date

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from fastapi.testclient import TestClient

from app.daily_color_context import build_public_rule_input
from app.daily_color_rule_config import (
    ElementRelationScores,
    PublicFactorWeights,
    PublicRuleConfiguration,
    SOLAR_TERMS,
    TendencyThresholds,
)
from app.daily_color_rule_engine import calculate_public_rule
from app.daily_color_rule_template import build_rule_template, parse_rule_template
from app.daily_color_research_v1 import PUBLIC_RESEARCH_CONFIG
from app.main import app


ADMIN_HEADERS = {"X-Admin-Key": "dev-admin-key", "X-Admin-Name": "pytest-rule-admin"}


def term_map() -> dict[str, str]:
    """只用于自动化测试的完整映射，不代表项目正式专业口径。"""
    result = {}
    elements = ("木", "火", "土", "金", "水")
    for index, term in enumerate(SOLAR_TERMS):
        result[term] = elements[index % len(elements)]
    return result


def make_test_configuration(version: str = "pytest-public-rule-v1") -> PublicRuleConfiguration:
    """测试配置仅验证引擎数学与确定性，不能进入真实业务数据库。"""
    return PublicRuleConfiguration(
        version=version,
        status="test_only",
        algorithm_summary="pytest只使用日干关系验证规则执行骨架",
        factor_weights=PublicFactorWeights(
            year_stem=0, year_branch=0, month_stem=0, month_branch=0,
            day_stem=100, day_branch=0, solar_term=0,
        ),
        relation_scores=ElementRelationScores(
            same_element=40,
            candidate_generates_reference=50,
            reference_generates_candidate=0,
            candidate_controls_reference=-20,
            reference_controls_candidate=20,
        ),
        solar_term_elements=term_map(),
        tie_break_color_order=["白金", "绿金", "黑金", "红金", "黄金"],
        tendency_thresholds=TendencyThresholds(
            strong_support_min=45, support_min=30, balanced_min=10, caution_min=-10,
        ),
    )


def test_factor_weights_must_total_one_hundred():
    with pytest.raises(ValidationError, match="合计必须等于100"):
        PublicFactorWeights(
            year_stem=0, year_branch=0, month_stem=0, month_branch=0,
            day_stem=50, day_branch=0, solar_term=0,
        )


def test_reviewed_configuration_requires_reviewer_references_and_cases():
    data = make_test_configuration().model_dump()
    data["status"] = "expert_reviewed"
    with pytest.raises(ValidationError, match="审核人"):
        PublicRuleConfiguration(**data)


def test_public_engine_is_deterministic_and_keeps_test_only_status():
    config = make_test_configuration()
    rule_input = build_public_rule_input(date(2026, 8, 19), rule_version=config.version)
    first = calculate_public_rule(rule_input, config)
    second = calculate_public_rule(rule_input, config)

    assert first == second
    assert first.result.rule_status == "test_only"
    assert [item.color for item in first.result.ranking.items] == ["黑金", "绿金", "黄金", "红金", "白金"]
    assert [item.rule_score for item in first.result.ranking.items] == [50, 40, 20, 0, -20]
    assert first.result.config_fingerprint == config.fingerprint()
    assert all(trace.contributions[0].factor_code == "PUBLIC_DAY_STEM" for trace in first.traces)


def test_rule_input_and_configuration_versions_must_match():
    config = make_test_configuration()
    wrong_input = build_public_rule_input(date(2026, 8, 19), rule_version="another-version")
    with pytest.raises(ValueError, match="版本不一致"):
        calculate_public_rule(wrong_input, config)


def test_formal_public_rule_uses_only_beijing_day_branch_fixed_relations():
    """乙酉日只取酉金：我生、同我、克我、生我、我克依次排名。"""
    rule_input = build_public_rule_input(
        date(2026, 9, 8), rule_version=PUBLIC_RESEARCH_CONFIG.version
    )
    result = calculate_public_rule(rule_input, PUBLIC_RESEARCH_CONFIG)

    assert rule_input.calendar.pillars.day.text == "乙酉"
    assert rule_input.calendar.pillars.day.branch_primary_element == "金"
    assert [item.color for item in result.result.ranking.items] == [
        "黑金", "白金", "红金", "黄金", "绿金",
    ]
    assert all(item.basis_codes == ["PUBLIC_DAY_BRANCH"] for item in result.result.ranking.items)


def test_excel_template_contains_public_and_personal_review_sheets():
    import io

    workbook = load_workbook(io.BytesIO(build_rule_template()))
    assert {
        "填写说明", "版本信息", "公共因子权重", "五行关系分值", "节气五行映射",
        "同分排序", "趋势阈值", "公共测试案例", "个人规则待确认", "个人测试案例",
    } == set(workbook.sheetnames)
    assert workbook["公共因子权重"]["C2"].value is None
    assert workbook["五行关系分值"]["C2"].value is None


def test_filled_test_only_excel_can_be_parsed_and_fingerprinted():
    import io

    config = make_test_configuration()
    workbook = load_workbook(io.BytesIO(build_rule_template()))
    version = workbook["版本信息"]
    version["B2"] = config.version
    version["B3"] = config.status
    version["B5"] = config.algorithm_summary

    weights = config.factor_weights.by_code()
    for row in workbook["公共因子权重"].iter_rows(min_row=2):
        row[2].value = weights[row[0].value]

    scores = config.relation_scores.by_code()
    for row in workbook["五行关系分值"].iter_rows(min_row=2):
        row[2].value = scores[row[0].value]

    for row in workbook["节气五行映射"].iter_rows(min_row=2):
        row[1].value = config.solar_term_elements[row[0].value]
    for index, color in enumerate(config.tie_break_color_order, start=2):
        workbook["同分排序"].cell(row=index, column=2, value=color)
    for row in workbook["趋势阈值"].iter_rows(min_row=2):
        row[2].value = getattr(config.tendency_thresholds, row[0].value)

    stream = io.BytesIO()
    workbook.save(stream)
    parsed = parse_rule_template(stream.getvalue())

    assert parsed == config
    assert parsed.fingerprint() == config.fingerprint()


def test_admin_can_download_template_and_blank_template_fails_validation():
    with TestClient(app) as client:
        downloaded = client.get("/api/admin/daily-color-rules/template.xlsx", headers=ADMIN_HEADERS)
        assert downloaded.status_code == 200
        assert downloaded.content.startswith(b"PK")

        blank = client.post(
            "/api/admin/daily-color-rules/validate",
            headers=ADMIN_HEADERS,
            files={"file": ("rules.xlsx", downloaded.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert blank.status_code == 422
        assert blank.json()["detail"] == "版本号不能为空"
