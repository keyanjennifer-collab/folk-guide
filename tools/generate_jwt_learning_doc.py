"""生成结合当前项目实际代码的 JWT 与微信登录学习文档。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from generate_code_guide_doc import ROOT, set_styles


OUTPUT = ROOT / "docs" / "JWT与微信小程序登录_代码学习手册_V1.0.docx"


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


doc = Document()
set_styles(doc)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("JWT 与微信小程序登录\n代码学习手册")
subtitle = doc.add_paragraph("结合五行五色项目实际代码｜V1.0｜2026年8月13日")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("一、JWT是什么", level=1)
doc.add_paragraph("JWT（JSON Web Token）可以理解为后端签发的一张有防伪签名、并且会过期的登录凭证。用户完成微信登录后，小程序保存JWT；以后访问档案、AI问答和本人历史时，把JWT放进Authorization请求头。后端验证签名和有效期，再从sub字段读取内部user_id。")
doc.add_paragraph("JWT通常由Header、Payload、Signature三段组成。Payload可以被读取，因此不能放AppSecret、手机号、生辰资料、密码等敏感数据。本项目只放内部user_id和过期时间。")

doc.add_heading("二、本项目代码在哪里", level=1)
table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"
table.rows[0].cells[0].text = "文件"; table.rows[0].cells[1].text = "作用"
for path, purpose in [
    ("backend/app/auth.py", "create_token签发JWT，current_user验证JWT并加载用户"),
    ("backend/app/main.py", "POST /api/auth/wechat登录，GET /api/users/me验证账号"),
    ("backend/app/services.py", "用微信临时code换openid；手机号code换号码"),
    ("miniprogram/services/api.ts", "保存JWT，并在Authorization: Bearer中自动携带"),
    ("miniprogram/app.ts", "应用启动时从本地缓存恢复token"),
]:
    cells = table.add_row().cells; cells[0].text = path; cells[1].text = purpose

doc.add_heading("三、完整登录流程", level=1)
doc.add_paragraph("wx.login() → 小程序取得临时code → POST /api/auth/wechat → Python后端调用微信code2session → 得到openid → users表查询或创建账号 → create_token(user.id) → 前端保存JWT → 后续请求携带Bearer JWT → current_user验证后得到当前用户。")

doc.add_heading("四、JWT不是什么", level=1)
bullets(doc, [
    "JWT不是数据库，用户档案和问答记录仍保存在数据库。",
    "JWT默认不是加密文本，不能存放隐私信息。",
    "JWT不是永久登录，本项目开发配置有效期为7天，过期后重新微信登录。",
    "JWT不能代替HTTPS；正式服务器必须配置HTTPS。",
])

doc.add_heading("五、手机号为什么是另一套授权", level=1)
doc.add_paragraph("wx.login只确认微信用户身份，不会自动返回手机号。用户必须主动点击open-type=getPhoneNumber按钮。按钮返回一次性code，后端再调用微信接口换取手机号并保存。前端不能直接提交一个手机号让后端相信，否则任何人都能伪造绑定。")

doc.add_heading("六、你现在需要完成的微信平台事项", level=1)
bullets(doc, [
    "准备一个未被其他公众平台占用的邮箱，并完成微信公众平台注册。",
    "注册类型选择小程序，按实际经营主体选择个人、个体工商户或企业；若计划商城支付，优先使用可认证的经营主体。",
    "完成管理员微信绑定、主体信息和小程序名称等资料。",
    "在开发管理中取得AppID；AppSecret只保存到后端环境变量，不能发到前端或提交代码仓库。",
    "在微信开发者工具中填写AppID，当前项目project.config.json的appid还是空值。",
    "准备已备案HTTPS域名和服务器后，把域名加入request合法域名；127.0.0.1只能用于本机开发。",
    "配置用户隐私保护指引，说明手机号和生辰档案的用途、保存和删除方式。",
    "手机号能力、微信认证和微信支付是否可用受主体类型及微信当期规则影响，应以注册后台显示为准。",
])

doc.add_heading("七、安全检查清单", level=1)
bullets(doc, [
    "生产环境更换JWT_SECRET和ADMIN_API_KEY。",
    "AppSecret仅放服务器.env或密钥管理服务。",
    "接口始终通过current_user取得user_id，不接受前端自报user_id。",
    "手机号展示时脱敏，注销账号时一并删除个人数据。",
    "正式上线前从SQLite迁移到PostgreSQL，并配置备份。",
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
