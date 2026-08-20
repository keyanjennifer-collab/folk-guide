"""生成全项目逐层代码与参数、链路学习手册。"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from generate_code_guide_doc import ROOT, set_styles


OUTPUT = ROOT / "docs" / "五行五色财库香小程序_全项目代码参数与链路详解_V2.0.docx"


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def steps(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(doc, headers, rows, widths=None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for index, header in enumerate(headers):
        result.rows[0].cells[index].text = header
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    if widths:
        for row in result.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return result


def code(doc, value):
    p = doc.add_paragraph()
    p.style = "No Spacing"
    run = p.add_run(value)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


doc = Document()
set_styles(doc)
section = doc.sections[0]
section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("五行五色财库香小程序\n全项目代码、参数与链路详解")
subtitle = doc.add_paragraph("微信原生小程序 + Python FastAPI｜学习手册 V2.0｜2026年8月13日")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("阅读定位：假设读者能看懂基础Python，但对前端、JWT、数据库ORM和完整Web链路还不熟悉。本手册按照“界面事件 → 请求 → 后端路由 → 服务函数 → 数据库 → 返回界面”的顺序解释。")

doc.add_heading("目录与阅读建议", level=1)
bullets(doc, [
    "第一次阅读：第1—5章，理解架构、文件类型、配置、数据库和请求机制。",
    "学习登录档案：第6—9章；学习今日五色：第10章；学习知识库与AI：第11—13章。",
    "前端逐页阅读：第14章；不完善项和正式上线差距：第16章。",
    "代码里的参数名保持英文，中文解释说明它是什么、从哪里来、影响什么。",
])

doc.add_heading("第1章 项目到底由哪些部分组成", level=1)
doc.add_paragraph("项目采用“微信原生薄前端 + Python厚后端”。小程序负责展示、收集点击和输入、保存JWT；Python负责身份、权限、历法、数据校验、知识库、AI次数和数据库写入。")
table(doc, ["目录/文件", "职责", "运行位置"], [
    ("miniprogram/", "微信小程序界面、页面状态、按钮事件和API请求", "微信开发者工具/用户手机"),
    ("backend/app/", "FastAPI接口、业务规则、数据库模型、微信和AI适配", "Python服务器"),
    ("backend/folk_guide.db", "SQLite开发数据库", "开发电脑；生产要换PostgreSQL"),
    ("backend/storage/knowledge/", "上传的知识原文件", "开发电脑；生产要换对象存储"),
    ("preview/index.html", "不用微信工具也能看的浏览器静态演示", "本机8088端口"),
    ("docs/", "方案、学习手册、实施与审查报告", "项目资料目录"),
    ("tools/", "生成Word和Excel等开发辅助脚本", "开发电脑"),
])
doc.add_paragraph("关键原则：前端永远不能拥有AppSecret、数据库密码和真实管理员Key；前端传来的user_id、手机号或价格也不能直接相信。")

doc.add_heading("第2章 程序如何启动", level=1)
doc.add_heading("2.1 Python后端", level=2)
code(doc, "cd D:\\soft\\folk-guide\\backend\n.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8000")
table(doc, ["参数", "含义"], [
    ("-m uvicorn", "让Python以模块方式运行ASGI服务器uvicorn"),
    ("app.main:app", "导入app包中的main.py，再取其中名为app的FastAPI对象"),
    ("--host 127.0.0.1", "只监听本机；局域网和公网不能访问"),
    ("--port 8000", "HTTP端口；接口文档为http://127.0.0.1:8000/docs"),
])
doc.add_heading("2.2 浏览器预览", level=2)
code(doc, "python -m http.server 8088 --bind 127.0.0.1 --directory D:\\soft\\folk-guide\\preview")
doc.add_paragraph("浏览器地址为 http://127.0.0.1:8088/index.html。它只是静态交互演示，不能等价替代微信登录、手机号授权、订阅消息和支付测试。")
doc.add_heading("2.3 微信开发者工具", level=2)
steps(doc, ["导入 D:\\soft\\folk-guide。", "project.config.json读取miniprogramRoot和AppID。", "工具编译app.json列出的页面。", "开发工具请求本机API；真机必须改HTTPS域名，因为手机中的127.0.0.1是手机自己。"])

doc.add_heading("第3章 配置参数逐项解释", level=1)
table(doc, ["参数", "当前作用", "正式环境要求"], [
    ("APP_NAME", "FastAPI文档标题", "可改品牌正式名称"),
    ("ENVIRONMENT", "development允许模拟微信和开发CORS", "生产设production并关闭模拟分支"),
    ("DATABASE_URL", "SQLAlchemy连接字符串；默认sqlite:///./folk_guide.db", "改PostgreSQL连接并安全保管密码"),
    ("JWT_SECRET", "HS256签名密钥，防止伪造JWT", "强随机值；不能用默认值"),
    ("JWT_EXPIRE_MINUTES", "JWT有效分钟数；当前10080即7天", "结合风险设计访问/刷新令牌"),
    ("WECHAT_APP_ID", "小程序公开身份，已确认为wx349bf6810a9a559a", "前后端AppID必须一致"),
    ("WECHAT_APP_SECRET", "后端调用微信接口的秘密", "只在服务器.env/密钥管理，绝不进前端"),
    ("LLM_API_KEY", "未来大模型服务密钥", "供应商确定后填写"),
    ("LLM_BASE_URL", "未来大模型API根地址", "供应商确定后填写"),
    ("LLM_MODEL", "未来模型名称", "记录版本并做成本统计"),
    ("ADMIN_API_KEY", "原型运营后台共享Key", "上线前换管理员账号与权限系统"),
    ("API_BASE", "小程序请求后端的基础地址，当前127.0.0.1:8000", "真机/上线改备案HTTPS域名"),
])
doc.add_paragraph("Settings.model_config中的env_file='.env'表示Pydantic会从backend/.env读取配置；extra='ignore'表示.env出现未定义字段时忽略。get_settings使用lru_cache，避免每个请求重复读文件。")

doc.add_heading("第4章 Python后端分层", level=1)
table(doc, ["层", "主要文件", "职责"], [
    ("路由层", "main.py、*_routes.py", "接HTTP参数、做身份依赖、把业务错误转状态码"),
    ("Schema层", "schemas.py、*_schemas.py", "Pydantic校验类型、长度、枚举和返回结构"),
    ("服务层", "services.py、*_service.py", "真正业务规则、状态流转和第三方调用"),
    ("模型层", "models.py", "SQLAlchemy表、字段、外键、索引和唯一约束"),
    ("基础设施", "database.py、config.py、auth.py", "连接、配置、JWT、会话"),
    ("适配器", "vector_store.py、llm_provider.py", "隔离向量库和模型供应商"),
])
doc.add_paragraph("FastAPI的Depends会自动执行依赖。例如user: User = Depends(current_user)表示在业务函数运行前先验证Bearer JWT并加载User；db: Session = Depends(get_db)表示为请求创建数据库会话，请求结束自动关闭。")

doc.add_heading("第5章 数据库10张表与每个字段", level=1)
tables = {
"users（账号）": [
    ("id", "内部用户主键；所有个人业务表用它关联"), ("openid", "微信用户在本小程序内的唯一标识，不返回前端"),
    ("phone_number", "微信手机号接口换取的号码，可为空"), ("phone_bound_at", "手机号绑定时间"), ("created_at", "账号创建时间，也是新客3天起点"),
],
"birth_profiles（生辰档案）": [
    ("id", "档案主键"), ("user_id", "用户外键且unique，因此一人最多一份"), ("calendar_type", "solar公历或lunar农历"),
    ("is_leap_month", "农历是否闰月；公历会强制False"), ("birth_date", "用户输入日期；含义由calendar_type决定"),
    ("time_known", "是否知道时辰"), ("birth_time", "HH:MM；未知则None"), ("birth_city", "可选出生城市"),
    ("gender", "male/female/unspecified"), ("timezone", "当前仅Asia/Shanghai"), ("profile_version", "每次修改递增"),
    ("normalized_solar_date", "转换后的标准公历日期"), ("calendar_data_json", "生肖、节气、四柱等缓存JSON"),
    ("calculation_version", "历法算法版本，便于将来失效重算"), ("created_at/updated_at", "创建和更新时间"),
],
"daily_guidance（个人每日缓存）": [("id", "主键"), ("user_id", "所属用户"), ("guidance_date", "建议日期"), ("payload_json", "颜色、宜事、提醒JSON"), ("created_at", "生成时间")],
"chat_messages（旧版问答）": [("id", "主键"), ("user_id", "所属用户"), ("question", "问题"), ("answer", "旧规则答案"), ("created_at", "时间；正式接口稳定后应废弃本表")],
"ai_service_grants（AI权益）": [("id", "主键"), ("user_id", "所属用户"), ("grant_type", "权益类型，如paid_30_days"), ("start_at/end_at", "有效区间"), ("created_at", "写入时间")],
"ai_conversation_messages（正式AI历史）": [
    ("id", "消息主键并返回前端用于反馈"), ("user_id", "数据隔离条件"), ("question/answer", "问题与答案"),
    ("category", "culture_knowledge/profile_culture/seven_day_comparison/high_risk"), ("references_json", "回答时引用快照"),
    ("model_name", "生成答案的模型或安全规则名称"), ("safety_status", "safe或blocked"),
    ("feedback", "helpful/unhelpful"), ("feedback_note", "反馈说明"), ("created_at", "用于历史排序和每日计数"),
],
"public_guides（今日五色）": [("id", "主键"), ("guide_date", "日期唯一"), ("status", "draft/reviewing/scheduled/published/withdrawn"), ("version", "修改版本"), ("payload_json", "完整五色内容"), ("scheduled_at", "计划发布时间"), ("published_at", "实际发布时间"), ("created_at/updated_at", "审计时间")],
"public_guide_audits（五色审计）": [("id", "主键"), ("guide_id", "对应每日内容"), ("action", "create/update/review/publish等"), ("operator", "操作者"), ("before_status/after_status", "状态前后"), ("detail_json", "附加信息"), ("created_at", "操作时间")],
"knowledge_documents（知识文档）": [
    ("id", "文档主键"), ("title/source_name/author", "展示标题、来源和作者"), ("copyright_status", "public_domain/licensed/original/permission_pending"),
    ("original_filename/storage_key", "原名和服务端存储位置"), ("file_hash", "SHA-256去重"), ("file_size/mime_type", "体积和类型"),
    ("status", "uploaded/reviewing/approved/failed/disabled"), ("version", "重新解析/修改版本"),
    ("error_message", "失败原因"), ("reviewed_by/reviewed_at", "审核人和时间"), ("created_at/updated_at", "时间"),
],
"knowledge_chunks（知识片段）": [
    ("id", "片段主键"), ("document_id", "所属文档"), ("chunk_index", "文档内顺序且与document_id联合唯一"),
    ("heading", "章节标题"), ("content", "问答检索正文"), ("content_hash", "内容指纹"), ("character_count", "字符数"),
    ("review_status", "pending/approved/rejected/disabled"), ("vector_id", "外部向量库记录ID"),
    ("embedding_model", "向量模型版本"), ("created_at/updated_at", "时间"),
]}
for name, fields in tables.items():
    doc.add_heading(name, level=2)
    table(doc, ["字段", "作用"], fields)
doc.add_paragraph("关键约束：DailyGuidance的(user_id,guidance_date)唯一；KnowledgeChunk的(document_id,chunk_index)唯一；openid、guide_date、file_hash唯一。这些约束在并发或重复提交时守住数据一致性。")

doc.add_heading("第6章 JWT与身份链路", level=1)
steps(doc, [
    "前端ensureLogin检查app.globalData.token；有缓存时先复用。",
    "没有token时调用wx.login()取得短期code。",
    "POST /api/auth/wechat，LoginRequest限制code长度1—256。",
    "exchange_wechat_code在真实配置下调用微信jscode2session，得到openid；开发未配AppID时生成模拟openid。",
    "数据库按openid查询users；没有则创建。",
    "create_token把user.id写入JWT的sub，并写exp过期时间，用JWT_SECRET做HS256签名。",
    "前端把access_token写入内存和wx.setStorageSync('token')。",
    "后续request自动加Authorization: Bearer <JWT>。",
    "current_user验证签名、算法和有效期，从sub取user_id并加载User。失败返回401，前端clearLogin。",
])
doc.add_paragraph("参数解释：credentials是HTTPBearer解析的令牌；db是请求数据库会话；sub是JWT标准主体字段；exp是到期时间。JWT不存生辰、手机号和AppSecret。当前不足：无主动撤销和Refresh Token。")

doc.add_heading("第7章 手机号绑定链路", level=1)
steps(doc, [
    "用户在‘我的’点击open-type=getPhoneNumber按钮。",
    "微信回调bindPhone(event)，event.detail.code是一性授权code，不是手机号。",
    "前端POST /api/users/me/phone并自动携带JWT。",
    "PhoneCodeInput限制code长度1—512。",
    "exchange_wechat_phone_code先用AppID/AppSecret取得access_token，再调用getuserphonenumber。",
    "后端读取phone_info.phoneNumber，保存users.phone_number和phone_bound_at。",
    "返回CurrentUserOutput；前端只展示前三位****后四位。",
])
doc.add_paragraph("为什么不接受前端直接传手机号：任何人都能伪造明文手机号。当前不足：微信access_token没有共享缓存，手机号唯一/换绑规则尚未确定。")

doc.add_heading("第8章 档案和公农历链路", level=1)
table(doc, ["BirthProfileInput参数", "规则"], [
    ("calendar_type", "solar或lunar；默认solar"), ("is_leap_month", "仅农历有效"), ("birth_date", "必填日期"),
    ("time_known", "是否知道时辰"), ("birth_time", "00:00—23:59；填写后自动视为已知"),
    ("birth_city", "最多64字，可空"), ("gender", "male/female/unspecified"), ("timezone", "当前仅Asia/Shanghai"),
])
steps(doc, [
    "profile页面onLoad先ensureLogin，再GET /api/profiles/current。404表示尚无档案。",
    "用户修改页面data；updateLocalCompleteness只做界面预览。",
    "save检查日期与隐私勾选，再PUT /api/profiles/current。",
    "Pydantic normalize_profile消除矛盾：公历清闰月；未知时辰清birth_time；城市去空白。",
    "save_profile新增或更新BirthProfile，修改时profile_version加一。",
    "calculate_birth_calendar调用lunar-python进行公农历、生肖、星座、节气和四柱计算。",
    "未知时辰时不生成时柱；结果写入calendar_data_json并标calculation_version。",
    "档案变化删除DailyGuidance旧缓存，避免继续使用旧资料。",
    "返回BirthProfileOutput，包含完整度、缺失项、full/simplified和calendar。",
])
doc.add_paragraph("删除链路：确认弹窗 → DELETE /api/profiles/current → 删除档案与每日缓存 → 返回204。档案删除不等于注销账号。")

doc.add_heading("第9章 个人每日建议链路", level=1)
steps(doc, [
    "GET /api/daily要求JWT和已存在BirthProfile。",
    "用user_id + date.today查询DailyGuidance缓存。",
    "命中则loads_payload直接返回，避免同日重复计算。",
    "未命中则build_daily_guidance以出生日期和当天生成稳定演示结果。",
    "dumps_payload写入JSON，保存后返回DailyGuidanceOutput。",
])
doc.add_paragraph("当前真实状态：这是确定性演示规则，不是正式个人五色算法；正式规则要版本化、审核、回溯，并明确文化口径。")

doc.add_heading("第10章 今日五色完整运营链路", level=1)
doc.add_heading("10.1 数据参数", level=2)
table(doc, ["参数", "作用/限制"], [
    ("rank", "1—5且一天完整不重复"), ("color", "白金、绿金、黑金、红金、黄金完整出现"),
    ("element", "与颜色固定对应"), ("smoothness", "五档描述枚举"), ("suitable", "至少一项"),
    ("resistance/advice", "各1—500字"), ("product_code", "香品内部代码，最多64"),
    ("incense_name/scent", "商品名最多100，香气最多300"), ("guide_date", "每天唯一"),
    ("weekday/lunar_date/solar_term/day_ganzhi", "首页日期文化信息"),
    ("share_title/share_summary/push_summary", "分享和推送文案"), ("rule_version", "内容生成或人工规则版本"),
])
doc.add_heading("10.2 运营状态", level=2)
code(doc, "draft → reviewing → scheduled → published → withdrawn")
steps(doc, [
    "运营页面或Excel把内容POST/import到后台，require_admin验证X-Admin-Key。",
    "Pydantic检查五色、五行和排名完整性；save_draft写payload_json并生成审计。",
    "review提交审核；schedule写scheduled_at或publish立即发布。",
    "每次状态变化写PublicGuideAudit，保存operator和前后状态。",
    "公开GET /api/public-guides/today只返回published；前端home.onShow映射成guides。",
    "请求失败时首页保留内置demo，contentSource显示演示内容。",
])
doc.add_paragraph("当前不足：到期发布由读取请求触发，不是真正定时器；生产应使用调度任务和数据库锁。")

doc.add_heading("第11章 知识文档入库链路", level=1)
table(doc, ["上传参数", "作用"], [
    ("file", "支持.docx/.md/.txt，最大10MB"), ("title", "运营展示标题"), ("source_name", "资料来源"),
    ("copyright_status", "版权状态枚举"), ("author", "作者，可空"), ("X-Admin-Key", "管理员认证"),
])
steps(doc, [
    "POST /api/admin/knowledge/documents：safe_filename移除路径和危险字符。",
    "读取文件，校验扩展名、空文件、10MB上限和版权状态。",
    "sha256_bytes计算指纹；file_hash已存在则409拒绝重复。",
    "原文件写storage/knowledge/YYYY/MM；元数据写KnowledgeDocument，status=uploaded。",
    "POST /{id}/parse按扩展名调用parse_docx/parse_markdown/parse_text。",
    "normalize_text统一换行和多余空白；ParsedBlock保留heading/paragraph/table结构。",
    "chunk_blocks以章节和中文句末为边界，目标约700字、尽量不低于250字。",
    "replace_chunks删除旧片段并写新片段，全部review_status=pending；文档进入reviewing。",
    "运营查看chunks，可PUT修订；修改后重新pending并清vector_id。",
    "POST /approve全部或选择片段通过，文档变approved；未通过片段rejected。",
])
doc.add_paragraph("关键边界：上传≠可检索，解析≠可检索，只有文档approved且片段approved才可能进入问答。用户聊天和生辰档案永远不进入公共知识库。")

doc.add_heading("第12章 向量同步与混合检索", level=1)
doc.add_heading("12.1 VectorStore参数", level=2)
table(doc, ["方法/参数", "作用"], [
    ("upsert(records)", "新增或覆盖VectorRecord；记录含id、text、metadata"),
    ("query(text, limit)", "用自然语言检索，limit控制返回数量"),
    ("delete(ids)", "幂等删除外部向量"), ("health_check()", "快速健康检查"),
    ("vector_id", "远端记录主键"), ("embedding_model", "生成向量的模型版本"),
])
steps(doc, [
    "审核后调用POST /sync-vectors；未approved返回409。",
    "sync_approved_chunks只查询approved片段，组装稳定ID knowledge:<chunk_id>:<hash前12位>。",
    "store.upsert写入；返回ID保存到KnowledgeChunk.vector_id和embedding_model。",
    "提问时hybrid_search同时执行keyword_search和store.query。",
    "用RRF倒数排名融合：每一路第rank名贡献1/(60+rank)，避免供应商分数尺度不同。",
    "最后再次查数据库，强制片段approved且文档approved，返回正文和来源。",
    "disable_document先删远端向量，再清映射并把文档/片段disabled。",
])
doc.add_paragraph("当前MemoryVectorStore使用字符二元组模拟相似度，重启即清空，不是真Embedding。正式版需要外部持久化向量库、任务队列、失败重试和远端对账。")

doc.add_heading("第13章 正式AI问答链路", level=1)
table(doc, ["接口参数/返回", "含义"], [
    ("question", "1—500字自然语言问题"), ("question_type", "normal或seven_day_comparison，可空"),
    ("message_id", "保存后的消息ID，用于反馈"), ("category", "问题类别"), ("blocked", "是否安全拦截"),
    ("citations", "document_id、chunk_id、title、heading、source_name"),
    ("remaining_today", "对应类别今日剩余次数"), ("model_name", "回答实现版本"), ("disclaimer", "固定风险提示"),
])
steps(doc, [
    "前端应先GET /api/ai/quota显示权益和剩余次数。",
    "POST /api/ai/chat进入chat_ai，再调用answer_ai_question。",
    "classify_question先执行高风险词硬拦截；七日问题分到seven_day_comparison；档案词分到profile_culture。",
    "高风险直接返回安全说明，不检索、不调用模型、不消耗次数。",
    "quota_for_user优先查有效AIServiceGrant；否则账号创建72小时内为new_user_3_days。",
    "新客普通20/日、七日2/日；正式30天普通50/日、七日5/日。",
    "次数不足返回429；权益过期返回403。",
    "hybrid_search取最多5段审核资料；AnswerProvider.generate生成答案。",
    "答案、分类、引用快照、模型名和安全状态写AIConversationMessage。",
    "返回答案、引用和剩余次数。history只按当前user_id查询；feedback也同时匹配message_id和user_id。",
])
doc.add_paragraph("当前LocalAnswerProvider只是把检索正文整理成开发预览。小程序chat/index.ts目前仍使用ANSWERS模拟表和setTimeout，尚未接真实/api/ai/*，这是下一步最重要工作。")

doc.add_heading("第14章 小程序前端逐文件解析", level=1)
doc.add_heading("14.1 app.json / app.ts / api.ts", level=2)
table(doc, ["参数/函数", "作用"], [
    ("pages", "页面注册顺序，第一项为启动页"), ("tabBar.list", "四个底部Tab路径和标题"),
    ("navigationBarBackgroundColor等", "全局导航视觉"), ("globalData.token", "当前内存JWT"),
    ("request<T>(path,method,data)", "泛型T描述返回类型；自动拼URL、加JWT、处理2xx/401"),
    ("ensureLogin()", "无token时wx.login并换JWT"), ("clearLogin()", "清内存和Storage"),
    ("getCurrentUser()", "验证JWT并取账号状态"), ("bindWechatPhone(code)", "提交手机号临时code"),
])
doc.add_heading("14.2 首页 home", level=2)
table(doc, ["data/方法", "作用"], [
    ("expandedRank", "当前展开的颜色排名；0表示全收起"), ("contentSource", "published或demo"),
    ("dateLabel/calendarLabel", "公历/农历/节气/干支展示"), ("shareTitle", "微信分享标题"),
    ("guides", "五色排名详情数组"), ("onShow", "每次进入调用公开今日接口"),
    ("toggleGuide(event)", "从dataset.rank读取排名并展开/收起"), ("toIncense", "switchTab到财库香"),
    ("onShareAppMessage", "返回分享title和path"),
])
doc.add_heading("14.3 AI国学 chat", level=2)
table(doc, ["data/方法", "作用/当前状态"], [
    ("serviceStatus", "guest/no_profile/trial/active/expired界面状态，当前写死trial"),
    ("trialDaysLeft/dailyLimit/remainingQuestions", "当前写死的演示权益"),
    ("comparisonLimit/remainingComparisons", "七日比较演示次数"), ("messages", "聊天气泡数组"),
    ("ANSWERS", "固定问题模拟答案；应删除并改真实接口"), ("sendQuestion", "当前本地查ANSWERS、350ms后插入回答并减次数"),
    ("feedback", "当前只改本地messages反馈；应调用后端feedback接口"),
])
doc.add_heading("14.4 档案 profile", level=2)
table(doc, ["data/方法", "作用"], [
    ("loading/saving/deleting", "防重复操作和界面加载状态"), ("hasProfile", "是否已有数据库档案"),
    ("calendarType/isLeapMonth", "公历农历与闰月"), ("timeKnown/birthTime", "时辰状态"),
    ("privacyAgreed", "保存前隐私说明勾选"), ("onLoad", "登录、读取档案、404视为新建"),
    ("updateLocalCompleteness", "50+时辰25+城市15+性别10的界面预览"),
    ("save", "组装snake_case JSON并PUT"), ("deleteProfile", "二次确认并DELETE"),
])
doc.add_heading("14.5 财库香 caikuxiang", level=2)
table(doc, ["data/方法", "作用/当前状态"], [
    ("PRODUCTS", "5个单色产品静态数组"), ("cart", "本地购物车数组"),
    ("selectedProduct", "商品详情弹层选中项"), ("addItem", "已有商品加quantity，否则新增"),
    ("changeQuantity", "根据dataset.id/delta增减，0时移除"),
    ("cartCount/cartTotal", "由购物车计算的展示值"), ("checkout", "当前只提示尚未接订单支付"),
])
doc.add_heading("14.6 我的 settings", level=2)
table(doc, ["data/方法", "作用/当前状态"], [
    ("isLoggedIn", "初始false，onShow用JWT验证"), ("phoneBound/phoneDisplay", "手机号状态和脱敏展示"),
    ("wechatPhoneAvailable", "后端是否有完整微信配置"), ("realLogin", "真实wx.login链路入口"),
    ("applyAccount", "把CurrentUser映射到页面data"), ("bindPhone", "处理getPhoneNumber回调"),
    ("logout", "只清本地JWT，不注销账号"),
    ("订单/卡包/提醒相关data和方法", "目前多为前端演示，尚无订单与服务卡后端"),
])
doc.add_paragraph("WXML负责结构和数据绑定：{{字段}}显示data；wx:if条件渲染；wx:for循环；bindtap绑定事件；data-xxx进入event.currentTarget.dataset。WXSS负责样式，rpx会随屏幕宽度缩放，750rpx约等于设计稿全宽。")

doc.add_heading("第15章 HTTP参数、状态码和关键术语", level=1)
table(doc, ["术语/状态", "解释"], [
    ("GET", "读取；原则上不改变业务数据"), ("POST", "创建或执行动作"), ("PUT", "完整更新某资源"), ("DELETE", "删除"),
    ("200", "成功并有JSON"), ("204", "成功但无响应正文"), ("401", "JWT无效/过期"), ("403", "身份有效但无权限或权益"),
    ("404", "资源不存在或为防越权而隐藏"), ("409", "状态冲突/重复"), ("413", "文件过大"), ("422", "参数或业务校验失败"),
    ("429", "次数用完/限流"), ("503", "依赖能力未配置或暂不可用"),
    ("Query", "URL ?q=...参数"), ("Path", "URL路径中的{id}"), ("Header", "Authorization、X-Admin-Key等"),
    ("Form/File", "multipart上传字段"), ("JSON body", "普通接口请求体"), ("response_model", "返回数据校验和文档模型"),
    ("Session", "SQLAlchemy数据库会话/事务单元"), ("select", "构建查询"), ("scalar/scalars", "取单值/对象序列"),
    ("commit", "提交事务"), ("refresh", "从数据库刷新自动生成字段"), ("index", "加速查询"), ("unique", "拒绝重复值"),
])

doc.add_heading("第16章 当前哪些链路已真实、哪些仍是演示", level=1)
table(doc, ["功能", "当前程度", "尚缺"], [
    ("微信登录/JWT", "后端和前端链路已有，已配置AppID/AppSecret后可真实调用", "开发者工具真实联调、令牌撤销/刷新、HTTPS"),
    ("手机号", "后端真实微信接口与前端按钮骨架已有", "微信能力验证、token缓存、换绑规则"),
    ("档案/历法", "数据库和接口真实可用", "历法产品口径专业复核、隐私授权版本记录"),
    ("今日五色", "后台、Excel、发布和首页公开接口真实可用", "正式内容生成规则和独立调度"),
    ("知识库", "上传、解析、切片、审核、关键词检索真实可用", "PDF/OCR、对象存储、全文检索"),
    ("向量检索", "接口和混合检索框架可用", "当前内存模拟，缺外部向量库与Embedding"),
    ("AI问答后端", "权益、安全、历史、引用接口真实可用", "当前本地答案Provider，缺正式模型"),
    ("AI前端", "高级感交互原型", "尚未连接/api/ai/*，仍是ANSWERS模拟"),
    ("财库香商城", "商品和购物车前端原型", "商品库、库存、订单、物流、售后、支付"),
    ("服务卡/提醒", "前端原型和AI权益表部分基础", "发卡、激活、赠送、模板消息后端"),
])

doc.add_heading("第17章 一次完整请求如何排查", level=1)
doc.add_paragraph("以“保存档案失败”为例：")
steps(doc, [
    "看profile/index.ts save是否通过日期和privacyAgreed检查。",
    "看request是否使用PUT /api/profiles/current并携带Authorization。",
    "浏览器/开发工具网络面板看HTTP状态和detail。",
    "401查JWT；422查BirthProfileInput和历法ValueError；500查服务器日志。",
    "后端main.save_profile是否进入；calculate_birth_calendar是否成功。",
    "查看birth_profiles是否写入，calendar_data_json/calculation_version是否存在。",
    "确认DailyGuidance旧缓存是否删除，返回BirthProfileOutput是否通过。",
])
doc.add_paragraph("同样方法适用于所有链路：先页面事件，再request，再路由Depends，再Schema，再服务函数，再数据库和外部接口，最后返回状态码。")

doc.add_heading("第18章 测试体系", level=1)
table(doc, ["测试文件", "覆盖"], [
    ("test_health.py", "启动和健康接口"), ("test_api_flow.py", "登录、档案、每日、旧问答、注销主流程"),
    ("test_calendar_service.py", "公农历和历法接口"), ("test_profile.py", "未知时辰、版本、非法日期"),
    ("test_public_guides.py", "五色状态机、Excel、定时发布"), ("test_knowledge.py", "文档解析、审核、检索、停用"),
    ("test_vector_retrieval.py", "未审核不能向量化、混合检索、停用退出"),
    ("test_ai_chat.py", "权益、安全、历史、反馈隔离"), ("test_user_account.py", "账号状态、手机号边界、注销清AI数据"),
    ("conftest.py", "测试强制独立数据库和模拟微信，不读取真实密钥"),
])
code(doc, "cd backend\n$env:PYTHONDONTWRITEBYTECODE='1'\n$env:PYTEST_DISABLE_PLUGIN_AUTOLOAD='1'\n.\\.venv\\Scripts\\python.exe -m pytest -q -p no:cacheprovider")
doc.add_paragraph("当前结果：23 passed。警告主要是datetime.utcnow未来弃用，已列入时区统一改造。")

doc.add_heading("第19章 推荐学习与开发顺序", level=1)
steps(doc, [
    "先读database.py、models.py，理解数据存在哪里。",
    "读schemas.py，理解接口允许什么输入、返回什么。",
    "读auth.py、services.exchange_wechat_code和main.wechat_login，掌握登录。",
    "读profile页面和main.save_profile，完整追一次前后端链路。",
    "读public_guide的schema → service → routes，理解运营状态机。",
    "读knowledge_service → routes → retrieval_service，理解RAG资料准备。",
    "读ai_service → ai_routes → llm_provider，理解问答编排。",
    "下一步实际开发：把chat/index.ts从ANSWERS改成ensureLogin + quota/chat/history/feedback。",
    "之后选择外部向量库、Embedding和大模型，再实现生产适配器。",
])

doc.add_heading("附录A 后端接口总表", level=1)
api_rows = [
    ("GET", "/health", "无", "健康检查"), ("POST", "/api/auth/wechat", "code", "微信登录/JWT"),
    ("GET", "/api/users/me", "JWT", "当前账号"), ("POST", "/api/users/me/phone", "JWT+phone code", "绑手机号"),
    ("GET/PUT/DELETE", "/api/profiles/current", "JWT", "读取/保存/删档案"), ("POST", "/api/calendar/convert", "JWT+档案参数", "历法预览"),
    ("GET", "/api/daily", "JWT", "个人每日建议"), ("POST", "/api/chat", "JWT+question", "旧版问答"),
    ("DELETE", "/api/account", "JWT", "注销账号"),
    ("GET", "/api/public-guides/today", "无", "今日公开五色"), ("GET", "/api/public-guides/{date}", "无", "指定日公开五色"),
    ("多接口", "/api/admin/public-guides/*", "Admin Key", "五色运营、导入、发布和审计"),
    ("多接口", "/api/admin/knowledge/*", "Admin Key", "知识上传、解析、审核、向量和检索"),
    ("GET", "/api/ai/quota", "JWT", "AI权益次数"), ("POST", "/api/ai/chat", "JWT+问题", "正式AI问答"),
    ("GET", "/api/ai/history", "JWT+limit", "本人AI历史"), ("PUT", "/api/ai/messages/{id}/feedback", "JWT+rating/note", "本人反馈"),
]
table(doc, ["方法", "路径", "认证/参数", "用途"], api_rows)

doc.add_heading("附录B 后端函数参数总索引", level=1)
function_rows = [
    ("admin_auth", "require_admin(x_admin_key,x_admin_name)", "两个值来自HTTP请求头；Key用于认证，Name用于审计", "操作者名称；失败403"),
    ("auth", "create_token(user_id)", "user_id为数据库内部用户主键", "签名JWT字符串"),
    ("auth", "current_user(credentials,db)", "credentials为Bearer令牌；db为请求会话", "当前User；失败401"),
    ("database", "get_db()", "无；FastAPI依赖函数", "yield Session并在请求后关闭"),
    ("calendar", "_split_pillar(value)", "value为甲子等两字柱", "干支字典"),
    ("calendar", "_jie_qi_payload(jie_qi)", "lunar-python节气对象", "节气名称和时间字典"),
    ("calendar", "calculate_birth_calendar(data)", "BirthProfileInput", "完整历法字典；非法农历抛ValueError"),
    ("calendar", "apply_calendar_calculation(profile,data)", "ORM档案对象+输入模型", "计算结果并写档案缓存"),
    ("calendar", "load_calendar_payload(profile)", "ORM档案对象", "解析后的缓存字典"),
    ("profile", "profile_completeness(profile)", "ORM档案", "(百分比,缺失字段列表)"),
    ("profile", "profile_output(profile)", "ORM档案", "API返回字典"),
    ("services", "exchange_wechat_code(code)", "wx.login短期code", "openid；微信失败抛HTTPException"),
    ("services", "exchange_wechat_phone_code(code)", "getPhoneNumber一次性code", "手机号；未配置503"),
    ("services", "build_daily_guidance(profile,today)", "档案+date日期", "每日建议字典"),
    ("services", "answer_question(question)", "旧版问题字符串", "(答案,分类,来源列表)"),
    ("services", "dumps_payload(payload)", "Python字典", "保留中文的JSON字符串"),
    ("services", "loads_payload(payload)", "JSON字符串", "Python字典"),
    ("main", "lifespan(app)", "FastAPI传入应用；函数不使用故命名_", "启动建表/迁移，上下文结束关闭"),
    ("main", "wechat_login(data,db)", "LoginRequest+数据库会话", "TokenResponse"),
    ("main", "current_account(user)", "Depends注入当前User", "不含openid的账号状态"),
    ("main", "bind_phone(data,user,db)", "PhoneCodeInput+当前用户+会话", "更新后的账号状态"),
    ("main", "get_profile(user,db)", "当前用户+会话", "本人档案；无档案404"),
    ("main", "convert_calendar(data,_)", "档案输入+仅用于鉴权的User", "历法预览，不保存"),
    ("main", "save_profile(data,user,db)", "档案输入+当前用户+会话", "保存后的档案"),
    ("main", "delete_profile(user,db)", "当前用户+会话", "204"),
    ("main", "get_daily(user,db)", "当前用户+会话", "缓存或新生成的每日建议"),
    ("main", "chat(data,user,db)", "旧ChatRequest+当前用户+会话", "旧ChatResponse"),
    ("main", "delete_account(user,db)", "当前用户+会话", "清个人数据后204"),
    ("knowledge", "sha256_bytes(content)", "文件bytes", "64位SHA-256十六进制"),
    ("knowledge", "normalize_text(value)", "任意文本", "统一空白后的文本"),
    ("knowledge", "parse_docx/markdown/text(path)", "服务端Path", "ParsedBlock列表"),
    ("knowledge", "parse_file(path)", "带扩展名Path", "分派解析器结果"),
    ("knowledge", "split_long_text(text,target=700)", "文本+目标字符数", "较短文本片段列表"),
    ("knowledge", "chunk_blocks(blocks)", "ParsedBlock列表", "(heading,content)列表"),
    ("knowledge", "replace_chunks(db,document,chunks)", "会话+文档+切片列表", "无；等待调用者commit"),
    ("knowledge", "document_output(db,document)", "会话+文档", "含切片统计的返回字典"),
    ("knowledge", "keyword_search(db,query,limit=8)", "会话+关键词+上限", "按命中分数排序的片段"),
    ("knowledge_routes", "safe_filename(filename)", "客户端文件名", "无目录的安全文件名"),
    ("knowledge_routes", "upload_document(file,title,source_name,copyright_status,author,_,db)", "multipart字段；_是只用于认证的管理员名", "KnowledgeDocumentOutput"),
    ("knowledge_routes", "list_documents(status,_,db)", "可选状态筛选+管理员+会话", "文档列表"),
    ("knowledge_routes", "parse/list/update/approve/disable(document_id或chunk_id,...)", "路径ID、审核/修改模型、管理员、会话", "相应文档或片段"),
    ("knowledge_routes", "search_knowledge(q,limit,_,db)", "Query问题、1—20上限、管理员、会话", "关键词结果"),
    ("knowledge_routes", "sync_document_vectors(document_id,_,db)", "文档ID+管理员+会话", "同步数量和模型"),
    ("retrieval", "sync_approved_chunks(db,document,store)", "会话+approved文档+VectorStore", "同步条数"),
    ("retrieval", "remove_document_vectors(db,document_id,store)", "会话+文档ID+向量库", "无；清远端和映射"),
    ("retrieval", "hybrid_search(db,query,store,limit=8)", "会话+问题+向量库+上限", "RRF融合结果"),
    ("vector", "_tokens(text)", "文本", "字符二元组集合"),
    ("vector", "VectorRecord(id,text,metadata)", "外部ID、正文、过滤元数据", "不可变记录对象"),
    ("vector", "VectorMatch(id,score,metadata)", "命中ID、相似度、元数据", "不可变命中对象"),
    ("vector", "set_vector_store(store)", "实现VectorStore协议的实例", "替换共享适配器"),
    ("ai_service", "classify_question(question,requested_type=None)", "问题+前端显式类型", "(类别,是否拦截)"),
    ("ai_service", "quota_for_user(db,user,now=None)", "会话+用户+可注入测试时间", "QuotaState"),
    ("ai_service", "QuotaState.remaining(category)", "问题类别", "该类剩余次数"),
    ("ai_service", "citations_from_contexts(contexts)", "混合检索结果", "去除内部字段后的引用列表"),
    ("ai_service", "answer_ai_question(db,user,question,requested_type,store,provider)", "会话、用户、问题、类型、向量库、模型", "(消息,提问前QuotaState,是否拦截)"),
    ("ai_routes", "quota_output(quota)", "QuotaState", "AIQuotaOutput字典"),
    ("ai_routes", "chat_ai(data,user,db)", "AIChatInput+当前用户+会话", "AIChatOutput；403/429"),
    ("ai_routes", "history(limit,user,db)", "1—100条上限+当前用户+会话", "本人历史列表"),
    ("ai_routes", "feedback(message_id,data,user,db)", "消息ID+AIFeedbackInput+用户+会话", "更新后的本人历史项"),
    ("public_service", "payload_dict/guide_output(guide)", "PublicGuide ORM对象", "内容字典/完整返回字典"),
    ("public_service", "audit(db,guide,action,operator,before,detail=None)", "会话、内容、动作、人、原状态、可选详情", "添加审计对象，暂不commit"),
    ("public_service", "save_draft(db,data,operator)", "会话+PublicGuideInput+操作者", "(guide,是否新建)"),
    ("public_service", "set_status(db,guide,target,operator,action,scheduled_at=None)", "目标状态及审计信息", "更新后的guide"),
    ("public_service", "publish_due_guides(db)", "会话", "无；发布到期内容"),
    ("public_service", "build_excel_template()", "无", "xlsx文件bytes"),
    ("public_service", "parse_excel(content)", "Excel bytes", "PublicGuideInput列表"),
]
table(doc, ["模块", "函数签名", "参数来源/作用", "返回/副作用"], function_rows)

doc.add_heading("附录C 前端事件参数总索引", level=1)
table(doc, ["参数/对象", "来自哪里", "如何使用"], [
    ("event.currentTarget.dataset", "WXML元素的data-*属性", "读取rank、question、id、delta、gender等业务参数"),
    ("event.detail.value", "input/switch/picker/checkbox组件", "文本、布尔值、选择下标或选中数组"),
    ("event.detail.code", "getPhoneNumber按钮回调", "一次性code，交给后端换手机号"),
    ("path", "request第1参数", "相对API路径，例如/api/users/me"),
    ("method", "request第2参数", "GET/POST/PUT/DELETE，默认GET"),
    ("data", "request第3参数", "发送的JSON对象，可选"),
    ("T", "request<T>泛型", "仅TypeScript编译期描述返回类型，不会发送到服务器"),
    ("this.data", "Page页面响应式状态", "WXML通过{{}}读取；必须用setData更新界面"),
    ("loading/saving/sending/deleting", "页面布尔状态", "禁重复提交并显示进度"),
    ("wx:key", "WXML循环参数", "帮助框架稳定识别列表项，减少重复渲染"),
    ("wx:if/wx:else", "WXML条件", "按登录、展开、弹层等状态创建/销毁结构"),
    ("bindtap/catchtap", "事件绑定", "bind允许冒泡；catch阻止继续冒泡"),
    ("open-type", "微信按钮开放能力", "share分享；getPhoneNumber申请手机号授权"),
])

for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = "五行五色财库香｜全项目代码参数与链路详解 V2.0"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.text = "学习手册｜2026-08-13"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
