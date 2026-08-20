from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "今日五色_每日更新功能方案_V1.0.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], "385F4D")
        if widths:
            table.rows[0].cells[index].width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Cm(widths[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F2F5F3")
    document.add_paragraph()
    return table


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_steps(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_code_block(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3EF")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(text.splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(56, 95, 77)
    document.add_paragraph()


def build_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size, color in [
        ("Title", 24, "27322D"),
        ("Heading 1", 16, "385F4D"),
        ("Heading 2", 12.5, "4D6659"),
    ]:
        style = styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("今日五色｜每日更新功能方案")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("五行五色财库香项目 · 产品与技术讨论稿 V1.0")
    run.font.color.rgb = RGBColor(113, 128, 120)
    run.font.size = Pt(11)
    document.add_paragraph()

    summary = document.add_table(rows=1, cols=1)
    summary_cell = summary.cell(0, 0)
    set_cell_shading(summary_cell, "E9F0EB")
    p = summary_cell.paragraphs[0]
    p.add_run("核心原则：").bold = True
    p.add_run("规则引擎计算排名，内容模板生成建议，人工审核后定时发布；AI只负责受约束的语言解释，不临时决定公共五色排名。")
    document.add_paragraph()

    document.add_heading("01｜功能目标", level=1)
    document.add_paragraph("“今日五色”是项目的免费日活与自然分享入口。每日更新功能需要保证用户每天打开小程序时，都能快速、稳定地看到当天公共五色排名及可执行的生活建议，并能自然进入分享、登录和香品浏览路径。")
    add_bullets(document, [
        "形成每日回访理由，提高次日和7日留存。",
        "提供无需登录即可完整查看的公共五色内容。",
        "让每日一份数据可复用于小程序、分享卡、朋友圈、短视频和公众号。",
        "用对应香品完成自然商品连接，但不承诺确定性结果。",
        "为未来个人五色服务提供公共环境基准。",
    ])

    document.add_heading("02｜推荐的每日更新流程", level=1)
    add_code_block(document, "每天日期变化\n    ↓\n历法程序计算当天信息\n    ↓\n五行规则计算五色排名\n    ↓\n内容模板生成详细建议\n    ↓\n运营后台人工审核\n    ↓\n次日 00:00 自动发布\n    ↓\n用户打开首页读取当天结果\n    ↓\n在微信允许范围内发送订阅提醒")
    document.add_paragraph("公共排名必须可追溯。系统需要记录计算规则版本、编辑记录、审核人和发布时间，以便出现争议或错误时快速定位。")

    document.add_heading("03｜每天需要生成的数据", level=1)
    add_table(document, ["数据层级", "内容"], [
        ["基础信息", "日期、星期、农历、节气、当日干支"],
        ["五色排名", "绿金、黑金、黄金、白金、红金的第1至第5明确排序"],
        ["顺畅度", "今天很顺、比较合适、平稳一般、会比较累、成效偏弱"],
        ["行动内容", "适合事项、可能的阻力、行动调整建议"],
        ["商品连接", "每种颜色对应的固定财库香商品编号"],
        ["分发内容", "分享标题、分享摘要、分享图、推送摘要"],
        ["追溯信息", "规则版本、内容版本、审核状态、发布时间"],
    ], [3.2, 12.5])

    document.add_heading("04｜固定关系与每日变化", level=1)
    add_table(document, ["永久固定", "每天变化"], [
        ["白金＝金；绿金＝木；黑金＝水；红金＝火；黄金＝土", "五种颜色从第1到第5的排序"],
        ["每种颜色对应的香品、配方和商品编号", "顺畅度、适合事项、阻力和行动建议"],
        ["公共内容的表达边界与免责声明", "分享文案、推送摘要和当日展示重点"],
    ], [7.8, 7.8])

    document.add_heading("05｜公共排名如何产生", level=1)
    document.add_paragraph("公共五色排名应由确定性规则计算，不交给大模型自由判断。规则需要综合年、月、日干支、当前节气、五行旺衰以及生、克、泄、耗关系。")
    document.add_paragraph("在进入自动计算前，需要国学老师将经验整理成技术可执行的规则表，包括：")
    add_bullets(document, [
        "不同干支、节气对五行分别增加或减少多少权重。",
        "五行生克、泄耗如何转换为五色评分。",
        "总分相同时采用什么排序规则。",
        "综合排名与合作、求财、沟通等场景标签如何关联。",
        "极端日期、规则冲突和缺失数据如何处理。",
        "至少10至20个老师确认过的测试案例。",
    ])
    add_table(document, ["示例条件", "影响对象", "分值示例", "说明"], [
        ["当前节气木旺", "绿金／木", "+15", "仅作规则表结构示例，非最终算法"],
        ["当日五行生木", "绿金／木", "+20", "实际权重由老师确认"],
        ["金克木", "绿金／木", "-15", "需明确不同场景是否一致"],
    ], [5.2, 3.3, 2.3, 5.0])

    document.add_heading("06｜第一阶段：人工录入＋自动发布", level=1)
    document.add_paragraph("在老师规则尚未完全工程化之前，第一版采用人工录入每日结果，是风险最低且最快可验证产品价值的方案。")
    add_steps(document, [
        "老师提前给出未来7天或30天的五色排序和内容要点。",
        "运营人员在后台录入每种颜色的顺畅度、适合事项、阻力和行动建议。",
        "系统自动检查五种颜色是否齐全、排名是否重复、必填内容是否缺失。",
        "内容保存为草稿，经过预览与人工审核后进入待发布状态。",
        "系统在目标日期00:00按北京时间自动发布。",
        "首页只读取当天已发布版本；运营人员可以紧急修正，但必须保留修改记录。",
    ])
    document.add_paragraph("这一阶段优先验证用户是否每日查看、是否分享、是否点击香品，以及每日内容生产需要多少时间；不必等待完整规则引擎。")

    document.add_heading("07｜内容状态与后台操作", level=1)
    add_code_block(document, "草稿 draft\n    ↓\n待审核 reviewing\n    ↓\n待发布 scheduled\n    ↓\n已发布 published\n    ↓\n已撤回 withdrawn")
    add_table(document, ["后台能力", "用途"], [
        ["新建或复制", "新建某日内容，或复制前一天的页面结构作为编辑起点"],
        ["排名调整", "拖动五种颜色，明确设置第1至第5名"],
        ["内容编辑", "维护适合事项、阻力、建议、分享文案和推送摘要"],
        ["手机预览", "发布前查看用户实际看到的页面"],
        ["定时发布", "按北京时间自动发布指定日期版本"],
        ["紧急修正", "发布后修正错误，同时保存操作人、时间和修改前内容"],
        ["数据查看", "查看访问、分享、登录和商品点击等效果数据"],
    ], [4.0, 11.8])

    document.add_heading("08｜小程序首页如何获取每日内容", level=1)
    document.add_paragraph("用户每次进入“今日五色”页面时，小程序请求当天公共内容接口。后端统一按 Asia/Shanghai 时区确定日期，避免手机时区或系统时间造成内容错位。")
    add_code_block(document, "GET /api/public-guides/today")
    add_steps(document, [
        "前端先显示当天本地缓存，提升打开速度。",
        "后台请求当天最新内容及内容版本号。",
        "版本有更新时替换缓存和页面内容。",
        "跨过零点或缓存日期不是今天时，强制重新请求。",
        "公共内容完整展示，不要求用户登录。",
    ])

    document.add_heading("09｜缺失与异常兜底", level=1)
    document.add_paragraph("每日内容不能因为运营遗漏而出现空白页，也不能直接用昨天排名冒充今天结果。推荐兜底顺序：")
    add_steps(document, [
        "优先读取当天正式发布内容。",
        "没有正式内容时，读取当天已审核的备用基础模板。",
        "仍然没有时，只展示准确的日期、历法基础信息和“详细指南正在更新”提示。",
        "同时向运营后台发送内容缺失告警。",
    ])
    document.add_paragraph("用户端建议提示：今日详细指南正在更新，稍后再来看看。传统历法内容仅供文化参考。")

    document.add_heading("10｜分享卡与多平台复用", level=1)
    document.add_paragraph("每天正式发布后，系统基于同一份结构化数据生成不同渠道需要的内容，避免团队每天重复生产。")
    add_bullets(document, [
        "小程序首页完整排名与详细建议。",
        "微信小程序分享标题、摘要及五色排名图。",
        "朋友圈和微信群可转发图片。",
        "小红书图文基础素材。",
        "15至30秒短视频排名播报字幕。",
        "公众号段落与订阅消息摘要。",
    ])

    document.add_heading("11｜每日订阅提醒", level=1)
    document.add_paragraph("微信小程序不能像普通App一样默认无限发送每日推送。用户必须主动授权订阅消息，发送模板、次数和触发条件受微信平台规则约束。")
    add_steps(document, [
        "页面记录用户的订阅意愿和偏好时间，例如07:00、08:00或09:00。",
        "申请符合实际用途的微信订阅消息模板。",
        "用户主动授权后，系统记录可用发送机会。",
        "在平台允许范围内发送“今日五色已更新”提醒。",
        "用户点击消息直接进入当天公共五色页面。",
    ])
    document.add_paragraph("偏好时间只是产品设置，不能承诺微信一定每天在该时间发送；正式实现前必须按当时微信订阅消息规则验证。")

    document.add_heading("12｜AI在每日更新中的边界", level=1)
    add_table(document, ["AI可以做", "AI不能做"], [
        ["把结构化结果转成自然、易懂的大白话", "临时决定公共五色排名"],
        ["根据固定模板生成不同渠道的文案草稿", "凭空创造干支、节气或规则依据"],
        ["在禁用词和字数约束下润色建议", "生成绝对吉凶、保证发财或改命话术"],
        ["协助生成分享摘要并交给人工审核", "绕过人工审核直接发布高风险内容"],
    ], [7.8, 7.8])

    document.add_heading("13｜推荐实施顺序", level=1)
    add_table(document, ["阶段", "实现内容", "阶段目标"], [
        ["第一阶段", "人工录入、后台审核、按日期自动发布、首页读取", "先跑通每日内容生产与发布闭环"],
        ["第二阶段", "分享卡、访问统计、商品点击统计、订阅提醒", "验证自然传播和香品导流"],
        ["第三阶段", "历法引擎、五色评分规则、自动生成草稿", "降低人工计算成本并保证可追溯"],
        ["第四阶段", "AI解释、渠道文案生成、审核辅助", "提高内容生产效率但不改变规则结果"],
    ], [2.6, 8.4, 4.8])

    document.add_heading("14｜下一轮需要共同确认", level=1)
    add_bullets(document, [
        "每天五种颜色分别需要哪些必填字段。",
        "第一批每日五色原始结果由谁提供，以什么格式交付。",
        "内容需要提前准备7天还是30天。",
        "谁负责录入、谁负责审核、最晚几点完成次日内容。",
        "是否允许发布后修正，以及用户端是否显示“内容已更新”。",
        "第一版分享卡需要展示完整五色还是只展示前三名。",
        "微信订阅消息可使用的模板和实际发送限制。",
    ])

    document.add_heading("结论", level=1)
    document.add_paragraph("每日更新功能应先采用“人工提供结果、后台结构化录入、人工审核、系统定时发布”的方式上线。等用户价值和内容生产流程验证成立，再将老师规则工程化，由程序自动计算排名和生成草稿。这样既能快速推进，也能避免AI乱算、内容不可追溯和每天漏更的问题。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("五行五色财库香项目｜今日五色每日更新功能方案 V1.0")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(130, 140, 134)
    return document


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)
