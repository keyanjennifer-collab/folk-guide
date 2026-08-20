"""生成后端代码注释与待完善项审查报告。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from generate_code_guide_doc import ROOT, set_styles


OUTPUT = ROOT / "docs" / "后端代码注释与待完善项_审查报告_V1.0.docx"


def bullets(doc, items):
    """添加项目符号列表。"""
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    """添加有序列表。"""
    for item in items:
        doc.add_paragraph(item, style="List Number")


doc = Document()
set_styles(doc)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("后端代码注释与待完善项\n审查报告")
subtitle = doc.add_paragraph("五行五色财库香小程序｜V1.0｜2026年8月13日")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("一、结论", level=1)
doc.add_paragraph("当前后端已经具备可持续开发的原型基础：微信登录与JWT、本人档案、历法转换、今日五色运营、知识文档处理、向量检索接口、AI问答权益与历史均可运行。自动测试为23 passed。它适合继续开发和小范围联调，但尚不等于可直接承载真实用户和交易的生产系统。")
doc.add_paragraph("本轮已给核心后端模块补充模块说明、类和函数说明、关键判断、数据流、安全边界以及TODO（上线前）。注释重点解释“为什么这样写”和“何时必须替换”，避免逐行翻译代码。")

doc.add_heading("二、本轮已经直接修复", level=1)
bullets(doc, [
    "注销账号现在会清理正式AI问答记录和AI服务权益，不再只清理旧版聊天与每日建议。",
    "新增测试环境强制隔离：测试不读取真实微信密钥、不调用真实微信接口、不使用真实管理员Key。",
    "测试改用test_folk_guide.db，不再污染日常folk_guide.db。",
    "为JWT签发、管理员比较、微信登录、手机号换取、知识审核和向量同步补充安全原因说明。",
    "为本地向量库、模型占位、SQLite、定时发布等开发实现写明正式替换条件。",
])

doc.add_heading("三、注释覆盖的主要模块", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "模块"
table.rows[0].cells[1].text = "已说明内容"
for module, notes in [
    ("auth.py / admin_auth.py", "JWT字段、算法限制、恒定时间比较、令牌撤销和后台权限差距"),
    ("config.py / database.py", "环境变量、密钥默认值、请求会话、SQLite与PostgreSQL边界"),
    ("models.py", "用户隐私、档案、AI权益、问答历史、公共内容、知识片段和JSON字段用途"),
    ("calendar_service.py", "历法库职责、未知时辰边界、算法口径待复核项"),
    ("knowledge_*.py", "上传、去重、解析、切片、审核、停用和只检索approved资料"),
    ("vector_store.py / retrieval_service.py", "可插拔接口、RRF融合、远端一致性和任务队列需求"),
    ("ai_*.py / llm_provider.py", "安全分类、权益计算、引用、历史隔离和正式模型替换项"),
    ("public_guide_*.py", "状态机、审计、Excel导入和自动发布的生产差距"),
    ("main.py / services.py", "应用装配、微信流程、手机号、缓存、注销和CORS"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = module, notes

doc.add_heading("四、P0：真实用户上线前必须完成", level=1)
numbered(doc, [
    "生产配置校验：禁止使用默认JWT_SECRET和ADMIN_API_KEY；密钥放入服务器密钥管理，启动时检测弱配置并拒绝启动。",
    "后台权限系统：共享Admin Key改为管理员账号、角色权限、登录过期、操作审计和必要的双因素验证。",
    "数据库生产化：SQLite迁移PostgreSQL，启用外键，使用Alembic版本化迁移，并建立备份、恢复和监控。",
    "HTTPS和域名：API部署到备案HTTPS域名，微信后台配置request合法域名；真机不能使用127.0.0.1。",
    "JWT撤销能力：加入jti或token_version，账号注销、风控或密钥变更后让旧令牌立即失效。",
    "隐私和数据治理：补充授权记录、隐私政策版本、档案和手机号用途、删除流程、保留期限及订单法定留存规则。",
    "接口防滥用：微信登录、AI问答、上传和管理接口加入限流、异常告警和结构化审计日志。",
    "统一时间：数据库使用带时区UTC时间，业务按Asia/Shanghai计算每日次数和发布时间。",
])

doc.add_heading("五、P1：接入正式AI前必须完成", level=1)
numbered(doc, [
    "持久化向量库：MemoryVectorStore重启会清空，不能用于生产；应接入已选供应商并保存索引。",
    "向量任务队列：审核通过后自动排队向量化，记录pending/running/failed，提供有限重试和失败人工处理。",
    "同步一致性：处理远端写入成功但数据库提交失败、停用删除失败等情况，增加补偿任务和定期对账。",
    "Embedding和模型适配器：加入超时、有限重试、并发控制、费用统计、熔断和无模型时的降级回答。",
    "安全体系：当前关键词规则只是一层硬边界，还要补同义词、上下文分类、提示注入防护、输出审核和人工抽检。",
    "引用可信度：检查生成答案是否得到检索片段支持，资料不足时拒绝推测，并提供可查看的原文引用。",
    "次数并发原子性：当前先统计再写消息，并发请求可能同时通过限额；正式版需要数据库原子扣减或Redis计数。",
    "AI前端联调：小程序AI页面仍有模拟内容，需要连接quota/chat/history/feedback真实接口并处理403、429和网络失败。",
])

doc.add_heading("六、P1：微信和账号功能待完善", level=1)
bullets(doc, [
    "微信access_token目前每次绑定手机号都重新获取，应缓存并在失效时刷新，多实例使用共享缓存。",
    "需要在真实开发者工具中验证code2session错误码、网络超时和AppSecret重置场景。",
    "手机号是否唯一、换绑方式、一个号码能否关联多个微信账号，需要先确定业务规则再加数据库约束。",
    "JWT当前有效期7天，没有刷新令牌；可根据体验与风险设计短Access Token加Refresh Token。",
    "微信不会自动返回昵称头像；如确有需要，应让用户主动填写或按微信现行能力单独授权。",
])

doc.add_heading("七、P2：稳定性和维护性优化", level=1)
bullets(doc, [
    "今日五色到期发布当前由读取接口触发，应迁移到独立调度任务并加数据库锁。",
    "知识原文件当前保存在本机，应迁移对象存储，增加病毒扫描、孤儿文件清理和备份。",
    "关键词检索使用SQLite contains，资料变多后应使用PostgreSQL全文检索或独立搜索服务。",
    "历史列表改为游标分页；知识文档和运营列表也应加分页。",
    "大量datetime.utcnow调用已出现弃用警告，应在时区改造时统一替换。",
    "旧版/api/chat与正式/api/ai/chat并存，前端切换后应废弃旧接口并安排兼容期。",
    "每日个人建议目前仍是确定性演示逻辑，真实规则、版本、审核和回溯机制尚未建立。",
    "日志需要脱敏：不得记录AppSecret、完整JWT、openid、手机号、生辰详情和完整用户问题。",
])

doc.add_heading("八、推荐实施顺序", level=1)
numbered(doc, [
    "先完成AI国学前端与现有真实后端接口联调。",
    "确定大模型、Embedding和向量数据库供应商。",
    "实现持久化向量适配器、任务队列、失败重试和对账。",
    "接入正式模型并完成安全、引用和费用控制。",
    "上线准备时集中完成PostgreSQL、Alembic、HTTPS、权限、限流、监控和隐私合规。",
    "最后再接订单、微信支付和真实服务卡，避免支付建立在不稳定的账号与数据基础上。",
])

doc.add_heading("九、如何阅读代码中的TODO", level=1)
doc.add_paragraph("搜索“TODO（上线前）”可以找到生产安全差距；搜索“TODO（接入向量供应商）”或“TODO（接入模型供应商）”可以找到外部AI服务接入点。TODO不是当前代码报错，而是明确记录当前原型在什么条件下必须升级。")

doc.add_heading("十、验证结果", level=1)
bullets(doc, [
    "后端自动测试：23 passed。",
    "测试使用独立SQLite数据库和模拟微信环境。",
    "测试流程不会读取、打印或覆盖真实AppSecret。",
    "新增注销测试验证用户、AI权益和正式AI问答记录均被清理。",
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
