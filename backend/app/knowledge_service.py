"""知识文档解析、清洗、切片和关键词检索。"""

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from docx import Document as DocxDocument
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .config import get_settings
from .models import KnowledgeChunk, KnowledgeDocument
from .pdf_service import extract_pdf


ALLOWED_EXTENSIONS = {".docx", ".md", ".txt", ".pdf"}
# 公共知识库只接受版权状态明确的资料；用户聊天和生辰资料不得走此导入流程。
COPYRIGHT_VALUES = {"public_domain", "licensed", "original", "permission_pending"}
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_PDF_FILE_SIZE = 100 * 1024 * 1024
TARGET_CHARS = 700
MIN_CHARS = 250


@dataclass
class ParsedBlock:
    """解析阶段的中间结构：标题、正文或表格。"""
    kind: str
    text: str
    heading_level: int | None = None
    # PDF保留物理页码；Word、Markdown和TXT没有可靠页码，因此为None。
    page_number: int | None = None
    extraction_method: str | None = None
    ocr_confidence: float | None = None


@dataclass
class ParsedDocument:
    """一份文件的结构块和解析质量元数据。"""
    blocks: list[ParsedBlock]
    extraction_method: str
    page_count: int | None = None
    needs_ocr: bool = False
    unreadable_pages: list[int] | None = None
    warning: str | None = None


@dataclass(frozen=True)
class ChunkDraft:
    """准备写入数据库的切片，包含PDF来源页和提取方式。"""
    heading: str | None
    content: str
    page_start: int | None = None
    page_end: int | None = None
    extraction_method: str | None = None
    ocr_confidence: float | None = None


def sha256_bytes(content: bytes) -> str:
    """计算文件指纹，用于识别完全相同的重复上传。"""
    return hashlib.sha256(content).hexdigest()


def normalize_text(value: str) -> str:
    """统一空白字符，避免同一内容因排版差异产生不同切片。"""
    value = value.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def parse_docx(path: Path) -> list[ParsedBlock]:
    """按 Word 标题、正文、表格提取结构块，不处理图片 OCR。

    TODO：若运营资料大量来自扫描件，再单独增加PDF/OCR流水线和人工校对状态；
    不应在这里悄悄加入低准确率OCR，避免错误资料进入公共知识库。
    """
    try:
        document = DocxDocument(path)
    except (BadZipFile, ValueError) as exc:
        raise ValueError("Word文件损坏或不是有效的.docx") from exc
    blocks: list[ParsedBlock] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
        if match:
            blocks.append(ParsedBlock("heading", text, min(int(match.group(1)), 6), extraction_method="docx"))
        else:
            blocks.append(ParsedBlock("paragraph", text, extraction_method="docx"))
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [normalize_text(cell.text).replace("\n", " / ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(ParsedBlock("table", "\n".join(rows), extraction_method="docx"))
    return blocks


def parse_markdown(path: Path) -> list[ParsedBlock]:
    """识别 Markdown 的 # 标题，空行作为普通段落边界。"""
    blocks: list[ParsedBlock] = []
    buffer: list[str] = []
    def flush():
        """把当前Markdown普通段落缓冲区写成一个结构块。"""
        if buffer:
            text = normalize_text("\n".join(buffer))
            if text:
                blocks.append(ParsedBlock("paragraph", text, extraction_method="markdown"))
            buffer.clear()
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if match:
            flush()
            blocks.append(ParsedBlock(
                "heading", normalize_text(match.group(2)), len(match.group(1)),
                extraction_method="markdown",
            ))
        elif not line.strip():
            flush()
        else:
            buffer.append(line)
    flush()
    return blocks


def parse_text(path: Path) -> list[ParsedBlock]:
    """按空行拆分 UTF-8 文本；不猜测不明确的旧编码。"""
    raw = path.read_text(encoding="utf-8-sig")
    return [
        ParsedBlock("paragraph", part, extraction_method="text")
        for part in (normalize_text(x) for x in re.split(r"\n\s*\n", raw)) if part
    ]


def parse_file(path: Path, pdf_parser_mode: str = "auto") -> ParsedDocument:
    """根据扩展名分派解析器，并统一返回解析质量元数据。"""
    extension = path.suffix.lower()
    if extension == ".docx":
        return ParsedDocument(parse_docx(path), "docx")
    if extension == ".md":
        return ParsedDocument(parse_markdown(path), "markdown")
    if extension == ".txt":
        return ParsedDocument(parse_text(path), "text")
    if extension == ".pdf":
        settings = get_settings()
        result = extract_pdf(
            path,
            parser_mode=pdf_parser_mode,
            max_pages=settings.pdf_max_pages,
            min_text_chars=settings.pdf_min_text_chars_per_page,
            mineru_enabled=settings.mineru_enabled,
            mineru_base_url=settings.mineru_base_url,
            mineru_timeout_seconds=settings.mineru_timeout_seconds,
        )
        blocks = [ParsedBlock(
            kind=block.kind,
            text=block.text,
            heading_level=block.heading_level,
            page_number=block.page_number,
            extraction_method=block.extraction_method,
            ocr_confidence=block.ocr_confidence,
        ) for block in result.blocks]
        return ParsedDocument(
            blocks=blocks,
            extraction_method=result.extraction_method,
            page_count=result.page_count,
            needs_ocr=bool(result.scanned_pages),
            unreadable_pages=result.scanned_pages,
            warning="；".join(result.warnings) or None,
        )
    raise ValueError("仅支持.docx、.md、.txt和.pdf")


def split_long_text(text: str, target: int = TARGET_CHARS) -> list[str]:
    """优先按中文标点断句，单句过长时才按字符数硬切。"""
    if len(text) <= target:
        return [text]
    # 中文句末标点优先作为边界，既减少语义截断，也让引用片段更自然。
    sentences = [part.strip() for part in re.split(r"(?<=[。！？；\n])", text) if part.strip()]
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        if current and len(current) + len(sentence) > target:
            chunks.append(current)
            current = ""
        if len(sentence) > target:
            for start in range(0, len(sentence), target):
                part = sentence[start:start + target]
                if current:
                    chunks.append(current)
                    current = ""
                chunks.append(part)
        else:
            current += sentence
    if current:
        chunks.append(current)
    return chunks


def chunk_blocks(blocks: list[ParsedBlock]) -> list[ChunkDraft]:
    """尽量在章节和句子边界切片，并保留章节名、PDF页码和提取方式。"""
    result: list[ChunkDraft] = []
    current_heading: str | None = None
    current_parts: list[tuple[str, int | None, str | None, float | None]] = []
    current_length = 0
    def flush():
        """结束当前切片，并重置累计字符数。"""
        nonlocal current_parts, current_length
        if current_parts:
            pages = [page for _, page, _, _ in current_parts if page is not None]
            methods = {method for _, _, method, _ in current_parts if method}
            confidences = [score for _, _, _, score in current_parts if score is not None]
            result.append(ChunkDraft(
                heading=current_heading,
                content="\n\n".join(part for part, _, _, _ in current_parts),
                page_start=min(pages) if pages else None,
                page_end=max(pages) if pages else None,
                extraction_method=(next(iter(methods)) if len(methods) == 1 else "mixed") if methods else None,
                # 一个片段跨越多个OCR元素时取最低置信度，避免平均值掩盖低质量内容。
                ocr_confidence=min(confidences) if confidences else None,
            ))
            current_parts = []
            current_length = 0
    for block in blocks:
        if block.kind == "heading":
            flush()
            current_heading = block.text
            continue
        for part in split_long_text(block.text):
            if current_parts and current_length + len(part) > TARGET_CHARS and current_length >= MIN_CHARS:
                flush()
            current_parts.append((part, block.page_number, block.extraction_method, block.ocr_confidence))
            current_length += len(part)
    flush()
    return result


def replace_chunks(db: Session, document: KnowledgeDocument, chunks: list[ChunkDraft]) -> None:
    """删除该文档旧切片并写入新切片；新切片统一等待人工审核。"""
    # 重新解析属于全量替换：旧向量随旧片段一起失效，之后必须重新审核、同步。
    db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id).delete()
    for index, chunk in enumerate(chunks):
        normalized = normalize_text(chunk.content)
        db.add(KnowledgeChunk(
            document_id=document.id,
            chunk_index=index,
            heading=chunk.heading,
            content=normalized,
            content_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
            character_count=len(normalized),
            review_status="pending",
            page_start=chunk.page_start,
            page_end=chunk.page_end,
            extraction_method=chunk.extraction_method,
            ocr_confidence=chunk.ocr_confidence,
        ))


def document_output(db: Session, document: KnowledgeDocument) -> dict:
    """组装文档返回值，并实时统计总片段和审核通过片段数量。"""
    chunk_count = db.scalar(select(func.count(KnowledgeChunk.id)).where(KnowledgeChunk.document_id == document.id)) or 0
    approved = db.scalar(select(func.count(KnowledgeChunk.id)).where(
        KnowledgeChunk.document_id == document.id,
        KnowledgeChunk.review_status == "approved",
    )) or 0
    return {
        "id": document.id, "title": document.title, "source_name": document.source_name,
        "author": document.author, "copyright_status": document.copyright_status,
        "original_filename": document.original_filename, "file_hash": document.file_hash,
        "file_size": document.file_size, "mime_type": document.mime_type, "status": document.status,
        "page_count": document.page_count, "extraction_method": document.extraction_method,
        "needs_ocr": document.needs_ocr, "unreadable_pages_json": document.unreadable_pages_json,
        "parse_warning": document.parse_warning,
        "version": document.version, "error_message": document.error_message,
        "reviewed_by": document.reviewed_by, "reviewed_at": document.reviewed_at,
        "chunk_count": chunk_count, "approved_chunk_count": approved,
        "created_at": document.created_at, "updated_at": document.updated_at,
    }


def keyword_search(db: Session, query: str, limit: int = 8) -> list[dict]:
    """SQLite 阶段的基础关键词召回，只检索已审核且未停用的公共资料。"""
    terms = [term for term in re.split(r"\s+", query.strip()) if term]
    if not terms:
        return []
    # SQLite contains只提供基础子串匹配。
    # TODO（上线前）：PostgreSQL全文检索或独立搜索服务，加入分词、同义词和索引。
    conditions = [KnowledgeChunk.content.contains(term) for term in terms]
    rows = db.execute(
        select(KnowledgeChunk, KnowledgeDocument)
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeChunk.document_id)
        .where(
            KnowledgeChunk.review_status == "approved",
            KnowledgeDocument.status == "approved",
            or_(*conditions),
        )
        .limit(limit)
    ).all()
    results = []
    for chunk, document in rows:
        matches = sum(chunk.content.count(term) for term in terms)
        results.append({
            "chunk_id": chunk.id, "document_id": document.id, "title": document.title,
            "heading": chunk.heading, "content": chunk.content, "source_name": document.source_name,
            "page_start": chunk.page_start, "page_end": chunk.page_end,
            "score": float(matches) / max(len(chunk.content), 1) * 1000,
        })
    return sorted(results, key=lambda item: item["score"], reverse=True)[:limit]
"""知识文档解析、清洗、切片和关键词检索的纯业务函数。"""
